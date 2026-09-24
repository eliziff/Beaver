"""Helpers for what python-docx lacks natively. Every helper edits ordinary OOXML;
in Review mode the recorder turns the result into native revisions afterwards."""
from __future__ import annotations
import copy
import re

from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import XmlPart
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Mm
from docx.table import Table, _Row
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from ooxml import XMLNS, accepted_text

W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
NSDECL = 'xmlns:w="%s"' % W_NS
DOC = None   # set by the runner


def _el(obj):
    for attr in ('_p', '_tbl', '_r', '_tr', '_tc', '_element'):
        if hasattr(obj, attr): return getattr(obj, attr)
    return obj


def text(obj) -> str:
    """Accepted-view text of a paragraph, table, run, section header/footer or element."""
    return accepted_text(_el(obj))


def blocks(container=None) -> list:
    """Body paragraphs and tables in document order (also for a cell or header)."""
    container = container or DOC
    return list(container.iter_inner_content())


def find(needle: str, container=None, regex=False) -> list[Paragraph]:
    """Paragraphs (including table cells, in order) whose accepted text contains needle (or matches a regex)."""
    container = container or DOC
    match = (lambda s: re.search(needle, s)) if regex else (lambda s: needle in s)
    root = DOC.element.body if container is DOC else _el(container)
    parent = DOC._body if container is DOC else container
    return [Paragraph(p, parent) for p in root.iter(qn('w:p')) if match(accepted_text(p))]


def isolate(paragraph: Paragraph, needle: str, occurrence: int = 0) -> list[Run]:
    """Split runs so that exactly `needle` (the n-th occurrence) is covered by whole runs; returns them.
    Searches the accepted view: runs in insertions and hyperlinks count, deleted runs do not."""
    runs = [r for r in paragraph._p.xpath('.//w:r[not(ancestor::w:del or ancestor::w:moveFrom)]')
            if r.xpath('ancestor::w:p[1]')[0] is paragraph._p and (r.xpath('./w:t|./w:tab|./w:br') or not r.xpath('./*[not(self::w:rPr)]'))]
    full, spans = '', []
    for r in runs:
        t = Run(r, paragraph).text
        spans.append((len(full), len(full) + len(t), r)); full += t
    starts = [m.start() for m in re.finditer(re.escape(needle), full)]
    if len(starts) <= occurrence: raise ValueError('Text not found in paragraph: %r' % needle[:80])
    start, end = starts[occurrence], starts[occurrence] + len(needle)

    def split(r, at):
        """Split run r at offset `at` inside it; returns the right half."""
        left_text = Run(r, paragraph).text
        right = copy.deepcopy(r); r.addnext(right)
        Run(r, paragraph).text, Run(right, paragraph).text = left_text[:at], left_text[at:]
        return right
    out = []
    for s, e, r in spans:
        if e <= start or s >= end: continue
        if s < start: r = split(r, start - s); s = start
        if e > end: split(r, end - s)
        out.append(Run(r, paragraph))
    return out


def replace_text(paragraph: Paragraph, old: str, new: str, count: int = 0) -> int:
    """Replace text in a paragraph, keeping the formatting of the first replaced run. count=0 replaces all."""
    done = 0
    while not count or done < count:
        try: runs = isolate(paragraph, old, done * new.count(old))
        except ValueError: break
        runs[0].text = new
        for r in runs[1:]: r._r.getparent().remove(r._r)
        done += 1
    return done


def delete(obj) -> None:
    """Remove a paragraph, table, row or run (Review mode records a tracked deletion)."""
    el = _el(obj)
    el.getparent().remove(el)


def insert_paragraph_after(anchor, text: str = '', style=None) -> Paragraph:
    """New paragraph after a paragraph or table."""
    p = OxmlElement('w:p'); _el(anchor).addnext(p)
    para = Paragraph(p, anchor._parent)
    if text: para.add_run(text)
    if style is not None: para.style = style
    return para


def move_after(obj, anchor):
    """Move a paragraph/table (e.g. one created with doc.add_table) to follow anchor."""
    _el(anchor).addnext(_el(obj)); return obj


def insert_table_after(anchor, rows: int, cols: int, style=None) -> Table:
    """New table after anchor; style 'Table Grid' is created when the document lacks it."""
    table = DOC.add_table(rows=rows, cols=cols)
    if style == 'Table Grid': table_grid_style()
    if style is not None: table.style = style
    return move_after(table, anchor)


def table_grid_style():
    """Add Word's built-in 'Table Grid' style (single borders) if missing."""
    styles = DOC.styles.element
    if styles.xpath('./w:style[@w:styleId="TableGrid"]'): return
    borders = ''.join('<w:%s w:val="single" w:sz="4" w:space="0" w:color="auto"/>' % side
                      for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'))
    styles.append(parse_xml(
        '<w:style %s w:type="table" w:styleId="TableGrid"><w:name w:val="Table Grid"/><w:basedOn w:val="TableNormal"/>'
        '<w:uiPriority w:val="39"/><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr>'
        '<w:tblPr><w:tblBorders>%s</w:tblBorders></w:tblPr></w:style>' % (NSDECL, borders)))


LIST_KINDS = {  # numFmt and label per level (cycled); '#' is the level's own number
    'decimal': [('decimal', '#.'), ('lowerLetter', '(#)'), ('lowerRoman', '(#)')],
    'upperLetter': [('upperLetter', '#.'), ('decimal', '(#)'), ('lowerRoman', '(#)')],
    'lowerLetter': [('lowerLetter', '(#)'), ('lowerRoman', '(#)'), ('decimal', '#.')],
    'upperRoman': [('upperRoman', '#.'), ('upperLetter', '#.'), ('decimal', '(#)')],
    'bullet': [('bullet', '•'), ('bullet', '◦'), ('bullet', '▪')],
}


def numbered_list(items, kind: str = 'decimal', start: int = 1, indent_mm: float = 6.35) -> int:
    """Make paragraphs ONE list via numbering.xml (abstractNum + num). items: paragraphs or (paragraph, level) pairs.
    kind: decimal | legal (1. / 1.1. / 1.1.1.) | bullet | upperLetter | lowerLetter | upperRoman. Returns the numId."""
    numbering = DOC.part.numbering_part.element
    abstract_id = max([int(x) for x in numbering.xpath('./w:abstractNum/@w:abstractNumId')] + [-1]) + 1
    num_id = max([int(x) for x in numbering.xpath('./w:num/@w:numId')] + [0]) + 1
    levels = []
    for level in range(9):
        if kind == 'legal': fmt, label = 'decimal', '.'.join('%%%d' % (i + 1) for i in range(level + 1)) + '.'
        else:
            fmt, label = LIST_KINDS[kind][level % 3]
            label = label.replace('#', '%%%d' % (level + 1))
        twips = lambda mm: int(round(mm * 56.6929))
        levels.append('<w:lvl w:ilvl="%d"><w:start w:val="%d"/><w:numFmt w:val="%s"/><w:lvlText w:val="%s"/><w:lvlJc w:val="left"/>'
                      '<w:pPr><w:ind w:left="%d" w:hanging="%d"/></w:pPr></w:lvl>'
                      % (level, start if level == 0 else 1, fmt, label, twips(indent_mm * (level + 1)), twips(indent_mm)))
    abstract = parse_xml('<w:abstractNum %s w:abstractNumId="%d"><w:multiLevelType w:val="multilevel"/>%s</w:abstractNum>'
                         % (NSDECL, abstract_id, ''.join(levels)))
    existing = numbering.xpath('./w:abstractNum')
    existing[-1].addnext(abstract) if existing else numbering.insert(0, abstract)
    num = parse_xml('<w:num %s w:numId="%d"><w:abstractNumId w:val="%d"/></w:num>' % (NSDECL, num_id, abstract_id))
    nums = numbering.xpath('./w:num')
    nums[-1].addnext(num) if nums else abstract.addnext(num)
    for item in items:
        para, level = item if isinstance(item, tuple) else (item, 0)
        numpr = para._p.get_or_add_pPr().get_or_add_numPr()
        numpr.get_or_add_ilvl().val = level
        numpr.get_or_add_numId().val = num_id
    return num_id


PAPER = {'letter': (215.9, 279.4), 'legal': (215.9, 355.6), 'a4': (210, 297), 'a3': (297, 420)}


def set_page(section, orientation=None, size=None, margins_mm=None, columns=None):
    """orientation 'portrait'|'landscape' swaps width/height correctly; size 'letter'|'legal'|'a4'|'a3'|(w_mm,h_mm);
    margins_mm: number or dict(top,bottom,left,right,header,footer,gutter); columns: int."""
    from docx.enum.section import WD_ORIENT
    w_mm, h_mm = (PAPER[size] if isinstance(size, str) else size) if size else (section.page_width.mm, section.page_height.mm)
    if orientation:
        landscape = orientation == 'landscape'
        w_mm, h_mm = (max(w_mm, h_mm), min(w_mm, h_mm)) if landscape else (min(w_mm, h_mm), max(w_mm, h_mm))
        section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = Mm(w_mm), Mm(h_mm)
    if margins_mm is not None:
        values = margins_mm if isinstance(margins_mm, dict) else dict(top=margins_mm, bottom=margins_mm, left=margins_mm, right=margins_mm)
        for side, value in values.items(): setattr(section, side + ('_margin' if side in ('top', 'bottom', 'left', 'right') else '_distance' if side in ('header', 'footer') else ''), Mm(value))
    if columns:
        sect = section._sectPr
        cols = sect.find(qn('w:cols'))
        if cols is None:
            cols = OxmlElement('w:cols')
            later = {qn('w:' + n) for n in ('formProt', 'vAlign', 'noEndnote', 'titlePg', 'textDirection', 'bidi', 'rtlGutter', 'docGrid', 'printerSettings', 'sectPrChange')}
            follow = next((c for c in sect if c.tag in later), None)
            follow.addprevious(cols) if follow is not None else sect.append(cols)
        cols.set(qn('w:num'), str(columns)); cols.set(qn('w:space'), '720')
    return section


def section_range(first, last, start: str = 'new_page'):
    """Make blocks first..last (paragraphs/tables, in order) their own section; returns its Section for set_page().
    start: new_page | continuous | odd_page | even_page. Headers/footers stay shared."""
    from docx.enum.section import WD_SECTION_START
    from docx.section import Section
    def ends_section(el):
        return el is not None and el.tag == qn('w:p') and el.find(qn('w:pPr') + '/' + qn('w:sectPr')) is not None
    def break_after(el):
        sect = copy.deepcopy(_governing(el))
        holder = el if el.tag == qn('w:p') else None
        if holder is None: holder = OxmlElement('w:p'); el.addnext(holder)
        holder.get_or_add_pPr().insert_element_before(sect, 'w:pPrChange')
        return sect
    first_el, last_el = _el(first), _el(last)
    prev = first_el.getprevious()
    if prev is not None and not ends_section(prev): break_after(prev)
    nxt = last_el.getnext()
    # A table's section already ends at a following empty section-break paragraph (e.g. from an earlier call).
    closed = ends_section(last_el) or last_el.tag == qn('w:tbl') and ends_section(nxt) and not nxt.xpath('string(.)')
    sect = break_after(last_el) if nxt is not None and nxt.tag != qn('w:sectPr') and not closed else _governing(last_el)
    section = Section(sect, DOC.part)
    section.start_type = getattr(WD_SECTION_START, start.upper())
    return section


def _governing(el):
    """The sectPr that governs block el."""
    node = el
    while node is not None:
        if node.tag == qn('w:p'):
            s = node.find(qn('w:pPr') + '/' + qn('w:sectPr'))
            if s is not None: return s
        node = node.getnext()
    return DOC.element.body.find(qn('w:sectPr'))


def add_field(paragraph: Paragraph, instr: str, result: str = '') -> list[Run]:
    """Append a complex field (e.g. 'PAGE', 'NUMPAGES', 'TOC \\o "1-3" \\h \\z \\u', 'DATE \\@ "d MMMM yyyy"', 'REF bm \\h')."""
    runs = []
    for part in ('begin', instr, 'separate', result, 'end'):
        r = paragraph.add_run()._r
        if part in ('begin', 'separate', 'end'):
            fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), part); r.append(fc)
            if part == 'begin': fc.set(qn('w:dirty'), 'true')
        elif part is instr:
            it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = ' %s ' % instr; r.append(it)
        else:
            t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve'); t.text = result; r.append(t)
        runs.append(Run(r, paragraph))
    return runs


def update_fields_on_open():
    """Ask Word to refresh fields (TOC, page refs) when the file is opened."""
    settings = DOC.settings.element
    if settings.find(qn('w:updateFields')) is None:
        el = OxmlElement('w:updateFields'); el.set(qn('w:val'), 'true')
        _settings_insert(settings, el, SETTINGS_AFTER_UPDATE)


def add_toc(anchor, levels: str = '1-3', title: str | None = 'Contents') -> Paragraph:
    """Insert an optional title and a TOC field after anchor (paragraph/table). Word fills it on open."""
    if title:
        heading = insert_paragraph_after(anchor, title)
        try: heading.style = DOC.styles['TOC Heading']
        except KeyError: pass
        anchor = heading
    p = insert_paragraph_after(anchor)
    add_field(p, 'TOC \\o "%s" \\h \\z \\u' % levels, 'Update field to build the table of contents.')
    update_fields_on_open()
    return p


def page_numbers(section, template: str = 'Page {PAGE} of {NUMPAGES}', align: str = 'center', where: str = 'footer') -> Paragraph:
    """Replace the section's header/footer text with template; {PAGE}/{NUMPAGES}/{SECTIONPAGES} become fields."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    story = getattr(section, where)
    story.is_linked_to_previous = False
    p = story.paragraphs[0] if story.paragraphs else story.add_paragraph()
    for extra in story.paragraphs[1:]: delete(extra)
    for child in [c for c in p._p if c.tag != qn('w:pPr')]: p._p.remove(child)
    for piece in re.split(r'(\{[A-Z]+\})', template):
        if re.fullmatch(r'\{[A-Z]+\}', piece): add_field(p, piece[1:-1], '1')
        elif piece: p.add_run(piece)
    p.alignment = getattr(WD_ALIGN_PARAGRAPH, align.upper())
    return p


NOTE_CT = {'footnote': 'application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml',
           'endnote': 'application/vnd.openxmlformats-officedocument.wordprocessingml.endnotes+xml'}


def _notes_part(kind):
    for rel in DOC.part.rels.values():
        if rel.reltype == getattr(RT, kind.upper() + 'S') and not rel.is_external: return rel.target_part
    xml = ('<w:{0}s {1}><w:{0} w:type="separator" w:id="-1"><w:p><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr>'
           '<w:r><w:separator/></w:r></w:p></w:{0}><w:{0} w:type="continuationSeparator" w:id="0"><w:p><w:pPr><w:spacing w:after="0" '
           'w:line="240" w:lineRule="auto"/></w:pPr><w:r><w:continuationSeparator/></w:r></w:p></w:{0}></w:{0}s>').format(kind, NSDECL)
    part = XmlPart(PackURI('/word/%ss.xml' % kind), NOTE_CT[kind], parse_xml(xml), DOC.part.package)
    DOC.part.relate_to(part, getattr(RT, kind.upper() + 'S'))
    return part


def add_footnote(paragraph: Paragraph, note: str, after: str | None = None, kind: str = 'footnote') -> int:
    """Add a footnote (or kind='endnote') whose reference follows `after` text in paragraph (default: its end)."""
    part = _notes_part(kind)
    root = part.element
    note_id = max([int(x) for x in root.xpath('./w:%s/@w:id' % kind, namespaces={'w': W_NS})] + [0]) + 1
    root.append(parse_xml(
        '<w:{0} {1} w:id="{2}"><w:p><w:pPr><w:pStyle w:val="{3}Text"/></w:pPr><w:r><w:rPr><w:vertAlign w:val="superscript"/></w:rPr>'
        '<w:{0}Ref/></w:r><w:r><w:t xml:space="preserve"> </w:t></w:r><w:r><w:t xml:space="preserve">{4}</w:t></w:r></w:p></w:{0}>'
        .format(kind, NSDECL, note_id, kind.capitalize(), _escape(note))))
    ref = parse_xml('<w:r %s><w:rPr><w:vertAlign w:val="superscript"/></w:rPr><w:%sReference w:id="%d"/></w:r>' % (NSDECL, kind, note_id))
    if after: isolate(paragraph, after)[-1]._r.addnext(ref)
    else: paragraph._p.append(ref)
    return note_id


def content_control(obj, tag: str | None = None, alias: str | None = None, placeholder: bool = False):
    """Wrap a paragraph (block control) or a list of runs (inline control) in a plain-text content control."""
    sdt = OxmlElement('w:sdt'); pr = OxmlElement('w:sdtPr')
    for name, value in (('w:alias', alias), ('w:tag', tag)):
        if value: el = OxmlElement(name); el.set(qn('w:val'), value); pr.append(el)
    pr.append(OxmlElement('w:text')) if not isinstance(obj, Paragraph) else None
    content = OxmlElement('w:sdtContent'); sdt.append(pr); sdt.append(content)
    elements = [_el(obj)] if isinstance(obj, (Paragraph, Table)) else [_el(r) for r in obj]
    elements[0].addprevious(sdt)
    for el in elements: content.append(el)
    return sdt


def xml(obj, limit: int = 6000) -> str:
    """OOXML of a paragraph, run, table, section or element, without namespace declarations."""
    from lxml import etree
    return XMLNS.sub('', etree.tostring(_el(obj), pretty_print=True, encoding='unicode'))[:limit]


def _escape(s):
    return s.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')


SETTINGS_AFTER_TRACK = ['doNotTrackMoves', 'doNotTrackFormatting', 'documentProtection', 'autoFormatOverride', 'styleLockTheme',
                        'styleLockQFSet', 'defaultTabStop', 'autoHyphenation', 'consecutiveHyphenLimit', 'hyphenationZone',
                        'doNotHyphenateCaps', 'showEnvelope', 'summaryLength', 'clickAndTypeStyle', 'defaultTableStyle',
                        'evenAndOddHeaders', 'bookFoldRevPrinting', 'bookFoldPrinting', 'bookFoldPrintingSheets',
                        'drawingGridHorizontalSpacing', 'drawingGridVerticalSpacing', 'displayHorizontalDrawingGridEvery',
                        'displayVerticalDrawingGridEvery', 'doNotUseMarginsForDrawingGridOrigin', 'drawingGridHorizontalOrigin',
                        'drawingGridVerticalOrigin', 'doNotShadeFormData', 'noPunctuationKerning', 'characterSpacingControl',
                        'printTwoOnOne', 'strictFirstAndLastChars', 'noLineBreaksAfter', 'noLineBreaksBefore', 'savePreviewPicture',
                        'doNotValidateAgainstSchema', 'saveInvalidXml', 'ignoreMixedContent', 'alwaysShowPlaceholderText',
                        'doNotDemarcateInvalidXml', 'saveXmlDataOnly', 'useXSLTWhenSaving', 'saveThroughXslt', 'showXMLTags',
                        'alwaysMergeEmptyNamespace', 'updateFields', 'hdrShapeDefaults', 'footnotePr', 'endnotePr', 'compat',
                        'docVars', 'rsids', 'mathPr', 'attachedSchema', 'themeFontLang', 'clrSchemeMapping',
                        'doNotIncludeSubdocsInStats', 'doNotAutoCompressPictures', 'forceUpgrade', 'captions',
                        'readModeInkLockDown', 'smartTagType', 'schemaLibrary', 'shapeDefaults', 'doNotEmbedSmartTags',
                        'decimalSymbol', 'listSeparator']
SETTINGS_AFTER_UPDATE = SETTINGS_AFTER_TRACK[SETTINGS_AFTER_TRACK.index('updateFields') + 1:]


def _settings_insert(settings, el, successors):
    names = {qn('w:' + n) for n in successors}
    follow = next((c for c in settings if c.tag in names), None)
    follow.addprevious(el) if follow is not None else settings.append(el)


def track_revisions_on():
    settings = DOC.settings.element
    if settings.find(qn('w:trackRevisions')) is None:
        _settings_insert(settings, OxmlElement('w:trackRevisions'), SETTINGS_AFTER_TRACK)
