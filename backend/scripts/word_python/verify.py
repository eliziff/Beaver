"""Trusted inspect and candidate verification. Never executes model code.

  verify.py inspect <docx>                      (stdin: {offset, limit, target})
  verify.py verify <source> <candidate> <mode>  (mode: tracked | direct)
"""
from __future__ import annotations
import difflib
import json
import os
import re
import sys
from hashlib import sha256

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from lxml import etree  # noqa: E402

import ooxml  # noqa: E402
from ooxml import W_NS, XMLNS, accepted_text, blocks, canon, main_part, parse, related, resolve, revision_counts, stories, story_root, w  # noqa: E402

EMU_PER_MM = 36000
TWIP_MM = 25.4 / 1440


def style_names(parts):
    path = related(parts, '/styles')
    if not path: return {}
    return {s.get(w('styleId')): (s.find(w('name')).get(w('val')) if s.find(w('name')) is not None else s.get(w('styleId')))
            for s in parse(parts[path]).iter(w('style'))}


def section_info(sect):
    def attr(tag, name):
        el = sect.find(w(tag))
        value = None if el is None else el.get(w(name))
        return None if value is None else value
    mm = lambda v: None if v is None else round(int(v) * TWIP_MM, 1)
    return {k: v for k, v in {
        'orientation': attr('pgSz', 'orient') or 'portrait',
        'page_mm': [mm(attr('pgSz', 'w')), mm(attr('pgSz', 'h'))],
        'margins_mm': {s: mm(attr('pgMar', s)) for s in ('top', 'bottom', 'left', 'right')},
        'start': attr('type', 'val') or 'nextPage',
        'columns': attr('cols', 'num'),
        'headers': [r.get(w('type')) for r in sect.findall(w('headerReference'))] or None,
        'footers': [r.get(w('type')) for r in sect.findall(w('footerReference'))] or None,
        'title_page': True if sect.find(w('titlePg')) is not None else None,
    }.items() if v is not None}


def body_sections(body):
    return [s for s in body.iter(w('sectPr')) if s.getparent().tag in (w('pPr'), w('body'))]


def inspect(path, request):
    parts = ooxml.read_package(path)
    main = main_part(parts)
    root = parse(parts[main])
    body = root.find(w('body'))
    names = style_names(parts)
    counts = revision_counts(root)
    offset, limit = int(request.get('offset') or 0), int(request.get('limit') or 40)
    target = request.get('target')
    tracked = bool(counts)
    paragraphs = body.findall(w('p')); tables = body.findall(w('tbl'))
    if target:
        kind, _, index = target.partition(':')
        if kind == 'paragraph':
            p = paragraphs[int(index)]
            return {'target': target, 'style': _style(p, names),
                    'text': accepted_text(p), 'xml': _xml(p)}
        if kind == 'table':
            t = tables[int(index)]
            rows = t.findall(w('tr'))
            grid = [[accepted_text(tc)[:120] for tc in tr.findall(w('tc'))] for tr in rows[offset:offset + limit]]
            return {'target': target, 'rows': len(rows), 'cells': grid, 'next_offset': offset + limit if offset + limit < len(rows) else None,
                    'table_properties_xml': _xml(t.find(w('tblPr')))}
        if kind == 'section':
            s = body_sections(body)[int(index)]
            return {'target': target, **section_info(s), 'xml': _xml(s)}
        if kind == 'styles':
            path = related(parts, '/styles')
            styles = list(parse(parts[path]).iter(w('style'))) if path else []
            used = {v for v in root.xpath('//w:pStyle/@w:val|//w:rStyle/@w:val|//w:tblStyle/@w:val', namespaces={'w': W_NS})}
            rows = [{'id': s.get(w('styleId')), 'name': names.get(s.get(w('styleId'))), 'type': s.get(w('type')),
                     'based_on': s.find(w('basedOn')).get(w('val')) if s.find(w('basedOn')) is not None else None,
                     'in_use': s.get(w('styleId')) in used or None, 'default': s.get(w('default')) == '1' or None} for s in styles]
            return {'target': target, 'styles': [{k: v for k, v in r.items() if v is not None} for r in rows[offset:offset + limit]],
                    'total': len(rows), 'next_offset': offset + limit if offset + limit < len(rows) else None}
        if kind == 'part':
            data = parts[index].decode('utf8', 'replace')
            size = limit * 1000
            return {'target': target, 'xml': data[offset * 1000:offset * 1000 + size], 'total_chars': len(data),
                    'next_offset': offset + limit if (offset + limit) * 1000 < len(data) else None}
        raise ValueError('Unknown target; omit target to page the body, or use paragraph:N, table:N, section:N, styles or part:word/<name>.xml')
    items, pi, ti, si = [], 0, 0, 0
    sections = body_sections(body)
    for b, child in enumerate(body):
        row = None
        if child.tag == w('p'):
            text = accepted_text(child)
            row = {'p': pi, 'style': _style(child, names), 'text': text[:300] + ('…' if len(text) > 300 else '')}
            numpr = child.find(w('pPr') + '/' + w('numPr'))
            if numpr is not None:
                row['list'] = 'num %s level %s' % (_val(numpr, 'numId') or '?', _val(numpr, 'ilvl') or '0')
            revs = sum(revision_counts(child).values())
            if revs: row['revisions'] = revs
            if child.find(w('pPr') + '/' + w('sectPr')) is not None: row['ends_section'] = si; si += 1
            if not text.strip() and 'list' not in row and 'ends_section' not in row: row.pop('text')
            pi += 1
        elif child.tag == w('tbl'):
            rows = child.findall(w('tr'))
            row = {'t': ti, 'table': '%dx%d' % (len(rows), max((len(r.findall(w('tc'))) for r in rows), default=0)),
                   'style': names.get(_val(child.find(w('tblPr')), 'tblStyle'), _val(child.find(w('tblPr')), 'tblStyle')),
                   'text': ooxml.table_text(child)}
            ti += 1
        elif child.tag == w('sdt'):
            row = {'content_control': accepted_text(child)[:200]}
        if row is not None: items.append({'b': b, **{k: v for k, v in row.items() if v is not None}})
    page = items[offset:offset + limit]
    comments = related(parts, '/comments')
    notes = related(parts, '/footnotes')
    return {
        'counts': {'paragraphs': pi, 'tables': ti, 'sections': len(sections),
                   'comments': len(parse(parts[comments]).findall(w('comment'))) if comments else 0,
                   'footnotes': max(0, len(parse(parts[notes]).findall(w('footnote'))) - 2) if notes else 0},
        **({'revisions': counts, 'view': 'accepted (text includes insertions, omits deletions)'} if tracked else {}),
        'sections': [{'section': i, **section_info(s)} for i, s in enumerate(sections)],
        'items': page, 'next_offset': offset + limit if offset + limit < len(items) else None,
    }


def _val(parent, tag):
    el = None if parent is None else parent.find(w(tag))
    return None if el is None else el.get(w('val'))


def _style(p, names):
    ref = p.find(w('pPr') + '/' + w('pStyle'))
    return names.get(ref.get(w('val')), ref.get(w('val'))) if ref is not None else 'Normal'


def _xml(node, limit=8000):
    if node is None: return None
    text = XMLNS.sub('', etree.tostring(node, encoding='unicode'))
    return text if len(text) <= limit else text[:limit] + '… [truncated]'


def _brief(block):
    kind = block[0]
    if kind == 'p': return repr(block[3][:70] + ('…' if len(block[3]) > 70 else ''))
    if kind == 'tbl': return 'table %d rows (%s)' % (len(block[2]), block[3][:60])
    return kind


_SECT = re.compile(r'<w:sectPr[ >].*</w:sectPr>', re.S)


def _diff_blocks(before, after, limit=20):
    """Blocks align by kind and text; same-text pairs report formatting, similar pairs report text."""
    out = []
    similar = lambda a, b: a[0] == b[0] and difflib.SequenceMatcher(None, a[3], b[3], autojunk=False).ratio() > 0.5

    def changed(j, a, b):
        ends = _SECT.search(b[1]) is not None
        props = _SECT.sub('', a[1]) != _SECT.sub('', b[1])
        if ends != (_SECT.search(a[1]) is not None):
            out.append('block %d %s a section: %s' % (j, 'now ends' if ends else 'no longer ends', _brief(b)))
        if a[3] == b[3]:
            if props or a[2] != b[2]: out.append('block %d formatting/style changed: %s' % (j, _brief(b)))
        else: out.append('block %d text%s changed: %s -> %s' % (j, ' and formatting' if props else '', _brief(a), _brief(b)))
    matcher = difflib.SequenceMatcher(None, [(b[0], b[3]) for b in before], [(b[0], b[3]) for b in after], autojunk=False)
    for op, i1, i2, j1, j2 in matcher.get_opcodes():
        if op == 'equal':
            for k in range(i2 - i1): changed(j1 + k, before[i1 + k], after[j1 + k])
            continue
        i, j = i1, j1
        while i < i2 or j < j2:
            if i < i2 and j < j2 and similar(before[i], after[j]): changed(j, before[i], after[j]); i += 1; j += 1
            elif j < j2 and (i >= i2 or any(similar(before[i], after[k]) for k in range(j + 1, j2))):
                out.append('inserted block %d: %s' % (j, _brief(after[j]))); j += 1
            else: out.append('removed block (was %d): %s' % (i, _brief(before[i]))); i += 1
    return out[:limit], len(out)


def _semantic(parts, name):
    """Canonical content of a part, ignoring how the package is re-serialized (content types, empty
    relationship files) and Review mode's own trackRevisions setting."""
    data = parts.get(name)
    if name == '[Content_Types].xml': return None
    if name.endswith('.rels'):
        return frozenset(tuple(r.get(k) for k in ('Id', 'Type', 'Target', 'TargetMode')) for r in parse(data)) if data else frozenset()
    if data is None or not name.endswith('.xml'): return data
    root = parse(data)
    for node in list(root.iter(w('trackRevisions'))): node.getparent().remove(node)
    return canon(root)


def _by_id(parts, suffix, tag, key):
    path = related(parts, suffix)
    return {} if not path else {n.get(w(key)): canon(n) for n in parse(parts[path]).iter(w(tag))}


def verify(source, candidate, mode):
    before, after = ooxml.read_package(source), ooxml.read_package(candidate)
    import docx
    reopened = docx.Document(candidate)
    _ = [p.text for p in reopened.paragraphs], [t.rows for t in reopened.tables], [s.page_width for s in reopened.sections]
    main_b, main_a = main_part(before), main_part(after)
    if main_b != main_a: raise ValueError('The main document part moved')
    changed_parts = sorted(n for n in set(before) | set(after) if _semantic(before, n) != _semantic(after, n))
    if not changed_parts:
        raise ValueError('The program changed nothing. If the document already has what was asked, no new preview is needed. '
                         'Otherwise: a program runs top to bottom as a function body (a def it declares is never called), '
                         'and find() returns [] when the text is absent.')
    story_b = stories(before)
    problems = []
    # Accepted-view summary of what changed.
    body_b = story_root(resolve(parse(before[main_b]), True))
    body_a = story_root(resolve(parse(after[main_a]), True))
    changes, change_count = _diff_blocks(blocks(body_b), blocks(body_a))
    for name in sorted(set(story_b) - {main_b}):
        if name in after and blocks(story_root(resolve(parse(before[name]), True))) != blocks(story_root(resolve(parse(after[name]), True))):
            changes.append('%s changed' % name); change_count += 1
    added_stories = sorted(set(stories(after)) - set(story_b))
    for name in added_stories: changes.append('%s added' % name); change_count += 1
    sections_b = [section_info(s) for s in body_sections(parse(before[main_b]).find(w('body')))]
    sections_a = [section_info(s) for s in body_sections(body_a)]
    styles_b, styles_a = _by_id(before, '/styles', 'style', 'styleId'), _by_id(after, '/styles', 'style', 'styleId')
    doc_defaults = lambda parts: canon(parse(parts[related(parts, '/styles')]).find(w('docDefaults'))) if related(parts, '/styles') else ''
    style_changes = {k: v for k, v in {
        'added': sorted(set(styles_a) - set(styles_b)) or None,
        'changed': sorted(k for k in styles_b if k in styles_a and styles_a[k] != styles_b[k]) or None,
        'removed': sorted(set(styles_b) - set(styles_a)) or None,
        'defaults_changed': doc_defaults(before) != doc_defaults(after) or None}.items() if v}
    nums_b, nums_a = _by_id(before, '/numbering', 'num', 'numId'), _by_id(after, '/numbering', 'num', 'numId')
    abstract_b, abstract_a = _by_id(before, '/numbering', 'abstractNum', 'abstractNumId'), _by_id(after, '/numbering', 'abstractNum', 'abstractNumId')
    comments_b, comments_a = _by_id(before, '/comments', 'comment', 'id'), _by_id(after, '/comments', 'comment', 'id')
    counts_b = revision_counts(parse(before[main_b])); counts_a = revision_counts(parse(after[main_a]))
    for name in story_b:
        if name != main_b and name in after:
            for k, v in revision_counts(parse(after[name])).items(): counts_a[k] = counts_a.get(k, 0) + v
            if name in before:
                for k, v in revision_counts(parse(before[name])).items(): counts_b[k] = counts_b.get(k, 0) + v
    if mode == 'tracked':
        if any(k not in comments_a or comments_a[k] != v for k, v in comments_b.items()):
            problems.append('existing comments changed or were removed')
        # Rejecting every revision must give back the source (with its own revisions rejected too).
        for name, kind in story_b.items():
            if name not in after: problems.append('%s was removed' % name); continue
            rb, ra = resolve(parse(before[name]), False), resolve(parse(after[name]), False)
            if kind in ('footnotes', 'endnotes'):
                nb, na = ooxml.note_blocks(rb), ooxml.note_blocks(ra)
                if any(na.get(k) != v for k, v in nb.items()): problems.append('existing %s changed without revisions' % kind)
                continue
            xb, xa = blocks(story_root(rb)), blocks(story_root(ra))
            if xb != xa:
                sm = difflib.SequenceMatcher(None, [(b[0], b[1], b[2]) for b in xb], [(b[0], b[1], b[2]) for b in xa], autojunk=False)
                op = next(o for o in sm.get_opcodes() if o[0] != 'equal')
                where = ('%s block %d: source %s, after rejecting revisions %s' %
                         (name, op[1], _brief(xb[op[1]]) if op[1] < len(xb) else 'end',
                          _brief(xa[op[3]]) if op[3] < len(xa) else 'nothing'))
                problems.append('edit is not a tracked revision (%s)' % where)
        edited = {k: v for k, v in style_changes.items() if k != 'added'}   # new styles leave the source intact
        if edited:
            problems.append('style definitions changed (%s); Word cannot track style edits, so use direct formatting or ask for direct editing'
                            % '; '.join('%s: %s' % (k, ', '.join(v) if isinstance(v, list) else 'document defaults') for k, v in edited.items()))
        if any(abstract_a.get(k) != v for k, v in abstract_b.items()) or any(nums_a.get(k) != v for k, v in nums_b.items()):
            problems.append('existing list definitions changed; create a new list instead')
    if problems:
        raise ValueError('Review mode verification failed: ' + '; '.join(problems[:4]))
    new_revisions = {k: v - counts_b.get(k, 0) for k, v in counts_a.items() if v > counts_b.get(k, 0)}
    if not change_count:  # page setup, style, comment or revision edits leave the accepted text as it was
        changes, change_count = ['%s changed' % name for name in changed_parts], len(changed_parts)
    return {k: v for k, v in {
        'reopened': True, 'python_docx_reopened': True,
        'mode': mode + '-candidate', 'review_verified': mode == 'tracked',
        'changes': changes, 'change_count': change_count, 'changes_truncated': change_count > len(changes) or None,
        'sections': sections_a if sections_a != sections_b else None,
        'styles': style_changes or None,
        'new_lists': len(nums_a) - len(nums_b) or None,
        'comments': len(comments_a) if comments_a != comments_b else None,
        'revision_count': sum(counts_a.values()), 'new_revision_count': sum(new_revisions.values()),
        'new_revisions': new_revisions or None,
        'changed_parts': changed_parts,
        'candidate_sha256': sha256(open(candidate, 'rb').read()).hexdigest(),
    }.items() if v is not None}


def main():
    command = sys.argv[1]
    try:
        if command == 'inspect':
            request = json.loads(sys.stdin.readline(1 << 20) or '{}')
            out = inspect(sys.argv[2], request)
        elif command == 'verify':
            out = verify(sys.argv[2], sys.argv[3], sys.argv[4])
        else: raise ValueError('unknown command')
        print(json.dumps({'ok': True, **out}, ensure_ascii=False))
    except Exception as error:  # noqa: BLE001
        print(json.dumps({'ok': False, 'error': str(error)[:1500]}, ensure_ascii=False))


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    main()
