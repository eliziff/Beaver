"""Build the harder "pointed at a file system" variant of a record.

Usage: python haystack.py <record_id> [--seed N] [--other 6] [--corpus 4] [--synthetic specs.json]

Output: %LOCALAPPDATA%/OpenLegalData/benchmarks/court-record-exhibits/haystack/<record_id>/
  files/<name>           every document, with names that give nothing away
  manifest.json          role and provenance of every file (never shown to a model)

Roles:
  exhibit        a stripped exhibit of this affidavit (label and gold as in gold.json)
  same_matter    a real document from the same proceeding that is not an exhibit
                 (from records/<id>/matter_docs.json); chronology rows citing it are neutral
  other_matter   a stripped exhibit of a different benchmark record
  corpus         an unrelated real legal PDF from experiments/legal_pdf_corpus
  synthetic      a generated DMS document: same people, companies and period, but
                 business irrelevant to the affidavit's matters (spec in --synthetic)

Synthetic specs are JSON: {"generator": "...", "created": "...", "documents": [
  {"kind": "email"|"letter"|"memo"|"invoice"|"minutes"|"text", "format": "pdf"|"eml"|"docx",
   "date": "YYYY-MM-DD", "time": "HH:MM", "from": "...", "to": ["..."], "cc": ["..."],
   "subject": "...", "body": "...", "letterhead": "...", "items": [[desc, amount], ...],
   "why_irrelevant": "..."}]}
"""
import email.message, email.utils, hashlib, json, os, random, re, shutil, sys, datetime
import fitz

BENCH = os.path.join(os.environ["LOCALAPPDATA"], "OpenLegalData", "benchmarks", "court-record-exhibits")
CORPUS = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), "experiments", "legal_pdf_corpus", "pdfs", "ca", "digitalborn")


def text_pdf(path, blocks):
    """Render (text, size, bold) blocks onto letter-size pages like a printed document."""
    doc = fitz.open()
    page, y = None, 72
    for text, size, bold in blocks:
        font = "hebo" if bold else "helv"
        for para in text.split("\n"):
            lines = wrap(para, size)
            for line in lines or [""]:
                if page is None or y > 740:
                    page, y = doc.new_page(width=612, height=792), 72
                page.insert_text((72, y), line, fontsize=size, fontname=font)
                y += size * 1.45
        y += size * 0.6
    doc.set_metadata({})
    doc.save(path, garbage=4, deflate=True)


def wrap(para, size, width=468):
    words, lines, cur = para.split(), [], ""
    for w in words:
        trial = (cur + " " + w).strip()
        if fitz.get_text_length(trial, fontsize=size) > width and cur:
            lines.append(cur)
            cur = w
        else:
            cur = trial
    if cur:
        lines.append(cur)
    return lines


def when(d):
    dt = datetime.datetime.strptime(d["date"] + " " + d.get("time", "10:15"), "%Y-%m-%d %H:%M")
    return dt


def render(d, path):
    kind, fmt = d["kind"], d.get("format", "pdf")
    dt = when(d)
    if kind == "email" and fmt == "eml":
        m = email.message.EmailMessage()
        m["From"], m["To"], m["Subject"] = d["from"], ", ".join(d["to"]), d["subject"]
        if d.get("cc"):
            m["Cc"] = ", ".join(d["cc"])
        m["Date"] = email.utils.format_datetime(dt.replace(tzinfo=datetime.timezone(datetime.timedelta(hours=-5))))
        m["Message-ID"] = email.utils.make_msgid(domain=d["from"].split("@")[-1].strip(">") if "@" in d["from"] else "mail.local")
        m.set_content(d["body"])
        open(path, "wb").write(bytes(m))
        return
    if fmt == "docx":
        import docx
        doc = docx.Document()
        if d.get("letterhead"):
            doc.add_paragraph(d["letterhead"]).runs[0].bold = True
        doc.add_paragraph(dt.strftime("%B %d, %Y").replace(" 0", " "))
        if d.get("to"):
            doc.add_paragraph("\n".join(d["to"]))
        if d.get("subject"):
            doc.add_paragraph(("Re: " if kind == "letter" else "Subject: ") + d["subject"]).runs[0].bold = True
        for para in d["body"].split("\n\n"):
            doc.add_paragraph(para)
        doc.core_properties.author = ""
        doc.save(path)
        return
    blocks = []
    if kind == "email":
        blocks = [(d["subject"], 13, True),
                  (f"From: {d['from']}\nSent: {dt.strftime('%A, %B %d, %Y %I:%M %p').replace(' 0', ' ')}\nTo: {'; '.join(d['to'])}" + (f"\nCc: {'; '.join(d['cc'])}" if d.get("cc") else "") + f"\nSubject: {d['subject']}", 10, False),
                  (d["body"], 10.5, False)]
    elif kind in ("letter", "memo"):
        if d.get("letterhead"):
            blocks.append((d["letterhead"], 12, True))
        if kind == "memo":
            blocks.append(("MEMORANDUM", 14, True))
            blocks.append((f"TO: {'; '.join(d.get('to', []))}\nFROM: {d.get('from', '')}\nDATE: {dt.strftime('%B %d, %Y').replace(' 0', ' ')}\nRE: {d.get('subject', '')}", 10.5, False))
        else:
            blocks.append((dt.strftime("%B %d, %Y").replace(" 0", " "), 10.5, False))
            blocks.append(("\n".join(d.get("to", [])), 10.5, False))
            if d.get("subject"):
                blocks.append(("Re: " + d["subject"], 10.5, True))
        blocks.append((d["body"], 10.5, False))
    elif kind == "invoice":
        blocks = [(d.get("letterhead", d.get("from", "")), 13, True), ("INVOICE", 16, True),
                  (f"Invoice date: {dt.strftime('%B %d, %Y').replace(' 0', ' ')}\nBill to: {'; '.join(d.get('to', []))}\nRe: {d.get('subject', '')}", 10.5, False),
                  ("\n".join(f"{desc}    ${amt:,.2f}" for desc, amt in d.get("items", [])) + (f"\n\nTOTAL    ${sum(a for _, a in d.get('items', [])):,.2f}" if d.get("items") else ""), 10.5, False),
                  (d.get("body", ""), 10, False)]
    elif kind == "minutes":
        blocks = [(d.get("letterhead", ""), 12, True), (d.get("subject", "Minutes of Meeting"), 14, True),
                  (f"Date: {dt.strftime('%B %d, %Y').replace(' 0', ' ')}\nPresent: {', '.join(d.get('to', []))}", 10.5, False), (d["body"], 10.5, False)]
    else:  # text messages, printed
        blocks = [(f"Messages with {', '.join(d.get('to', []))}", 13, True), (dt.strftime("%Y-%m-%d"), 10, False), (d["body"], 11, False)]
    text_pdf(path, [b for b in blocks if b[0]])


def main():
    args = sys.argv[1:]
    rid = args[0]
    opt = lambda n, v: args[args.index(n) + 1] if n in args else v
    rng = random.Random(int(opt("--seed", 7)) * 1_000_003 + int(hashlib.sha256(rid.encode()).hexdigest()[:8], 16))
    rec = os.path.join(BENCH, "records", rid)
    out = os.path.join(BENCH, "haystack", rid)
    if os.path.exists(out):
        shutil.rmtree(out)
    os.makedirs(os.path.join(out, "files"))
    gold = json.load(open(os.path.join(rec, "gold.json"), encoding="utf-8"))
    entries = []
    names = set()

    def place(src, ext, meta):
        while True:
            name = "".join(rng.choice("abcdefghjkmnpqrstuvwxyz23456789") for _ in range(8)) + ext
            if name not in names:
                break
        names.add(name)
        dst = os.path.join(out, "files", name)
        if callable(src):
            src(dst)
        else:
            shutil.copyfile(src, dst)
        entries.append({"file": name, **meta})

    for e in gold["exhibits"]:
        place(os.path.join(rec, "files", e["file"]), ".pdf", {"role": "exhibit", "label": e["label"], "record_file": e["file"]})
    md = os.path.join(rec, "matter_docs.json")
    for m in json.load(open(md, encoding="utf-8")) if os.path.exists(md) else []:
        src = os.path.join(BENCH, m["file"]) if m["file"].replace("\\", "/").startswith("raw/") else os.path.join(BENCH, "raw", "matter", rid, m["file"])
        if os.path.exists(src):
            place(src, os.path.splitext(src)[1], {"role": "same_matter", "provenance": m})
    others = [d for d in sorted(os.listdir(os.path.join(BENCH, "records"))) if d != rid and os.path.exists(os.path.join(BENCH, "records", d, "gold.json"))]
    small = lambda path, mb: os.path.getsize(path) < mb * 1_000_000  # keep haystacks light on disk
    pool = [(d, e) for d in others for e in json.load(open(os.path.join(BENCH, "records", d, "gold.json"), encoding="utf-8"))["exhibits"]
            if small(os.path.join(BENCH, "records", d, "files", e["file"]), 10)]
    for d, e in rng.sample(pool, min(int(opt("--other", 6)), len(pool))):
        place(os.path.join(BENCH, "records", d, "files", e["file"]), ".pdf", {"role": "other_matter", "provenance": {"record": d, "label": e["label"]}})
    corpus = sorted(os.path.join(dp, f) for dp, _, fs in os.walk(CORPUS) for f in fs if f.endswith(".pdf") and small(os.path.join(dp, f), 5))
    for src in rng.sample(corpus, min(int(opt("--corpus", 4)), len(corpus))):
        place(src, ".pdf", {"role": "corpus", "provenance": {"path": os.path.relpath(src, CORPUS)}})
    spec_path = opt("--synthetic", None)
    if spec_path:
        spec = json.load(open(spec_path, encoding="utf-8"))
        for i, d in enumerate(spec["documents"]):
            ext = {"eml": ".eml", "docx": ".docx"}.get(d.get("format", "pdf"), ".pdf")
            place(lambda dst, d=d: render(d, dst), ext, {"role": "synthetic", "provenance": {
                "generator": spec["generator"], "created": spec["created"], "spec": os.path.basename(spec_path), "index": i,
                "kind": d["kind"], "date": d["date"], "why_irrelevant": d.get("why_irrelevant", "")}})
    rng.shuffle(entries)
    json.dump({"record_id": rid, "files": entries}, open(os.path.join(out, "manifest.json"), "w", encoding="utf-8"), indent=1)
    roles = {}
    for e in entries:
        roles[e["role"]] = roles.get(e["role"], 0) + 1
    print(rid, len(entries), "files", roles, "->", out)


if __name__ == "__main__":
    main()
