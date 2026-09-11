from __future__ import annotations

import argparse
import hashlib
import json
import tempfile
import threading
import time
import traceback
from datetime import datetime, timezone
from functools import partial
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlsplit

import fitz
from docx import Document
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.common.action_chains import ActionChains
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.support.ui import Select, WebDriverWait


ROOT = Path(__file__).resolve().parents[1]
CHROME = next((path for path in (
    Path(r"C:\Program Files\Google\Chrome\Application\chrome.exe"),
    Path(r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe"),
) if path.is_file()), None)
CHROMEDRIVER = next(iter(sorted(Path.home().parent.glob(
    r"*/.cache/selenium/chromedriver/win64/*/chromedriver.exe"
), reverse=True)), None)
CHOICES = {
    "ab-kb-affidavit-exhibits": "Alberta",
    "fc-motion-record-moving": "Federal courts",
    "fca-motion-record-moving": "Federal courts",
    "ab-ca-appeal-record": "Alberta",
}
RESPONSIVE_SCROLL_RESETS: list[dict[str, object]] = []


class QuietStaticHandler(SimpleHTTPRequestHandler):
    def log_message(self, _format: str, *_args: object) -> None:
        pass

    def do_GET(self) -> None:
        if self.path == "/favicon.ico":
            self.send_response(204)
            self.end_headers()
            return
        super().do_GET()


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def make_pdf(path: Path, lines: list[str]) -> None:
    with fitz.open() as document:
        page = document.new_page(width=612, height=792)
        remaining = page.insert_textbox(
            fitz.Rect(72, 64, 540, 735), "\n".join(lines),
            fontsize=11, fontname="Times-Roman", lineheight=1.25,
        )
        assert remaining >= 0, f"Fixture text overflowed {path.name}."
        document.set_metadata({"title": lines[0], "author": "Beaver live proof"})
        document.save(path)


def make_long_pdf(path: Path, pages: int = 80) -> None:
    with fitz.open() as document:
        for number in range(1, pages + 1):
            page = document.new_page(width=612, height=792)
            page.insert_text((72, 72), f"LARGE FILE BUILD-RACE PROOF — PAGE {number}",
                             fontsize=11, fontname="Times-Roman")
        document.save(path)


def make_transcript_pdf(path: Path) -> None:
    with fitz.open() as document:
        for text in ("TRANSCRIPT OF PROCEEDINGS", "TABLE OF CONTENTS", "ORAL PROCEEDINGS TRANSCRIPT"):
            page = document.new_page(width=612, height=792)
            page.insert_text((72, 72), text, fontsize=11, fontname="Times-Roman")
        document.set_page_labels([
            {"startpage": 0, "prefix": "", "style": "", "firstpagenum": 1},
            {"startpage": 1, "prefix": "", "style": "r", "firstpagenum": 1},
            {"startpage": 2, "prefix": "", "style": "D", "firstpagenum": 1},
        ])
        document.save(path)


def make_fixtures(directory: Path) -> dict[str, Path]:
    directory.mkdir()
    fixtures = {
        "affidavit": directory / "Affidavit-of-Alexandra-Smith.pdf",
        "exhibit_b": directory / "Project-ledger.pdf",
        "ambiguous": directory / "Service-agreement.pdf",
        "notice": directory / "Notice-of-motion.pdf",
        "representations": directory / "Written-representations.pdf",
        "oral_request": directory / "Request-for-oral-hearing.docx",
        "large": directory / "Large-file-for-build-race.pdf",
        "appeal_pleading": directory / "Statement-of-Claim.pdf",
        "appeal_reasons": directory / "Reasons-for-Judgment.pdf",
        "appeal_order": directory / "Formal-Order.pdf",
        "appeal_notice": directory / "Notice-of-Appeal.pdf",
        "appeal_transcript_prior": directory / "Prior-Transcript.pdf",
        "appeal_transcript": directory / "EVK26DOEJ.pdf",
    }
    make_pdf(fixtures["affidavit"], [
        "COURT OF KING'S BENCH OF ALBERTA",
        "COURT FILE NUMBER: 2401-12345",
        "JUDICIAL CENTRE: CALGARY",
        "BETWEEN:",
        "NORTH PRAIRIE LTD.",
        "PLAINTIFF",
        "- and -",
        "RIVERSTONE INC.",
        "DEFENDANT",
        "2nd Affidavit of Alexandra Smith",
        "I, Alexandra Smith, of Calgary, Alberta, SWEAR AND SAY THAT:",
        "1. I am a director of the Plaintiff and have personal knowledge of these facts.",
        "2. The signed service agreement is attached and marked as Exhibit A.",
        "3. The project ledger is attached and marked as Exhibit B.",
        "SWORN BEFORE ME AT THE CITY OF CALGARY",
        "IN THE PROVINCE OF ALBERTA",
        "THIS 29TH DAY OF AUGUST, 2026",
        "A Commissioner for Oaths in and for Alberta",
    ])
    make_pdf(fixtures["exhibit_b"], [
        'This is Exhibit "B" referred to in the Affidavit.',
        "PROJECT LEDGER",
        "August 2026",
        "The ledger records the project invoices relied on by the deponent.",
    ])
    make_pdf(fixtures["ambiguous"], [
        "SERVICE AGREEMENT",
        "Executed January 15, 2026",
        "North Prairie Ltd. and Riverstone Inc. agree to the following services.",
    ])
    make_pdf(fixtures["notice"], [
        "FEDERAL COURT",
        "COURT FILE NUMBER: T-982-19",
        "BETWEEN:",
        "NORTH PRAIRIE LTD.",
        "APPLICANT",
        "- and -",
        "RIVERSTONE INC.",
        "RESPONDENT",
        "NOTICE OF MOTION",
        "The Applicant moves for procedural directions.",
    ])
    make_pdf(fixtures["representations"], [
        "WRITTEN REPRESENTATIONS OF THE APPLICANT",
        "The requested directions are just and proportionate.",
    ])
    word = Document()
    word.add_heading("REQUEST FOR ORAL HEARING", level=1)
    word.add_paragraph("The moving party requests an oral hearing under Rule 369.2(2).")
    word.save(fixtures["oral_request"])
    make_long_pdf(fixtures["large"])
    for key, heading in (
        ("appeal_pleading", "STATEMENT OF CLAIM"),
        ("appeal_reasons", "REASONS FOR JUDGMENT"),
        ("appeal_order", "FORMAL ORDER"),
        ("appeal_notice", "NOTICE OF APPEAL"),
    ):
        metadata = (["Trial Court File Number: 2301-00456",
            "Decision Maker Appealed From: The Honourable Justice A. Ng",
            "Decision Date: March 4, 2026", "Decision Filing Date: March 5, 2026"]
            if key == "appeal_reasons" else [])
        make_pdf(fixtures[key], [heading, "Court File Number 2403-00789",
                                 "North Prairie Ltd. v. Riverstone Inc.", *metadata])
    make_transcript_pdf(fixtures["appeal_transcript_prior"])
    make_transcript_pdf(fixtures["appeal_transcript"])
    return fixtures


def start_browser(profile: Path, headed: bool) -> webdriver.Chrome:
    if not CHROME:
        raise RuntimeError("Google Chrome is required for this live test.")
    options = webdriver.ChromeOptions()
    options.binary_location = str(CHROME)
    if not headed:
        options.add_argument("--headless=new")
    for argument in ("--disable-gpu", "--no-sandbox", "--disable-dev-shm-usage",
                     "--disable-crash-reporter", "--no-first-run"):
        options.add_argument(argument)
    options.add_argument(f"--user-data-dir={profile}")
    options.set_capability("goog:loggingPrefs", {"browser": "ALL"})
    options.add_experimental_option("prefs", {
        "download.prompt_for_download": False,
        "plugins.always_open_pdf_externally": True,
    })
    driver = webdriver.Chrome(
        options=options,
        service=Service(str(CHROMEDRIVER)) if CHROMEDRIVER else None,
    )
    driver.execute_cdp_cmd("Network.enable", {})
    driver.execute_cdp_cmd("Network.setCacheDisabled", {"cacheDisabled": True})
    driver.execute_cdp_cmd("Page.addScriptToEvaluateOnNewDocument", {
        "source": """
Object.defineProperty(window,'showOpenFilePicker',{value:undefined});
const nativeFileClick=HTMLInputElement.prototype.click;
HTMLInputElement.prototype.click=function(){
  if(this.type==='file'&&!this.isConnected){
    this.dataset.chromedriverFilePicker=''; document.body.append(this); return;
  }
  return nativeFileClick.call(this);
};
""",
    })
    driver.set_window_size(1440, 1000)
    return driver


def navigation_timing(driver: webdriver.Chrome) -> dict[str, object]:
    return driver.execute_script("""
const n=performance.getEntriesByType('navigation')[0];
return n&&{type:n.type,duration:n.duration,responseStart:n.responseStart,
  domContentLoaded:n.domContentLoadedEventEnd,load:n.loadEventEnd,
  transferSize:n.transferSize,decodedBodySize:n.decodedBodySize};
""")


def set_viewport(driver: webdriver.Chrome, width: int = 1440, height: int = 900,
                 dpr: int = 1) -> None:
    driver.execute_cdp_cmd("Emulation.setPageScaleFactor", {"pageScaleFactor": 1})
    driver.execute_cdp_cmd("Emulation.setDeviceMetricsOverride", {
        "width": width, "height": height, "deviceScaleFactor": dpr, "mobile": False,
    })
    WebDriverWait(driver, 5).until(
        lambda item: item.execute_script("return [innerWidth,innerHeight,devicePixelRatio]")
        == [width, height, dpr]
    )


def restore_viewport(driver: webdriver.Chrome) -> dict[str, object]:
    set_viewport(driver)
    retained = driver.execute_script("""
const main=document.querySelector('main'),workspace=document.querySelector('.court-records-workspace'),
  r=workspace?.getBoundingClientRect();
return {url:location.href,main:main?.scrollTop||0,workspace:workspace?.scrollTop||0,
  workspaceRect:r&&{top:r.top,bottom:r.bottom,height:r.height}};
""")
    if retained["main"]:
        RESPONSIVE_SCROLL_RESETS.append(retained)
    driver.execute_script("""
for(const node of [document.scrollingElement,document.querySelector('main'),
  document.querySelector('.court-records-workspace')]) if(node) node.scrollTop=0;
""")
    return retained


def choose_new_profile(
    driver: webdriver.Chrome, profile_id: str
) -> dict[str, object]:
    wait = WebDriverWait(driver, 20)
    driver.find_element(By.XPATH, "//button[normalize-space()='New court record']").click()
    dialog = wait.until(lambda item: item.find_element(By.CSS_SELECTOR, "dialog[open]"))
    jurisdiction = CHOICES[profile_id]
    proof: dict[str, object] = {"jurisdiction": jurisdiction, "profile": profile_id}
    dialog.find_element(By.XPATH,
        f".//*[@aria-label='Jurisdiction']//button[normalize-space()='{jurisdiction}']").click()
    assert driver.find_elements(By.CSS_SELECTOR, "dialog[open]") == [dialog], \
        "Jurisdiction replaced the dialog"
    wait.until(lambda _item: dialog.find_element(
        By.CSS_SELECTOR, f"button[data-choice='{profile_id}']")).click()
    wait.until(lambda item: not item.find_elements(By.CSS_SELECTOR, "dialog[open]") and
        item.find_elements(By.CSS_SELECTOR, "button[aria-label='Back from court record']"))
    return proof


def reset_workspace(
    driver: webdriver.Chrome, url: str, profile_id: str, chooser_output: Path | None = None
) -> tuple[float, dict[str, object], dict[str, object]]:
    started = time.perf_counter()
    parsed = urlsplit(url)
    driver.get("about:blank")
    driver.execute_cdp_cmd("Storage.clearDataForOrigin", {
        "origin": f"{parsed.scheme}://{parsed.netloc}",
        "storageTypes": "indexeddb,local_storage",
    })
    driver.get(url)
    wait = WebDriverWait(driver, 20)
    workspace = wait.until(lambda item: item.find_element(
        By.CSS_SELECTOR, "#court-record-workspace:not([inert])"
    ))
    timing = navigation_timing(driver)
    assert not workspace.find_elements(By.CSS_SELECTOR, "[data-court-record-chooser]")
    assistant = driver.find_elements(
        By.CSS_SELECTOR, ".court-records-workspace button[aria-label='Assistant']"
    )
    if parsed.path.endswith("court-records.html"):
        assert not assistant, "The standalone builder exposed Beaver's assistant."
    restore_viewport(driver)
    if chooser_output:
        assert driver.save_screenshot(str(chooser_output / "blank-start.png"))
    chooser_proof = choose_new_profile(driver, profile_id)
    return round((time.perf_counter() - started) * 1000, 1), timing, chooser_proof


def field(driver: webdriver.Chrome, field_id: str):
    return driver.find_element(By.ID, f"cover-{field_id}")


def select_filing_party(driver: webdriver.Chrome, name: str) -> None:
    group = driver.find_elements(By.XPATH,
        "//fieldset[.//legend[contains(normalize-space(.),'Filing parties')]]")
    if not group:
        return  # A single eligible party is selected automatically.
    label = next(item for item in group[0].find_elements(By.TAG_NAME, "label")
                 if item.text.casefold().startswith(name.casefold()))
    control = label.find_element(By.CSS_SELECTOR, "input[type='checkbox']")
    if not control.is_selected():
        control.click()


def enter(element, value: str) -> None:
    element.send_keys(Keys.CONTROL, "a")
    element.send_keys(value)


def reveal(driver: webdriver.Chrome, element) -> None:
    driver.execute_script(
        "for(let n=arguments[0];n;n=n.parentElement) if(n.tagName==='DETAILS') n.open=true;"
        "arguments[0].scrollIntoView({block:'center'});", element,
    )


def toggle_disclosure(driver: webdriver.Chrome, details) -> None:
    original = driver.execute_script("return arguments[0].open", details)
    summary = details.find_element(By.TAG_NAME, "summary")
    summary.click()
    WebDriverWait(driver, 5).until(
        lambda _item: driver.execute_script("return arguments[0].open", details) != original
    )
    summary.click()
    WebDriverWait(driver, 5).until(
        lambda _item: driver.execute_script("return arguments[0].open", details) == original
    )


def upload(driver: webdriver.Chrome, kind_id: str, path: Path) -> None:
    selector = f"[data-kind-id='{kind_id}']"
    before = len(driver.find_elements(By.CSS_SELECTOR, f"{selector} [data-entry-id]"))
    control = driver.find_element(By.ID, f"court-record-{kind_id}-file")
    dynamic = control.tag_name != "input"
    if dynamic:
        control.click()
        picker = WebDriverWait(driver, 5).until(lambda item: item.find_element(
            By.CSS_SELECTOR, "input[data-chromedriver-file-picker]"
        ))
    else:
        picker = control
    picker.send_keys(str(path.resolve()))

    def prepared(item: webdriver.Chrome) -> bool:
        return bool(item.execute_script("""
const rows=[...document.querySelectorAll(arguments[0])];
return rows.length>=arguments[2]&&rows.some(row=>[...row.querySelectorAll('span[title]')]
  .some(node=>node.title===arguments[1]))&&rows.every(row=>row.getAttribute('aria-busy')!=='true');
""", f"{selector} [data-entry-id]", path.name, max(1, before)))

    WebDriverWait(driver, 30).until(prepared)
    if dynamic:
        driver.execute_script("arguments[0].remove()", picker)


def article(driver: webdriver.Chrome, filename: str):
    return driver.find_element(
        By.XPATH, f"//span[@title={json.dumps(filename)}]/ancestor::article[1]"
    )


def exhibit_slot(driver: webdriver.Chrome, label: str):
    return driver.find_element(By.CSS_SELECTOR, f"[aria-label='Exhibit {label} slot']")


def unassigned_pool(driver: webdriver.Chrome):
    return driver.find_element(
        By.XPATH, "//h4[starts-with(normalize-space(),'Files')]/ancestor::div[1]"
    )


def contains_file(container, filename: str) -> bool:
    return bool(container.find_elements(By.CSS_SELECTOR, f"span[title='{filename}']"))


def drag(driver: webdriver.Chrome, source, target) -> dict[str, object]:
    reveal(driver, source)
    driver.execute_script("""
window.__courtRecordDrag=[];
for(const type of ['dragstart','dragenter','dragover','drop','dragend']) document.addEventListener(type,event=>
  window.__courtRecordDrag.push({type,target:event.target?.tagName||'',
    entry:event.target?.closest?.('[data-entry-id]')?.dataset.entryId||'',
    slot:event.target?.closest?.('[aria-label$=" slot"]')?.getAttribute('aria-label')||'',
    value:event.dataTransfer?.getData('text/x-court-record-entry')||'',trusted:event.isTrusted}),
    {capture:true,once:true});
""")
    handles = source.find_elements(By.CSS_SELECTOR, ".lucide-grip-vertical")
    handle = handles[0] if handles else source
    delta = driver.execute_script("""
const a=arguments[0].getBoundingClientRect(),b=arguments[1].getBoundingClientRect();
return {x:b.left+b.width/2-a.left-a.width/2,y:b.top+b.height/2-a.top-a.height/2};
""", handle, target)
    ActionChains(driver, duration=700).move_to_element(handle).click_and_hold() \
        .pause(.2).move_by_offset(8, 0).move_by_offset(delta["x"] - 8, delta["y"]) \
        .pause(.3).release().perform()
    pointer_events = driver.execute_script("return window.__courtRecordDrag")
    if any(event["type"] == "drop" for event in pointer_events):
        return {"route": "webdriver-pointer", "events": pointer_events}
    data = {"items": [{"mimeType": "text/x-court-record-entry",
                       "data": source.get_attribute("data-entry-id")}],
            "dragOperationsMask": 16}
    point = driver.execute_script("""
const r=arguments[0].getBoundingClientRect();return {x:r.left+r.width/2,y:r.top+r.height/2};
""", target)
    for event_type in ("dragEnter", "dragOver", "drop"):
        driver.execute_cdp_cmd("Input.dispatchDragEvent", {
            "type": event_type, "x": point["x"], "y": point["y"], "data": data,
        })
    return {"route": "cdp-browser-drag", "webdriver_events": pointer_events,
            "events": driver.execute_script("return window.__courtRecordDrag")}


def keyboard_contract(driver: webdriver.Chrome) -> dict[str, object]:
    driver.execute_script(
        "document.activeElement?.blur();document.body.tabIndex=-1;document.body.focus();"
    )
    driver.find_element(By.TAG_NAME, "body").send_keys(Keys.TAB)
    focus = driver.execute_script("""
const node=document.activeElement, rect=node.getBoundingClientRect(), style=getComputedStyle(node);
return {tag:node.tagName,text:node.textContent.trim(),width:rect.width,height:rect.height,
  outline:style.outlineWidth,shadow:style.boxShadow};
""")
    assert focus["tag"] == "A" and focus["text"] in {
        "Skip to builder", "Skip to content",
    }, focus
    assert focus["width"] >= 24 and focus["height"] >= 24, focus
    assert focus["outline"] != "0px" or focus["shadow"] != "none", focus
    driver.switch_to.active_element.send_keys(Keys.ENTER)
    if focus["text"] == "Skip to builder":
        WebDriverWait(driver, 5).until(
            lambda item: item.switch_to.active_element.get_attribute("id") == "court-record-workspace"
        )
    else:
        WebDriverWait(driver, 5).until(
            lambda item: urlsplit(item.current_url).fragment == "main-content"
        )
    driver.execute_script("document.body.removeAttribute('tabindex')")
    return focus


def viewport_contract(
    driver: webdriver.Chrome, output: Path, slug: str, focus_selector: str
) -> dict[str, object]:
    results: dict[str, object] = {}
    script = """
const visible=node=>{const s=getComputedStyle(node),r=node.getBoundingClientRect();
  return s.display!=='none'&&s.visibility!=='hidden'&&r.width>1&&r.height>1&&!node.classList.contains('sr-only')};
const name=node=>node.getAttribute('aria-label')||[...(node.labels||[])].map(x=>x.textContent).join(' ').trim()||node.textContent.trim()||node.title;
const controls=[...document.querySelectorAll("button,summary,select,input:not([type=file]),textarea,label:has(input[type=file]),a[href]")].filter(visible);
const targetRect=node=>['checkbox','radio'].includes(node.type)&&node.closest('label')
  ? node.closest('label').getBoundingClientRect():node.getBoundingClientRect();
const ids=[...document.querySelectorAll('[id]')].map(node=>node.id);
const save=document.querySelector('[data-save-filing-details]'), saveRect=save?.getBoundingClientRect(),
  setupRect=save?.closest('section')?.getBoundingClientRect();
return {width:innerWidth,overflow:document.documentElement.scrollWidth-document.documentElement.clientWidth,
  dpr:devicePixelRatio,header:(()=>{const r=document.querySelector('.court-records-workspace>header')?.getBoundingClientRect();return r&&{left:r.left,right:r.right,height:r.height}})(),
  duplicateIds:[...new Set(ids.filter((id,index)=>ids.indexOf(id)!==index))],
  unnamed:controls.filter(node=>!name(node)).map(node=>node.outerHTML.slice(0,120)),
  smallTargets:controls.filter(node=>{const r=targetRect(node);return r.width<24||r.height<24})
    .map(node=>{const r=targetRect(node);return `${node.tagName}:${name(node)}:${Math.round(r.width)}x${Math.round(r.height)}`}),
  smallFormText:innerWidth>640?[]:[...document.querySelectorAll('input:not([type=file]):not([type=checkbox]):not([type=radio]),select,textarea')].filter(visible)
    .filter(node=>parseFloat(getComputedStyle(node).fontSize)<16).map(node=>`${node.tagName}:${getComputedStyle(node).fontSize}`),
  positiveTabindex:[...document.querySelectorAll('[tabindex]')].filter(node=>Number(node.tabIndex)>0).map(node=>node.outerHTML.slice(0,120)),
  crampedEntryFields:[...document.querySelectorAll('[data-entry-id] > div:first-of-type label input')]
    .filter(visible).filter(node=>node.getBoundingClientRect().width<120)
    .map(node=>`${name(node)}:${Math.round(node.getBoundingClientRect().width)}px`),
  filingDetails:saveRect&&setupRect?{left:saveRect.left,right:saveRect.right,
    containerLeft:setupRect.left,containerRight:setupRect.right,
    contained:saveRect.left>=setupRect.left-1&&saveRect.right<=setupRect.right+1}:null,
  wideNodes:[...document.querySelectorAll('body *')].filter(visible)
    .filter(node=>node.getBoundingClientRect().right>document.documentElement.clientWidth+1)
    .map(node=>{const r=node.getBoundingClientRect(),s=getComputedStyle(node);return `${node.tagName}.${node.className}:left=${Math.round(r.left)},width=${Math.round(r.width)},right=${Math.round(r.right)},min=${s.minWidth}`}).slice(0,12)};
"""
    for label, width, height, dpr in (("1440", 1440, 1000, 1),
                                      ("320", 320, 800, 1),
                                      ("200pct", 640, 500, 2)):
        set_viewport(driver, width, height, dpr)
        metrics = driver.execute_script(script)
        driver.execute_script("scrollTo(0,0)")
        driver.save_screenshot(str(output / f"{slug}-{label}-top.png"))
        reveal(driver, driver.find_element(By.CSS_SELECTOR, focus_selector))
        metrics["scrollOwner"] = driver.execute_script("""
const main=document.querySelector('main'),workspace=document.querySelector('.court-records-workspace');
return {main:main?.scrollTop||0,workspace:workspace?.scrollTop||0};
""")
        driver.save_screenshot(str(output / f"{slug}-{label}-documents.png"))
        assert metrics["overflow"] <= 1, metrics
        assert not metrics["duplicateIds"], metrics
        assert not metrics["unnamed"], metrics
        assert not metrics["smallTargets"], metrics
        assert not metrics["smallFormText"], metrics
        assert not metrics["positiveTabindex"], metrics
        assert not metrics["crampedEntryFields"], metrics
        assert not metrics["filingDetails"] or metrics["filingDetails"]["contained"], metrics
        assert not metrics["wideNodes"], metrics
        assert metrics["scrollOwner"]["main"] == 0, metrics
        assert metrics["dpr"] == dpr, metrics
        assert metrics["header"]["left"] >= 0 and metrics["header"]["right"] <= width, metrics
        results[label] = metrics
    restored = restore_viewport(driver)
    assert restored["main"] == 0 and restored["workspaceRect"]["bottom"] > 0, restored
    results["restored_1440"] = restored
    return results


def build_and_download(
    driver: webdriver.Chrome, directory: Path, slug: str, page_count: int,
    busy_proof: dict[str, object] | None = None,
) -> dict[str, Path]:
    button = driver.find_element(By.CSS_SELECTOR, "[data-court-record-build]")
    panel = button.find_element(By.XPATH, "ancestor::aside[1]")
    reveal(driver, button)
    button.click()
    if busy_proof is not None:
        state = WebDriverWait(driver, 5).until(lambda _item: driver.execute_script("""
const build=arguments[0];
if (build.getAttribute('aria-busy') !== 'true') return null;
const workspace=document.getElementById('court-record-workspace');
const back=document.querySelector("button[aria-label='Back from court record']");
return {observed:true,workspace_inert:workspace.inert,
  back_disabled:back.disabled,build_disabled:build.disabled};
""", button))
        busy_proof.update(state)
    WebDriverWait(driver, 45).until(lambda _item:
        button.get_attribute("aria-busy") != "true" and "Build complete" in panel.text)
    preview = WebDriverWait(driver, 10).until(lambda item: item.find_element(
        By.CSS_SELECTOR, "[aria-label='Built court record preview']"
    ))
    WebDriverWait(driver, 20).until(
        lambda _item: len(preview.find_elements(By.CSS_SELECTOR, "[data-page-number]")) == page_count
    )
    assert not panel.find_elements(By.TAG_NAME, "iframe")
    directory.mkdir(parents=True)
    driver.execute_cdp_cmd("Page.setDownloadBehavior", {
        "behavior": "allow", "downloadPath": str(directory.resolve()),
    })
    reveal(driver, preview)
    assert driver.save_screenshot(str(directory.parent / f"{slug}-built-preview.png"))
    downloaded: dict[str, Path] = {}
    for candidate in panel.find_elements(By.CSS_SELECTOR, "button"):
        names = [span.text.strip() for span in candidate.find_elements(By.CSS_SELECTOR, "span")]
        filename = next((name for name in names if name.lower().endswith((".pdf", ".docx"))), "")
        if not filename:
            continue
        driver.execute_script("arguments[0].click()", candidate)
        target = directory / filename
        WebDriverWait(driver, 20).until(
            lambda _item: target.is_file() and not list(directory.glob("*.crdownload"))
        )
        downloaded[filename] = target
    assert downloaded, "The live build exposed no downloadable artifacts."
    restore_viewport(driver)
    return downloaded


def stale_build_contract(
    driver: webdriver.Chrome, large_file: Path, output: Path
) -> dict[str, object]:
    upload(driver, "other-filed-material", large_file)
    driver.execute_script("""
const original=Blob.prototype.arrayBuffer;
window.__courtBuildGate={original,waiting:false,release:null};
Blob.prototype.arrayBuffer=function(){
  const gate=window.__courtBuildGate;
  if(!gate.waiting){gate.waiting=true;return new Promise(resolve=>gate.release=resolve)
    .then(()=>original.call(this));}
  return original.call(this);
};
""")
    button = driver.find_element(By.CSS_SELECTOR, "[data-court-record-build]")
    panel = button.find_element(By.XPATH, "ancestor::aside[1]")
    reveal(driver, button)
    button.click()
    WebDriverWait(driver, 10).until(lambda _item:
        button.get_attribute("aria-busy") == "true" and
        driver.execute_script("return window.__courtBuildGate.waiting"))
    workspace = driver.find_element(By.ID, "court-record-workspace")
    locked = {
        "workspace_inert": driver.execute_script("return arguments[0].inert", workspace),
        "back_disabled": not driver.find_element(
            By.CSS_SELECTOR, "button[aria-label='Back from court record']"
        ).is_enabled(),
        "build_disabled": not button.is_enabled(),
    }
    assert all(locked.values()), locked
    assert driver.save_screenshot(str(output / "fca-build-busy.png"))
    changed = field(driver, "recordSubtitle")
    driver.execute_script("""
const input=arguments[0], value=arguments[1], set=Object.getOwnPropertyDescriptor(
  HTMLInputElement.prototype,'value').set;
set.call(input,value); input.dispatchEvent(new Event('input',{bubbles:true}));
""", changed, "Changed during build")
    driver.execute_script("""
const gate=window.__courtBuildGate; Blob.prototype.arrayBuffer=gate.original; gate.release();
""")
    WebDriverWait(driver, 60).until(
        lambda _item: button.get_attribute("aria-busy") != "true"
    )
    assert "changed while it was building" in panel.text.casefold(), panel.text
    assert not driver.find_elements(By.CSS_SELECTOR, "[aria-label='Built court record preview']")
    assert driver.save_screenshot(str(output / "fca-stale-build-rejected.png"))
    article(driver, large_file.name).find_element(
        By.CSS_SELECTOR, f"button[aria-label={json.dumps('Remove ' + large_file.name)}]"
    ).click()
    WebDriverWait(driver, 5).until(
        lambda item: not item.find_elements(By.XPATH, f"//span[@title={json.dumps(large_file.name)}]")
    )
    enter(changed, "Motion for procedural directions")
    return {**locked, "stale_result_rejected": True, "pages_during_race": 80}


def reopen_saved_contract(
    driver: webdriver.Chrome, profile_id: str, expected: dict[str, str],
    filenames: list[str], output: Path,
) -> dict[str, object]:
    title = driver.find_element(By.CSS_SELECTOR, "[data-workspace-header] h1").text
    started = time.perf_counter()
    back = WebDriverWait(driver, 10).until(lambda item: next((button for button in
        item.find_elements(By.CSS_SELECTOR, "button[aria-label='Back from court record']")
        if button.is_enabled() and button.is_displayed()), False))
    visible_wait_ms = round((time.perf_counter() - started) * 1000, 1)
    back.click()
    WebDriverWait(driver, 15).until(lambda item: item.find_elements(
        By.XPATH, "//button[normalize-space()='Open saved record']"
    ))
    driver.find_element(By.XPATH, "//button[normalize-space()='Open saved record']").click()
    dialog = WebDriverWait(driver, 10).until(
        lambda item: item.find_element(By.CSS_SELECTOR, "dialog[open]")
    )
    assert driver.save_screenshot(str(output / "saved-record-picker.png"))
    dialog.find_element(By.XPATH, ".//button[time]").click()
    WebDriverWait(driver, 30).until(lambda item: item.find_element(
        By.CSS_SELECTOR, "[data-court-record-chooser]"
    ).get_attribute("data-selected-profile") == profile_id)
    WebDriverWait(driver, 30).until(lambda item: all(
        row.get_attribute("aria-busy") != "true"
        for row in item.find_elements(By.CSS_SELECTOR, "[data-entry-id]")
    ))
    assert all(field(driver, key).get_attribute("value") == value
               for key, value in expected.items())
    assert all(driver.find_elements(By.CSS_SELECTOR, f"span[title='{name}']")
               for name in filenames)
    assert driver.find_element(By.CSS_SELECTOR, "[data-workspace-header] h1").text == title
    assert driver.save_screenshot(str(output / "saved-record-reopened.png"))
    return {"title": title, "files": filenames, "values": expected,
            "visible_wait_ms": visible_wait_ms}


def save_and_filter_contract(
    driver: webdriver.Chrome, output_file: Path, output: Path, notice_file: Path
) -> dict[str, object]:
    mode = "standalone" if urlsplit(driver.current_url).path.endswith(
        "court-records.html"
    ) else "beaver"
    save = driver.find_element(By.XPATH,
        "//button[normalize-space()='Save output' or normalize-space()='Save to Library']")
    save.click()
    WebDriverWait(driver, 20).until(lambda item: "1 file saved" in item.find_element(
        By.CSS_SELECTOR, "[aria-label='Build output']"
    ).text)
    back = WebDriverWait(driver, 10).until(lambda item: next((button for button in
        item.find_elements(By.CSS_SELECTOR, "button[aria-label='Back from court record']")
        if button.is_enabled() and button.is_displayed()), False))
    back.click()
    WebDriverWait(driver, 10).until(lambda item: item.find_elements(
        By.XPATH, "//button[normalize-space()='New court record']"
    ))
    choose_new_profile(driver, "fc-motion-record-moving")
    upload(driver, "notice-motion", notice_file)
    select_filing_party(driver, "North Prairie Ltd.")

    for key, value in {"courtFileNumber": "T-982-19", "counselName": "Jordan Lee",
                       "counselAddress": "100 Legal Avenue\nCalgary, Alberta",
                       "counselPhone": "403-555-0100",
                       "counselEmail": "jlee@example.test",
                       "recordSubtitle": "Motion for procedural directions"}.items():
        enter(field(driver, key), value)

    driver.find_element(By.CSS_SELECTOR, "[data-kind-id='moving-evidence'] button").click()
    dialog = WebDriverWait(driver, 20).until(
        lambda item: item.find_element(By.CSS_SELECTOR, "dialog[open]")
    )
    WebDriverWait(driver, 20).until(lambda _item: "No matching files." in dialog.text)
    assert output_file.name not in dialog.text
    assert driver.save_screenshot(str(output / "incompatible-output-excluded.png"))
    dialog.send_keys(Keys.ESCAPE)
    application_under = field(driver, "applicationUnder")
    reveal(driver, application_under)
    enter(application_under, "Federal Courts Act, section 18.1")
    build = driver.find_element(By.CSS_SELECTOR, "[data-court-record-build]")
    reveal(driver, build)
    build.click()
    missing_focus = WebDriverWait(driver, 5).until(lambda item:
        item.switch_to.active_element.get_attribute("id")
        if item.switch_to.active_element.get_attribute("id") ==
        "court-record-written-representations-file" else False)
    return {"available": True, "mode": mode, "output": output_file.name,
            "incompatible_output_excluded": True, "missing_document_focus": missing_focus}


def save_pdf_page(document: fitz.Document, page_number: int, output: Path) -> None:
    document[page_number].get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False).save(output)


def inspect_affidavit(path: Path, output: Path) -> dict[str, object]:
    with fitz.open(path) as document:
        pages = [page.get_text() for page in document]
        toc = [row[1] for row in document.get_toc()]
        assert document.page_count == 4, pages
        assert [page.get_label() for page in document] == ["1", "2", "3", "4"]
        assert not document.is_encrypted and all(text.strip() for text in pages)
        assert 'This is Exhibit "A" referred to in the Affidavit of:' in pages[1]
        assert "SERVICE AGREEMENT" in pages[2]
        assert 'This is Exhibit "B" referred to in the Affidavit.' in pages[3]
        assert "PROJECT LEDGER" in pages[3]
        assert toc.count("Exhibit A certificate") == 1, toc
        assert "Exhibit B certificate" not in toc, toc
        for page_number, name in ((0, "affidavit-output"), (1, "exhibit-a-certificate"),
                                  (3, "exhibit-b-with-certificate")):
            save_pdf_page(document, page_number, output / f"{name}.png")
        return {"filename": path.name, "sha256": sha256(path), "bytes": path.stat().st_size,
                "pages": document.page_count, "bookmarks": toc}


def inspect_federal(
    path: Path, output: Path, slug: str = "federal-motion",
    court: str = "FEDERAL COURT", pages_expected: int = 4,
    extra_text: tuple[str, ...] = (),
) -> dict[str, object]:
    with fitz.open(path) as document:
        text = "\n".join(page.get_text() for page in document)
        toc = [row[1] for row in document.get_toc()]
        links = sum(len(page.get_links()) for page in document)
        assert document.page_count == pages_expected, text
        assert [page.get_label() for page in document] == [
            str(value) for value in range(1, pages_expected + 1)
        ]
        assert not document.is_encrypted
        for value in (court, "MOTION RECORD", "INDEX", "NOTICE OF MOTION",
                      "WRITTEN REPRESENTATIONS OF THE APPLICANT", *extra_text):
            assert value in text, value
        assert links >= 2 and len(toc) >= 3, (links, toc)
        for page_number, name in ((0, f"{slug}-cover"), (1, f"{slug}-index")):
            save_pdf_page(document, page_number, output / f"{name}.png")
        return {"filename": path.name, "sha256": sha256(path), "bytes": path.stat().st_size,
                "pages": document.page_count, "bookmarks": toc, "internal_links": links}


def inspect_appeal_record(
    record_path: Path, transcript_path: Path, source_transcript: Path, output: Path
) -> dict[str, object]:
    with fitz.open(record_path) as record, fitz.open(transcript_path) as transcript:
        pages = [page.get_text() for page in record]
        text = "\n".join(pages)
        toc = [row[1] for row in record.get_toc()]
        links = sum(len(page.get_links()) for page in record)
        corner = record[0].get_pixmap(alpha=False).pixel(5, 5)[:3]
        index_text = " ".join(pages[1].split())
        assert record.page_count == 6, pages
        assert [page.get_label() for page in record] == [str(value) for value in range(1, 7)]
        assert corner[0] > 240 and corner[1] < 15 and corner[2] < 15, corner
        for value in ("COURT OF APPEAL OF ALBERTA", "APPEAL RECORD", "Table of Contents",
                      "STATEMENT OF CLAIM", "REASONS FOR JUDGMENT", "FORMAL ORDER",
                      "NOTICE OF APPEAL"):
            assert value in text, value
        cover_text = " ".join(pages[0].split())
        for value in ("Jordan Lee", "403-555-0100", "403-555-0101", "Riley Counsel",
                      "200 Court Street", "780-555-0123"):
            assert value in cover_text, (value, cover_text)
        for description in ("STATEMENT OF CLAIM", "REASONS FOR JUDGMENT", "Formal Order",
                            "NOTICE OF APPEAL"):
            assert description in index_text, (description, index_text)
        assert "Part 2 - Formal order or decision" not in index_text, index_text
        assert index_text.count("Part 2") == 1, index_text
        assert "ORAL PROCEEDINGS TRANSCRIPT" not in text
        assert links >= 4 and len(toc) >= 6, (links, toc)
        assert transcript.page_count == 3 and not transcript.is_encrypted
        assert [page.get_label() for page in transcript] == ["", "i", "1"]
        assert "ORAL PROCEEDINGS TRANSCRIPT" in transcript[2].get_text()
        assert transcript_path.name == source_transcript.name
        assert sha256(transcript_path) == sha256(source_transcript)
        save_pdf_page(record, 0, output / "appeal-record-cover.png")
        save_pdf_page(record, 1, output / "appeal-record-index.png")
        return {
            "record": {"filename": record_path.name, "sha256": sha256(record_path),
                       "pages": record.page_count, "bookmarks": toc,
                       "internal_links": links, "cover_rgb": list(corner),
                       "filename_descriptions": True},
            "transcript": {"filename": transcript_path.name,
                           "sha256": sha256(transcript_path), "pages": transcript.page_count,
                           "page_labels": [page.get_label() for page in transcript],
                           "preserved_byte_for_byte": True},
        }


def inspect_no_oral_record(path: Path, description: str, output: Path) -> dict[str, object]:
    with fitz.open(path) as document:
        text = "\n".join(page.get_text() for page in document)
        assert document.page_count == 6
        assert " ".join(description.split()) in " ".join(text.split()), text
        assert "ORAL PROCEEDINGS TRANSCRIPT" not in text
        save_pdf_page(document, 1, output / "appeal-record-no-oral-record-index.png")
        return {"filename": path.name, "sha256": sha256(path),
                "pages": document.page_count, "description": description}


def check_affidavit(
    driver: webdriver.Chrome, url: str, fixtures: dict[str, Path], output: Path
) -> dict[str, object]:
    profile_ms, cold_navigation, chooser = reset_workspace(
        driver, url, "ab-kb-affidavit-exhibits", output
    )
    assert len(driver.find_elements(By.CSS_SELECTOR, '#court-record-workspace [data-kind-id]')) == 1
    assert not driver.find_elements(By.CSS_SELECTOR, '.court-record-build-panel')
    assert not driver.find_elements(By.ID, 'cover-courtFileNumber')
    driver.save_screenshot(str(output / 'affidavit-initial.png'))
    upload(driver, "affidavit", fixtures["affidavit"])
    WebDriverWait(driver, 20).until(lambda item: field(item, "deponent").get_attribute("value"))
    assert driver.find_element(By.ID, "court-record-party-style").is_displayed()
    enter(field(driver, "registry"), "Edmonton")
    driver.find_element(By.XPATH, "//button[normalize-space()='Add plaintiff']").click()
    enter(driver.find_elements(By.CSS_SELECTOR, "[data-party-group='Plaintiff'] input")[-1], "Existing Plaintiff")
    expected = {
        "courtFileNumber": "2401-12345", "registry": "Edmonton", "affidavitNumber": "2",
        "deponent": "Alexandra Smith", "swornDate": "AUGUST 29, 2026", "swornPlace": "CALGARY",
    }
    WebDriverWait(driver, 20).until(lambda item: all(
        field(item, key).get_attribute("value") == value for key, value in expected.items()
    ))
    plaintiffs = [item.get_attribute("value") for item in driver.find_elements(
        By.CSS_SELECTOR, "[data-party-group='Plaintiff'] input")]
    defendants = [item.get_attribute("value") for item in driver.find_elements(
        By.CSS_SELECTOR, "[data-party-group='Defendant'] input")]
    assert plaintiffs == ["NORTH PRAIRIE LTD.", "Existing Plaintiff"], plaintiffs
    assert defendants == ["RIVERSTONE INC."], defendants
    select_filing_party(driver, "Existing Plaintiff")
    for label in ("A", "B"):
        WebDriverWait(driver, 10).until(lambda item, value=label: exhibit_slot(item, value))

    upload(driver, "exhibit", fixtures["exhibit_b"])
    upload(driver, "exhibit", fixtures["ambiguous"])
    assert contains_file(exhibit_slot(driver, "B"), fixtures["exhibit_b"].name)
    assert contains_file(unassigned_pool(driver), fixtures["ambiguous"].name)

    drag_events = drag(driver, article(driver, fixtures["ambiguous"].name),
                       exhibit_slot(driver, "A"))
    assert any(event["type"] == "drop" and event["trusted"]
               for event in drag_events["events"]), drag_events
    WebDriverWait(driver, 5).until(
        lambda item: contains_file(exhibit_slot(item, "A"), fixtures["ambiguous"].name)
    )
    drag(driver, article(driver, fixtures["ambiguous"].name), unassigned_pool(driver))
    WebDriverWait(driver, 5).until(
        lambda item: contains_file(unassigned_pool(item), fixtures["ambiguous"].name)
    )
    drag(driver, article(driver, fixtures["ambiguous"].name), exhibit_slot(driver, "A"))
    WebDriverWait(driver, 5).until(
        lambda item: contains_file(exhibit_slot(item, "A"), fixtures["ambiguous"].name)
    )
    assert contains_file(exhibit_slot(driver, "B"), fixtures["exhibit_b"].name)
    toggle_disclosure(driver, exhibit_slot(driver, "A").find_element(By.TAG_NAME, "details"))

    keyboard = keyboard_contract(driver)
    viewports = viewport_contract(driver, output, "affidavit", "[data-kind-id='exhibit']")
    busy: dict[str, object] = {}
    downloads = build_and_download(
        driver, output / "affidavit-downloads", "affidavit", 4, busy
    )
    pdf = next(path for path in downloads.values() if path.suffix.lower() == ".pdf")
    inspected = inspect_affidavit(pdf, output)
    reopened = reopen_saved_contract(
        driver, "ab-kb-affidavit-exhibits", expected,
        [fixtures[key].name for key in ("affidavit", "exhibit_b", "ambiguous")], output,
    )
    return {"profile": "ab-kb-affidavit-exhibits", "reset_and_profile_ms": profile_ms,
            "cold_navigation": cold_navigation, "chooser": chooser,
            "propagated": expected, "pointer_drag": drag_events,
            "keyboard_focus": keyboard, "viewports": viewports,
            "build_busy": busy, "saved_reopen": reopened, "output": inspected}


def complete_motion_cover(driver: webdriver.Chrome) -> None:
    for field_id, value in {
        "counselName": "Jordan Lee", "counselAddress": "100 Legal Avenue\nCalgary, Alberta",
        "counselPhone": "403-555-0100", "counselEmail": "jlee@example.test",
        "recordSubtitle": "Motion for procedural directions",
        "applicationUnder": "Federal Courts Act, section 18.1",
    }.items():
        reveal(driver, field(driver, field_id))
        enter(field(driver, field_id), value)


def check_federal(
    driver: webdriver.Chrome, url: str, fixtures: dict[str, Path], output: Path
) -> dict[str, object]:
    profile_ms, navigation, _chooser = reset_workspace(
        driver, url, "fc-motion-record-moving"
    )
    assert not driver.find_elements(By.CSS_SELECTOR, "[draggable='true']")
    upload(driver, "notice-motion", fixtures["notice"])
    WebDriverWait(driver, 10).until(
        lambda item: field(item, "courtFileNumber").get_attribute("value") == "T-982-19"
    )
    assert driver.find_element(By.CSS_SELECTOR, "[data-party-group='Applicant'] input").get_attribute("value") == "NORTH PRAIRIE LTD."
    assert driver.find_element(By.CSS_SELECTOR, "[data-party-group='Respondent'] input").get_attribute("value") == "RIVERSTONE INC."
    select_filing_party(driver, "North Prairie Ltd.")
    complete_motion_cover(driver)
    upload(driver, "written-representations", fixtures["representations"])
    assert not driver.find_elements(By.CSS_SELECTOR, "[draggable='true']")
    viewports = viewport_contract(
        driver, output, "federal-motion", "[data-kind-id='notice-motion']"
    )
    busy: dict[str, object] = {}
    downloads = build_and_download(
        driver, output / "federal-motion-downloads", "federal-motion", 4, busy
    )
    pdf = next(path for path in downloads.values() if path.suffix.lower() == ".pdf")
    inspected = inspect_federal(pdf, output)
    nested = save_and_filter_contract(driver, pdf, output, fixtures["notice"])
    return {"profile": "fc-motion-record-moving", "reset_and_profile_ms": profile_ms,
            "navigation": navigation,
            "viewports": viewports,
            "build_busy": busy, "saved_output_filter": nested, "output": inspected}


def check_federal_appeal(
    driver: webdriver.Chrome, url: str, fixtures: dict[str, Path], output: Path
) -> dict[str, object]:
    profile_ms, navigation, chooser = reset_workspace(
        driver, url, "fca-motion-record-moving"
    )
    upload(driver, "notice-motion", fixtures["notice"])
    WebDriverWait(driver, 10).until(
        lambda item: field(item, "courtFileNumber").get_attribute("value") == "T-982-19"
    )
    select_filing_party(driver, "North Prairie Ltd.")
    complete_motion_cover(driver)
    upload(driver, "written-representations", fixtures["representations"])
    race = stale_build_contract(driver, fixtures["large"], output)

    beaver = not urlsplit(driver.current_url).path.endswith("court-records.html")
    if beaver:
        upload(driver, "oral-hearing-request", fixtures["oral_request"])
        assert article(driver, fixtures["oral_request"].name)
    viewports = viewport_contract(
        driver, output, "fca-motion", "[data-kind-id='oral-hearing-request']"
        if beaver else "[data-kind-id='notice-motion']",
    )
    busy: dict[str, object] = {}
    pages = 5 if beaver else 4
    downloads = build_and_download(
        driver, output / "fca-motion-downloads", "fca-motion", pages, busy
    )
    pdf = next(path for path in downloads.values() if path.suffix.lower() == ".pdf")
    return {"profile": "fca-motion-record-moving", "reset_and_profile_ms": profile_ms,
            "navigation": navigation, "chooser": chooser, "stale_build": race,
            "docx_rendition": {"tested": beaver, "filename": fixtures["oral_request"].name
                               if beaver else None},
            "viewports": viewports, "build_busy": busy,
            "output": inspect_federal(
                pdf, output, "fca-motion", "FEDERAL COURT OF APPEAL", pages,
                ("REQUEST FOR ORAL HEARING",) if beaver else (),
            )}


def check_appeal_record(
    driver: webdriver.Chrome, url: str, fixtures: dict[str, Path], output: Path
) -> dict[str, object]:
    profile_ms, navigation, _chooser = reset_workspace(driver, url, "ab-ca-appeal-record")
    style = Select(driver.find_element(By.ID, "court-record-party-style"))
    style.select_by_visible_text(next(option.text for option in style.options
                                      if "plaintiff appeals" in option.text))
    for selector, value in (("[data-party-group='Appellant'] input", "North Prairie Ltd."),
                            ("[data-party-group='Respondent'] input", "Riverstone Inc.")):
        enter(driver.find_element(By.CSS_SELECTOR, selector), value)
    select_filing_party(driver, "North Prairie Ltd.")
    for field_id, value in {
        "courtFileNumber": "2403-00789", "registry": "Calgary",
        "counselName": "Jordan Lee", "counselAddress": "100 Legal Avenue\nCalgary, Alberta",
        "counselPhone": "403-555-0100", "counselEmail": "jlee@example.test",
        "counselFax": "403-555-0101",
    }.items():
        reveal(driver, field(driver, field_id))
        enter(field(driver, field_id), value)
    assert field(driver, "counselFax").get_dom_attribute("required") is not None
    for label, value in {
        "Lawyer or filing person": "Riley Counsel",
        "Address for service": "200 Court Street\nEdmonton, Alberta",
        "Telephone": "780-555-0123",
    }.items():
        control = driver.find_element(
            By.CSS_SELECTOR, f'[aria-label="{label} for Riverstone Inc."]'
        )
        assert control.get_dom_attribute("required") is not None
        reveal(driver, control)
        enter(control, value)
    upload(driver, "part-3-transcript", fixtures["appeal_transcript_prior"])
    for kind_id, fixture in (
        ("part-1-pleading", "appeal_pleading"),
        ("part-2-reasons", "appeal_reasons"),
        ("part-2-order", "appeal_order"),
        ("part-2-notice", "appeal_notice"),
        ("part-3-transcript", "appeal_transcript"),
    ):
        upload(driver, kind_id, fixtures[fixture])
    transcript_slot = driver.find_element(By.CSS_SELECTOR, "[data-kind-id='part-3-transcript']")
    assert not transcript_slot.find_elements(
        By.CSS_SELECTOR, f"span[title='{fixtures['appeal_transcript_prior'].name}']"
    )
    assert len(transcript_slot.find_elements(By.CSS_SELECTOR, "[data-entry-id]")) == 1
    pleading_date = article(driver, fixtures["appeal_pleading"].name).find_element(
        By.CSS_SELECTOR, "input[id$='-date']"
    )
    assert pleading_date.get_attribute("value") == ""
    enter(pleading_date, "January 15, 2026")
    propagated = {"lowerCourtFileNumber": "2301-00456",
                  "decisionMaker": "The Honourable Justice A. Ng",
                  "decisionDate": "March 4, 2026",
                  "decisionFileDate": "March 5, 2026"}
    WebDriverWait(driver, 10).until(lambda item: all(
        field(item, key).get_attribute("value") == value for key, value in propagated.items()))
    expected_titles = {
        "appeal_pleading": "STATEMENT OF CLAIM",
        "appeal_reasons": "REASONS FOR JUDGMENT",
        "appeal_order": "Formal Order",
        "appeal_notice": "NOTICE OF APPEAL",
        "appeal_transcript": "TRANSCRIPT OF PROCEEDINGS",
    }
    assert {key: article(driver, fixtures[key].name).find_element(
        By.CSS_SELECTOR, "input[id$='-title']").get_attribute("value")
        for key in expected_titles} == expected_titles
    assert not driver.find_elements(By.CSS_SELECTOR, "[draggable='true']")
    viewports = viewport_contract(
        driver, output, "appeal-record", "[data-kind-id='part-3-transcript']"
    )
    busy: dict[str, object] = {}
    downloads = build_and_download(
        driver, output / "appeal-record-downloads", "appeal-record", 6, busy
    )
    assert len(downloads) == 2, downloads
    transcript = downloads[fixtures["appeal_transcript"].name]
    record = next(path for path in downloads.values() if path != transcript)
    inspected = inspect_appeal_record(
        record, transcript, fixtures["appeal_transcript"], output)

    transcript_name = fixtures["appeal_transcript"].name
    article(driver, transcript_name).find_element(
        By.CSS_SELECTOR, f"button[aria-label='Remove {transcript_name}']",
    ).click()
    no_record = WebDriverWait(driver, 10).until(lambda item: item.find_element(
        By.CSS_SELECTOR, "[data-kind-id='part-3-no-oral-record']"
    ))
    description = "There is no oral record that can be transcribed for Part 3, Transcripts"
    description_input = no_record.find_element(By.CSS_SELECTOR, "input")
    assert description_input.get_attribute("value") == ""
    enter(description_input, description)
    no_record_busy: dict[str, object] = {}
    no_record_downloads = build_and_download(
        driver, output / "appeal-no-oral-record-downloads",
        "appeal-record-no-oral-record", 6, no_record_busy,
    )
    assert len(no_record_downloads) == 1, no_record_downloads
    no_record_pdf = next(iter(no_record_downloads.values()))
    return {"profile": "ab-ca-appeal-record", "reset_and_profile_ms": profile_ms,
            "navigation": navigation, "propagated": propagated, "viewports": viewports,
            "build_busy": busy, "outputs": inspected,
            "description_only": {"build_busy": no_record_busy,
                                 "output": inspect_no_oral_record(
                                     no_record_pdf, description, output)}}


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Exercise the production Court Record Builder through Chrome."
    )
    parser.add_argument("--url", help="Optional served court-records entry point.")
    parser.add_argument("--headed", action="store_true")
    parser.add_argument("--chooser-only", action="store_true", help="Check selection, cancellation and responsive dialog geometry without building PDFs.")
    parser.add_argument("--artifacts", type=Path, help="Screenshot and proof-file directory.")
    parser.add_argument("--built-asset", type=Path, help="Exact main bundle served for this proof.")
    parser.add_argument("--server-pid", type=int)
    parser.add_argument("--server-command")
    args = parser.parse_args()
    RESPONSIVE_SCROLL_RESETS.clear()

    with tempfile.TemporaryDirectory(prefix="beaver-court-record-browser-") as temporary:
        session = Path(temporary)
        output = args.artifacts.resolve() if args.artifacts else session / "artifacts"
        output.mkdir(parents=True, exist_ok=True)
        fixtures = {} if args.chooser_only else make_fixtures(session / "fixtures")

        server = None
        thread = None
        url = args.url
        if url and url.endswith("/court-records"):
            url += "/"  # Vite otherwise serves the standalone court-records.html entry.
        if not url:
            dist = ROOT / "frontend" / "dist"
            if not (dist / "court-records.html").is_file():
                raise RuntimeError("Build the production frontend before running this test.")
            server = ThreadingHTTPServer(("127.0.0.1", 0), partial(
                QuietStaticHandler, directory=str(dist)
            ))
            thread = threading.Thread(target=server.serve_forever, daemon=True)
            thread.start()
            url = f"http://127.0.0.1:{server.server_port}/court-records.html"

        driver = start_browser(session / "chrome-profile", args.headed)
        started = time.perf_counter()
        try:
            if args.chooser_only:
                proofs = {}
                for profile_id in CHOICES:
                    print(f"Court Records chooser: {profile_id}", flush=True)
                    case_output = output / profile_id
                    case_output.mkdir(exist_ok=True)
                    proofs[profile_id] = reset_workspace(driver, url, profile_id, case_output)[2]
                    courts = driver.find_elements(By.CSS_SELECTOR, "button[aria-label^='Change format:']")
                    if not courts:
                        assert len(driver.find_elements(By.CSS_SELECTOR, '#court-record-workspace [data-kind-id]')) == 1
                        continue
                    courts[0].click()
                    dialog = driver.find_element(By.CSS_SELECTOR, "dialog[open]")
                    dialog.find_element(By.XPATH,
                        ".//*[@aria-label='Jurisdiction']//button[normalize-space()='No court preset']").click()
                    assert driver.find_elements(By.CSS_SELECTOR, "dialog[open]") == [dialog]
                    dialog.send_keys(Keys.ESCAPE)
                    assert not driver.find_elements(By.CSS_SELECTOR, "dialog[open]")
                    assert driver.find_element(By.CSS_SELECTOR, "[data-court-record-chooser]").get_attribute("data-selected-profile") == profile_id
                errors = [row for row in driver.get_log("browser") if row["level"] == "SEVERE"]
                assert not errors, errors
                (output / "chooser-proof.json").write_text(json.dumps(proofs, indent=2), encoding="utf-8")
                print(f"PASS Court Records single-dialog selection; screenshots: {output}", flush=True)
                return 0
            result = {
                "schema_version": "beaver.court-record-browser-test.v4",
                "created_at": datetime.now(timezone.utc).isoformat(),
                "url": url,
                "affidavit": check_affidavit(driver, url, fixtures, output),
                "federal_motion": check_federal(driver, url, fixtures, output),
                "federal_appeal_motion": check_federal_appeal(
                    driver, url, fixtures, output),
                "alberta_appeal_record": check_appeal_record(
                    driver, url, fixtures, output),
            }
            logs = driver.get_log("browser")
            severe = [entry for entry in logs if entry["level"] == "SEVERE" and not (
                "/preparation?" in entry["message"] and "409 (Conflict)" in entry["message"]
            )]
            result["browser_logs"] = logs
            result["responsive_scroll_resets"] = RESPONSIVE_SCROLL_RESETS
            result["elapsed_seconds"] = round(time.perf_counter() - started, 3)
            if args.built_asset:
                loaded = driver.execute_async_script(r"""
const done=arguments[arguments.length-1],entry=performance.getEntriesByType('resource')
  .find(({name})=>/\/assets\/courtRecords-[^/]+\.js$/u.test(name));
if(!entry){done(null);return;}
fetch(entry.name,{cache:'no-store'}).then(async response=>{
  const bytes=await response.arrayBuffer(),digest=await crypto.subtle.digest('SHA-256',bytes);
  done({url:entry.name,sha256:[...new Uint8Array(digest)]
    .map(value=>value.toString(16).padStart(2,'0')).join('')});
}).catch(error=>done({error:String(error)}));
""")
                built_hash = sha256(args.built_asset)
                assert loaded and loaded.get("sha256") == built_hash, (loaded, built_hash)
                result["served_build"] = {"server_pid": args.server_pid,
                    "server_command": args.server_command,
                    "built_asset": str(args.built_asset.resolve()), "built_sha256": built_hash,
                    "loaded_url": loaded["url"], "loaded_sha256": loaded["sha256"]}
            failures = [*(f"SEVERE browser log: {item['message']}" for item in severe)]
            if RESPONSIVE_SCROLL_RESETS:
                failures.append(
                    f"Responsive scroll-owner transition needed {len(RESPONSIVE_SCROLL_RESETS)} ancestor reset(s)."
                )
            for name in ("affidavit", "federal_motion", "federal_appeal_motion",
                         "alberta_appeal_record"):
                busy = result[name]["build_busy"]
                if not busy.get("observed") or not all(busy.get(key) for key in (
                    "workspace_inert", "back_disabled", "build_disabled"
                )):
                    failures.append(f"{name}: build busy state was not observable and locked: {busy}")
            result["failures"] = failures
            receipt = output / "court-records-live-proof.json"
            receipt.write_text(json.dumps(result, indent=2), encoding="utf-8")
            print(json.dumps(result, indent=2))
            print(f"ARTIFACTS {output}")
            return 1 if failures else 0
        except Exception as error:
            driver.save_screenshot(str(output / "failure.png"))
            state = driver.execute_script("""
const main=document.querySelector('main'),root=document.querySelector('#root'),
  workspace=document.querySelector('.court-records-workspace');
return {url:location.href,readyState:document.readyState,historyState:history.state,
  activeElement:document.activeElement?.outerHTML?.slice(0,1000),
  bodyText:document.body.innerText,rootText:root?.innerText,
  mainHtml:main?.innerHTML,workspacePresent:!!workspace,
  workspaceHtml:workspace?.outerHTML,
  fileInputs:[...document.querySelectorAll('input[type=file]')].map(node=>({
    connected:node.isConnected,outerHTML:node.outerHTML,rect:(()=>{const r=node.getBoundingClientRect();
      return {x:r.x,y:r.y,width:r.width,height:r.height}})()})),
  navigation:performance.getEntriesByType('navigation').map(({name,duration,responseStart,
    domContentLoadedEventEnd,loadEventEnd})=>({name,duration,responseStart,
      domContentLoadedEventEnd,loadEventEnd})),
  resources:performance.getEntriesByType('resource').slice(-30).map(({name,initiatorType,
    duration,transferSize,decodedBodySize})=>({name,initiatorType,duration,transferSize,decodedBodySize}))};
""")
            state["browserLogs"] = driver.get_log("browser")
            state["exception"] = {
                "type": type(error).__name__, "message": str(error),
                "traceback": traceback.format_exc(),
            }
            (output / "failure-state.json").write_text(
                json.dumps(state, indent=2), encoding="utf-8"
            )
            (output / "failure.html").write_text(driver.page_source, encoding="utf-8")
            print("LIVE_FAILURE", driver.current_url)
            print(driver.find_element(By.TAG_NAME, "body").text[:4000]
                  .encode("ascii", "backslashreplace").decode())
            print(json.dumps(driver.get_log("browser"), indent=2))
            raise
        finally:
            driver.quit()
            if server:
                server.shutdown()
                server.server_close()
            if thread:
                thread.join(timeout=2)


if __name__ == "__main__":
    raise SystemExit(main())
