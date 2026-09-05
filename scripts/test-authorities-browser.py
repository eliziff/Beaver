from __future__ import annotations

import argparse
import base64
import hashlib
import json
import subprocess
import tempfile
import time
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import parse_qs, urlparse
from xml.sax.saxutils import escape

import fitz
from docx import Document
from selenium import webdriver
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.common.exceptions import StaleElementReferenceException
from selenium.webdriver.support.ui import Select
from selenium.webdriver.support.wait import WebDriverWait


OAKES = "R v Oakes, [1986] 1 SCR 103, 1986 CanLII 46 (SCC)"
FALSE_POSITIVE = "2024 ABKB 999"
RECONSTRUCTED = "2021 BCCA 222"
REAL_CITATIONS = ("2009 SCC 32", "2016 SCC 27", "2026 SCC 16", RECONSTRUCTED)
CITATION_RENDER_BUDGET_MS = 50
CHROME = next((path for path in (
    Path(r"C:\Program Files\Google\Chrome\Application\chrome.exe"),
    Path(r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe"),
) if path.is_file()), None)
CHROMEDRIVER = next(iter(sorted(Path.home().parent.glob(
    r"*/.cache/selenium/chromedriver/win64/*/chromedriver.exe"), reverse=True)), None)


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def pdf_fixture(path: Path, title: str, detail: str) -> Path:
    pdf = fitz.open()
    page = pdf.new_page(width=612, height=792)
    page.insert_text((72, 86), title, fontsize=18)
    page.insert_textbox(fitz.Rect(72, 125, 540, 700), detail, fontsize=11)
    pdf.set_metadata({"title": title, "subject": detail})
    pdf.save(path)
    pdf.close()
    return path


def add_footnote(path: Path, text: str) -> None:
    with zipfile.ZipFile(path) as source:
        files = {name: source.read(name) for name in source.namelist()}
    document = files["word/document.xml"].decode()
    assert "[[FOOTNOTE]]" in document
    files["word/document.xml"] = document.replace(
        "<w:t>[[FOOTNOTE]]</w:t>", '<w:footnoteReference w:id="2"/>').encode()
    types = files["[Content_Types].xml"].decode()
    files["[Content_Types].xml"] = types.replace("</Types>",
        '<Override PartName="/word/footnotes.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/></Types>').encode()
    rels = files["word/_rels/document.xml.rels"].decode()
    files["word/_rels/document.xml.rels"] = rels.replace("</Relationships>",
        '<Relationship Id="rIdAuthoritiesFootnote" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes" Target="footnotes.xml"/></Relationships>').encode()
    files["word/footnotes.xml"] = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f'<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
        f'<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>'
        f'<w:footnote w:id="2"><w:p><w:r><w:footnoteRef/></w:r><w:r><w:t xml:space="preserve">{escape(text)}</w:t></w:r></w:p></w:footnote>'
        f'</w:footnotes>').encode()
    patched = path.with_name(f"{path.stem}-patched.docx")
    with zipfile.ZipFile(patched, "w", zipfile.ZIP_DEFLATED) as target:
        for name, content in files.items():
            target.writestr(name, content)
    patched.replace(path)


def fixtures(directory: Path) -> tuple[Path, list[Path], Path]:
    source = directory / "real-authorities-smoke.docx"
    document = Document()
    document.add_heading("Written argument", level=1)
    document.add_paragraph("The governing framework is stated in R v Grant, 2009 SCC 32 at para 29.")
    document.add_paragraph(f"😀 The proportionality analysis originates in {OAKES}.")
    document.add_paragraph(f"The internal file marker {FALSE_POSITIVE} is not an authority.")
    document.add_paragraph("The current publisher source is Ahluwalia v Ahluwalia, 2026 SCC 16.")
    document.add_paragraph("The source-text fallback is Neufeld v Hansman, 2021 BCCA 222.")
    document.add_paragraph("😀 Background context. " + "The record supplies additional context. " * 45 +
        "Delay is addressed in R v Jordan, 2016 SCC 27 at para 47. " +
        "The conclusion follows from the cited framework. " * 45)
    document.add_paragraph("[[FOOTNOTE]]")
    document.save(source)
    add_footnote(source, f"See R v Grant, 2009 SCC 32; {OAKES}; Ibid at para 31; and R v Jordan, 2016 SCC 27.")
    pdfs = []
    for filename, title, citation in (
        ("2009scc32.pdf", "R v Grant", "2009 SCC 32\n\n[29] The governing framework is stated here."),
        ("R v Oakes.pdf", "R v Oakes", "[1986] 1 SCR 103"),
        ("R v Jordan.pdf", "R v Jordan", "2016 SCC 27"),
        ("R v Grant replacement.pdf", "R v Grant", "2009 SCC 32\n\nReplacement original."),
        ("Hearing transcript.pdf", "Hearing transcript", "Supplemental book material."),
        ("Hearing transcript replacement.pdf", "Hearing transcript",
         "Replacement supplemental book material."),
        ("Authorities Smoke Cover.pdf", "Authorities Smoke Book", "Custom cover."),
    ):
        pdfs.append(pdf_fixture(directory / filename, title,
            f"{citation}\n\nValid local source PDF used by the Authorities production smoke."))
    filing = pdf_fixture(directory / "appeal-factum.pdf", "Appeal Factum",
        f"R v Grant, 2009 SCC 32\n\n{OAKES}\n\nR v Jordan, 2016 SCC 27")
    return source, pdfs, filing


def chrome(profile: Path, headed: bool) -> webdriver.Chrome:
    if not CHROME:
        raise RuntimeError("Google Chrome is required.")
    options = webdriver.ChromeOptions()
    options.binary_location = str(CHROME)
    if not headed:
        options.add_argument("--headless=new")
    for flag in ("--disable-gpu", "--no-sandbox", "--disable-dev-shm-usage",
                 "--disable-crash-reporter", "--no-first-run"):
        options.add_argument(flag)
    options.add_argument(f"--user-data-dir={profile}")
    options.set_capability("goog:loggingPrefs", {"browser": "ALL", "performance": "ALL"})
    driver = webdriver.Chrome(options=options,
        service=Service(str(CHROMEDRIVER)) if CHROMEDRIVER else None)
    driver.execute_cdp_cmd("Network.enable", {})
    driver.execute_cdp_cmd("Network.setCacheDisabled", {"cacheDisabled": True})
    driver.execute_cdp_cmd("Page.addScriptToEvaluateOnNewDocument", {"source": r"""
window.__authoritiesSmoke={cls:0,shifts:[],longTasks:[],canliiClicks:[],canliiLinkEvents:[],openedUrls:[],
  requestTimings:[],importUi:null,fakeCanlii:null,errors:[],builds:0,buildStarts:0,
  gateBuild:false,releaseBuild:null,lastBuild:null,lastPrepared:null,lastBookPart:null};
try { new PerformanceObserver(list => list.getEntries().forEach(e => {
  if (!e.hadRecentInput) { window.__authoritiesSmoke.cls += e.value;
    window.__authoritiesSmoke.shifts.push({value:e.value,time:e.startTime}); }
})).observe({type:'layout-shift',buffered:true}); } catch {}
try { new PerformanceObserver(list => list.getEntries().forEach(e =>
  window.__authoritiesSmoke.longTasks.push({duration:e.duration,time:e.startTime})
)).observe({type:'longtask',buffered:true}); } catch {}
try { Object.defineProperty(window,'showOpenFilePicker',{configurable:true,value:undefined}); } catch {}
try { Object.defineProperty(window,'showDirectoryPicker',{configurable:true,value:async()=>{let scans=0;return ({
  requestPermission:async()=> 'granted',
  async *values() {
    const item=window.__authoritiesSmoke.fakeCanlii;
    if (!item || scans++ === 0) return;
    yield {kind:'file',name:item.name,getFile:async()=>new File([
      Uint8Array.from(atob(item.base64),character=>character.charCodeAt(0))
    ],item.name,{type:'application/pdf',lastModified:item.lastModified||1700000000000})};
  }
})}}); } catch {}
const nativeOpen=window.open;
window.open=(url,target,features)=>url==='about:blank'
  ? {opener:null,location:{replace:value=>{
      window.__authoritiesSmoke.openedUrls.push(value);
      if (String(value).includes('canlii.org')) window.__authoritiesSmoke.canliiClicks.push(value);
    }},close(){}}
  : nativeOpen.call(window,url,target,features);
addEventListener('click',event=>{
  const link=event.target.closest?.('a[href*="canlii.org"]');
  if (link) {
    window.__authoritiesSmoke.canliiLinkEvents.push({href:link.href,target:link.target,
      appPrevented:event.defaultPrevented});
    event.preventDefault();
  }
});
const nativeFetch=window.fetch;
const sourceOrigin=authority=>authority.source.kind==='attached'
  ? authority.source.sources.find(({origin})=>origin)?.origin||null : null;
window.fetch=async(...args)=>{
  const body=args[1]?.body,target=typeof args[0]==='string' ? args[0] : args[0]?.url||'';
  const method=(args[1]?.method||args[0]?.method||'GET').toUpperCase(),absolute=new URL(target,location.href).href;
  const isImport=(method==='POST'&&/\/api\/authorities$/.test(new URL(absolute).pathname))||
    absolute.includes('/authorities-runtime/import');
  const timing={url:absolute,method,startAt:performance.now(),responseAt:null,status:null};
  window.__authoritiesSmoke.requestTimings.push(timing);
  if (isImport) window.__authoritiesSmoke.importUi=timing;
  const isBuild=target.includes('/authorities-runtime/build')||/\/authorities\/[^/]+\/build$/.test(target);
  if (isBuild) {
    window.__authoritiesSmoke.buildStarts++;
    if (window.__authoritiesSmoke.gateBuild) await new Promise(resolve=>{
      window.__authoritiesSmoke.releaseBuild=()=>{window.__authoritiesSmoke.gateBuild=false;resolve();};
    });
  }
  if (target.includes('/authorities-runtime/build') && body instanceof FormData) {
    const draft=JSON.parse(body.get('draft'));
    window.__authoritiesSmoke.lastBuild={importKind:draft.import.kind,outputMode:draft.outputMode,
      insertIntoDocument:draft.insertIntoDocument,settings:draft.settings,
      sources:draft.authorityOrder.map(id=>draft.authorities[id].source.kind),
      sourceOrigins:draft.authorityOrder.map(id=>sourceOrigin(draft.authorities[id])),
      sourceProof:draft.authorityOrder.map(id=>({citation:draft.authorities[id].citation,
        origin:sourceOrigin(draft.authorities[id])})),
      bookParts:draft.bookParts,roles:JSON.parse(body.get('roles')),
      files:body.getAll('files').map(file=>({name:file.name,size:file.size,type:file.type}))};
  }
  const response=await nativeFetch(...args);
  timing.responseAt=performance.now();timing.status=response.status;
  if (response.ok && /\/api\/authorities\/[^/]+\/sources$/.test(response.url)) {
    const product=await response.clone().json(),draft=product.state;
    window.__authoritiesSmoke.lastPrepared={settings:draft.settings,
      sources:draft.authorityOrder.map(id=>draft.authorities[id].source.kind),
      sourceOrigins:draft.authorityOrder.map(id=>sourceOrigin(draft.authorities[id])),
      sourceProof:draft.authorityOrder.map(id=>({citation:draft.authorities[id].citation,
        origin:sourceOrigin(draft.authorities[id])}))};
  }
  if (response.ok && (/\/api\/authorities\/[^/]+\/book-parts\//.test(response.url)||
      response.url.includes('/authorities-runtime/book-part'))) {
    const value=await response.clone().json(),state=value.state||value;
    window.__authoritiesSmoke.lastBookPart=state.bookParts;
  }
  if (!response.ok) window.__authoritiesSmoke.errors.push({url:response.url,
    status:response.status,body:await response.clone().text()});
  if (response.url.includes('/authorities-runtime/build') ||
      /\/authorities\/[^/]+\/build$/.test(response.url)) window.__authoritiesSmoke.builds++;
  return response;
};
new MutationObserver(()=>{
  const timing=window.__authoritiesSmoke.importUi;
  if (timing?.responseAt&&!timing.readyAt&&
      document.querySelector('[role="listbox"][aria-label="Citations"] [role="option"]'))
    timing.readyAt=performance.now();
}).observe(document,{subtree:true,childList:true});
const nativeClick=HTMLInputElement.prototype.click;
HTMLInputElement.prototype.click=function() {
  if (this.type==='file' && !this.isConnected) {
    this.dataset.authoritiesSmokePicker='true';
    Object.assign(this.style,{position:'fixed',left:'-9999px',top:'0'});
    document.documentElement.append(this);
    this.addEventListener('change',()=>setTimeout(()=>this.remove()),{once:true});
    return;
  }
  return nativeClick.call(this);
};
"""})
    return driver


def wait(driver: webdriver.Chrome, seconds: int = 60) -> WebDriverWait:
    return WebDriverWait(driver, seconds, ignored_exceptions=(StaleElementReferenceException,))


def idle(driver: webdriver.Chrome) -> None:
    wait(driver, 120).until(lambda item: not item.find_elements(By.CSS_SELECTOR, "[class*='animate-spin']"))


def tab(driver: webdriver.Chrome, name: str):
    return driver.find_element(By.XPATH,
        f"//*[@role='tablist' and @aria-label='Authorities sections']//*[@role='tab' and normalize-space(.)='{name}']")


def click_button(driver: webdriver.Chrome, name: str, root=None):
    owner = root or driver
    return owner.find_element(By.XPATH, f".//button[normalize-space(.)='{name}']").click()


def choose(driver: webdriver.Chrome, label: str, option: str, root=None) -> None:
    owner = root or driver
    field = owner.find_element(By.XPATH,
        f".//label[.//select and starts-with(normalize-space(.), '{label}')]/select")
    Select(field).select_by_visible_text(option)


def upload(driver: webdriver.Chrome, label: str, paths: list[Path], root=None) -> None:
    owner = root or driver
    inputs = owner.find_elements(By.XPATH,
        f".//label[normalize-space(.)='{label}']/input[@type='file']")
    if inputs:
        picker = inputs[0]
    else:
        click_button(driver, label, owner)
        picker = wait(driver, 5).until(lambda item: item.find_element(
            By.CSS_SELECTOR, "input[type='file'][data-authorities-smoke-picker]"))
    picker.send_keys("\n".join(str(path.resolve()) for path in paths))


def upload_aria(driver: webdriver.Chrome, name: str, paths: list[Path]) -> None:
    controls = driver.find_elements(By.CSS_SELECTOR,
        f"input[type='file'][aria-label={json.dumps(name)}]")
    if controls:
        picker = controls[0]
    else:
        driver.find_element(By.CSS_SELECTOR, f"button[aria-label={json.dumps(name)}]").click()
        picker = wait(driver, 5).until(lambda item: item.find_element(
            By.CSS_SELECTOR, "input[type='file'][data-authorities-smoke-picker]"))
    picker.send_keys("\n".join(str(path.resolve()) for path in paths))


def authority_rows(driver: webdriver.Chrome):
    return driver.find_elements(By.XPATH,
        "//section[.//h2[normalize-space()='Authorities']]//article")


def set_authorities_court(driver: webdriver.Chrome, jurisdiction: str, court: str) -> None:
    current = driver.find_element(By.CSS_SELECTOR, "button[aria-label^='Court:']")
    if current.get_attribute("aria-label") == f"Court: {court}":
        return
    current.click()
    dialog = wait(driver, 5).until(lambda item: item.find_element(By.CSS_SELECTOR, "dialog[open]"))
    dialog.find_element(By.XPATH,
        f".//button[normalize-space()={json.dumps(jurisdiction)}]").click()
    if jurisdiction != court:
        match = "contains(normalize-space(),'King')" if "King" in court \
            else f"normalize-space()={json.dumps(court)}"
        wait(driver, 5).until(lambda item: item.find_element(By.XPATH,
            f"//dialog[@open]//button[{match}]")).click()
    expected = "King" if "King" in court else f"Court: {court}"
    wait(driver, 30).until(lambda item: expected in item.find_element(
        By.CSS_SELECTOR, "button[aria-label^='Court:']").get_attribute("aria-label"))
    idle(driver)


def fixture_proof(source: Path, pdfs: list[Path], filing: Path) -> dict[str, object]:
    text = "\n".join(paragraph.text for paragraph in Document(source).paragraphs)
    assert all(citation in text for citation in REAL_CITATIONS)
    with zipfile.ZipFile(source) as package:
        notes = package.read("word/footnotes.xml").decode()
        assert "Ibid at para 31" in notes and OAKES in notes and package.testzip() is None
    files = [*pdfs, filing]
    for path in files:
        with fitz.open(path) as pdf:
            assert pdf.is_pdf and not pdf.is_encrypted and pdf.page_count == 1 and pdf[0].get_text().strip()
    return {"source": {"file": source.name, "sha256": digest(source)},
            "pdfs": [{"file": path.name, "sha256": digest(path)} for path in files]}


def citation_options(driver: webdriver.Chrome):
    return driver.find_elements(By.CSS_SELECTOR, "[role='listbox'][aria-label='Citations'] [role='option']")


def draft_state(driver: webdriver.Chrome, draft_id: str) -> dict[str, object]:
    if urlparse(driver.current_url).path.rstrip("/").endswith("authorities.html"):
        result = driver.execute_async_script(r"""
const id=arguments[0],done=arguments[arguments.length-1],opening=
  indexedDB.open('beaver-work-products');
opening.onerror=()=>done({error:String(opening.error)});
opening.onsuccess=()=>{
  const request=opening.result.transaction('drafts').objectStore('drafts').get(id);
  request.onerror=()=>done({error:String(request.error)});
  request.onsuccess=()=>done(request.result?.state||null);
};
""", draft_id)
        assert result, {"draft": draft_id, "result": result}
        return result
    result = driver.execute_async_script(r"""
const id=arguments[0],done=arguments[arguments.length-1];
fetch(`/api/authorities/${encodeURIComponent(id)}`).then(async response=>done({
  status:response.status,value:response.ok ? await response.json() : await response.text()
})).catch(error=>done({error:String(error)}));
""", draft_id)
    assert result.get("status") == 200, result
    return result["value"]["state"]


def assert_rejected_false_positive(driver: webdriver.Chrome, draft_id: str) -> dict[str, int]:
    state = draft_state(driver, draft_id)
    occurrences = list(state["occurrences"].values())
    authorities = list(state["authorities"].values())
    assert not any(FALSE_POSITIVE in f'{item.get("text", "")} {item.get("citation", "")}'
                   for item in occurrences), occurrences
    assert not any(FALSE_POSITIVE in f'{item.get("citation", "")} {item.get("name", "")}'
                   for item in authorities), authorities
    return {"occurrences": len(occurrences), "authorities": len(authorities)}


def occurrence(driver: webdriver.Chrome, needle: str, location: str = ""):
    matches = [item for item in citation_options(driver)
               if needle.lower() in item.text.lower() and (not location or location in item.text)]
    assert matches, {"needle": needle, "locations": [item.text for item in citation_options(driver)]}
    return matches[0]


def review_surface(driver: webdriver.Chrome):
    return driver.find_element(By.CSS_SELECTOR,
        "[role='textbox'][aria-label='In-text citation context'],[role='textbox'][aria-label='Footnote context']")


def citation_import_timing(driver: webdriver.Chrome) -> dict[str, object]:
    value = wait(driver, 10).until(lambda item: item.execute_script(r"""
const timing=window.__authoritiesSmoke.importUi;if(!timing?.readyAt)return null;
const entry=performance.getEntriesByType('resource').filter(({name})=>name===timing.url).at(-1);
if(!entry?.responseEnd)return null;
return {url:timing.url,status:timing.status,
  fetchToHeadersMs:timing.responseAt-timing.startAt,
  requestToFirstByteMs:entry.responseStart-entry.requestStart,
  responseBodyMs:entry.responseEnd-entry.responseStart,
  requestToBodyEndMs:entry.responseEnd-entry.requestStart,
  uiAfterBodyEndMs:timing.readyAt-entry.responseEnd,
  endToEndMs:timing.readyAt-timing.startAt,
  transferSize:entry.transferSize,encodedBodySize:entry.encodedBodySize};
"""))
    assert 200 <= value["status"] < 300 and ("/api/authorities" in value["url"] or
        "/authorities-runtime/import" in value["url"]), value
    assert 0 <= value["uiAfterBodyEndMs"] <= CITATION_RENDER_BUDGET_MS, value
    for key in ("fetchToHeadersMs", "requestToFirstByteMs", "responseBodyMs",
                "requestToBodyEndMs", "uiAfterBodyEndMs", "endToEndMs"):
        value[key] = round(value[key], 3)
    value["gate"] = {"metric": "uiAfterBodyEndMs", "budgetMs": CITATION_RENDER_BUDGET_MS}
    value["reportOnly"] = ["fetchToHeadersMs", "requestToFirstByteMs", "responseBodyMs",
                           "requestToBodyEndMs", "endToEndMs"]
    driver.execute_script("sessionStorage.setItem('authoritiesImportTiming',JSON.stringify(arguments[0]))",
                          value)
    return value


def native_citation_performance() -> dict[str, object]:
    proof = Path(__file__).with_name("test-authorities-citation-performance.mjs")
    completed = subprocess.run(["node", str(proof)], cwd=proof.parent.parent,
                               capture_output=True, text=True, check=False)
    assert completed.returncode == 0, completed.stdout + completed.stderr
    return json.loads(completed.stdout)


def selection_offsets(driver: webdriver.Chrome, surface) -> dict[str, object] | None:
    return driver.execute_script(r"""
const root=arguments[0],selection=getSelection(); if(!selection?.rangeCount)return null;
const range=selection.getRangeAt(0); if(!root.contains(range.startContainer)||!root.contains(range.endContainer))return null;
const before=document.createRange();before.selectNodeContents(root);before.setEnd(range.startContainer,range.startOffset);
const through=document.createRange();through.selectNodeContents(root);through.setEnd(range.endContainer,range.endOffset);
return {start:before.toString().length,end:through.toString().length,text:range.toString(),collapsed:range.collapsed};
""", surface)


def pointer_select_text(driver: webdriver.Chrome, surface, needle: str,
                        backward: bool = False) -> dict[str, object]:
    driver.execute_script("arguments[0].scrollIntoView({block:'center',inline:'nearest'})", surface)
    geometry = driver.execute_script(r"""
const root=arguments[0],needle=arguments[1],walker=document.createTreeWalker(root,NodeFilter.SHOW_TEXT);
let nodes=[],text='',node;while(node=walker.nextNode()){nodes.push([node,text.length]);text+=node.data;}
const start=text.indexOf(needle),end=start+needle.length;if(start<0)return null;
const point=(index,right)=>{const hit=nodes.findLast(([,offset])=>offset<=index),local=index-hit[1],r=document.createRange();
  r.setStart(hit[0],local);r.setEnd(hit[0],local+1);const box=r.getClientRects()[0];
  return {x:right?box.right-1:box.left+1,y:box.top+box.height/2};};
const hits=[...root.querySelectorAll('mark')].filter(mark=>{const r=document.createRange();
  r.selectNodeContents(root);r.setEndBefore(mark);const s=r.toString().length;
  return s<end&&s+mark.textContent.length>start;}),box=root.getBoundingClientRect(),marked=hits[0]?.getBoundingClientRect();
return {start,end,a:point(start,false),b:point(end-1,true),marks:hits.length,
  visible:box.top>=0&&box.bottom<=innerHeight,
  centerDelta:marked?Math.abs((marked.top+marked.bottom-box.top*2-box.height)/2):null};
""", surface, needle)
    assert geometry and geometry["marks"] > 0 and geometry["visible"], {
        "needle": needle, "text": surface.text, "geometry": geometry}
    assert geometry["centerDelta"] is not None and geometry["centerDelta"] <= 24, geometry
    a, b = geometry["a"], geometry["b"]
    if backward:
        a, b = b, a
    driver.execute_cdp_cmd("Input.dispatchMouseEvent", {"type": "mouseMoved", **a})
    driver.execute_cdp_cmd("Input.dispatchMouseEvent", {"type": "mousePressed", "button": "left",
        "buttons": 1, "clickCount": 1, **a})
    for step in range(1, 13):
        driver.execute_cdp_cmd("Input.dispatchMouseEvent", {"type": "mouseMoved", "button": "left",
            "buttons": 1, "x": a["x"] + (b["x"] - a["x"]) * step / 12,
            "y": a["y"] + (b["y"] - a["y"]) * step / 12})
    driver.execute_cdp_cmd("Input.dispatchMouseEvent", {"type": "mouseReleased", "button": "left",
        "buttons": 0, "clickCount": 1, **b})
    actual = wait(driver, 5).until(lambda _item: selection_offsets(driver, review_surface(driver)))
    assert actual == {"start": geometry["start"], "end": geometry["end"],
                      "text": needle, "collapsed": False}, {"expected": geometry, "actual": actual}
    return {"start": actual["start"], "end": actual["end"],
            "direction": "backward" if backward else "forward", "marksCrossed": geometry["marks"]}


def set_cursor(driver: webdriver.Chrome, surface, needle: str, within: int) -> int:
    cursor = driver.execute_script(r"""
const root=arguments[0],target=arguments[1],within=arguments[2],walker=document.createTreeWalker(root,NodeFilter.SHOW_TEXT);
let nodes=[],text='',node;while(node=walker.nextNode()){nodes.push([node,text.length]);text+=node.data;}
const cursor=text.indexOf(target)+within,hit=nodes.findLast(([,offset])=>offset<=cursor);if(cursor<within||!hit)return -1;
const range=document.createRange();range.setStart(hit[0],cursor-hit[1]);range.collapse(true);
const selection=getSelection();selection.removeAllRanges();selection.addRange(range);
root.dispatchEvent(new KeyboardEvent('keyup',{bubbles:true,key:'ArrowRight'}));return cursor;
""", surface, needle, within)
    assert cursor >= 0
    return cursor


def authority_bounds(driver: webdriver.Chrome, surface, needle: str,
                     kind: str = "authority") -> dict[str, object]:
    result = driver.execute_script(r"""
const root=arguments[0],needle=arguments[1],marks=[...root.querySelectorAll(`[data-${arguments[2]}-span]`)]
  .filter(mark=>mark.textContent===needle);if(marks.length!==1)return {count:marks.length,text:root.textContent};
const before=document.createRange();before.selectNodeContents(root);before.setEndBefore(marks[0]);
return {count:1,start:before.toString().length,end:before.toString().length+marks[0].textContent.length,
  expectedStart:root.textContent.indexOf(needle),text:marks[0].textContent};
""", surface, needle, kind)
    assert result == {"count": 1, "start": result.get("expectedStart"),
        "end": result.get("expectedStart", -1) + len(needle.encode("utf-16-le")) // 2,
        "expectedStart": result.get("expectedStart"), "text": needle}, result
    return result


def header_rect(driver: webdriver.Chrome) -> dict[str, float]:
    return driver.execute_script(r"""
const header=document.querySelector('.authorities-workspace>[data-workspace-header]'),title=header?.querySelector('h1');
const h=header?.getBoundingClientRect(),t=title?.getBoundingClientRect();
return {x:h.x,y:h.y,width:h.width,height:h.height,titleX:t.x,titleY:t.y,titleHeight:t.height};
""")


def workspace_rect(driver: webdriver.Chrome) -> dict[str, object]:
    metrics = driver.execute_script(r"""
const root=document.querySelector('.authorities-workspace'),r=root.getBoundingClientRect(),
  bottom=document.elementFromPoint(r.left+2,innerHeight-2);
return {top:r.top,bottom:r.bottom,height:r.height,viewport:innerHeight,
  rootBackground:getComputedStyle(root).backgroundColor,
  bottomBackground:getComputedStyle(bottom).backgroundColor};
""")
    assert metrics["bottom"] >= metrics["viewport"] - 1, metrics
    assert metrics["rootBackground"] == metrics["bottomBackground"], metrics
    return metrics


def stable_headers(headers: dict[str, dict[str, float]]) -> None:
    values = list(headers.values()); base = values[0]
    for name, item in headers.items():
        for key in ("x", "y", "width", "height", "titleY", "titleHeight"):
            assert abs(item[key] - base[key]) <= 1, {"header": name, "field": key, "all": headers}


def header_cycle(driver: webdriver.Chrome, static: dict[str, float]) -> dict[str, dict[str, float]]:
    headers = {"static": static, "draft": header_rect(driver)}
    for name in ("Drafts", "Settings", "Automatic"):
        tab(driver, name).click()
        wait(driver, 5).until(lambda _item, name=name:
            tab(driver, name).get_attribute("aria-selected") == "true")
        headers[name.lower()] = header_rect(driver)
    stable_headers(headers)
    return headers


def assistant_header(driver: webdriver.Chrome, beaver: bool, output: Path) -> dict[str, object]:
    controls = driver.find_elements(By.CSS_SELECTOR,
        ".authorities-workspace>[data-workspace-header] button[aria-label='Assistant']")
    if not beaver:
        assert not controls
        return {"available": False}
    assert len(controls) == 1
    draft_id = parse_qs(urlparse(driver.current_url).query)["draft"][0]
    title = driver.find_element(By.CSS_SELECTOR,
        ".authorities-workspace>[data-workspace-header] h1").text
    product = driver.execute_async_script(r"""
const id=arguments[0],done=arguments[arguments.length-1];
fetch(`/api/work-products/${encodeURIComponent(id)}`).then(async response=>
  done({status:response.status,value:await response.json()})).catch(error=>done({error:String(error)}));
""", draft_id)
    assert product["status"] == 200 and product["value"]["id"] == draft_id and \
        product["value"]["title"] == title, product
    before = header_rect(driver); controls[0].click()
    def dock_is_visible(item: webdriver.Chrome) -> bool:
        try:
            return item.find_element(
                By.CSS_SELECTOR, "aside[aria-label='Assistant dock']").is_displayed()
        except StaleElementReferenceException:
            return False
    wait(driver, 10).until(dock_is_visible)
    dock = driver.find_element(By.CSS_SELECTOR,
        "aside[data-assistant-dock][aria-label='Assistant dock'][aria-hidden='false']")
    assert dock.find_element(By.CSS_SELECTOR,
        f"[aria-label={json.dumps(f'Current draft: {title}')}]").is_displayed()
    composer = wait(driver, 30).until(lambda _item: dock.find_element(
        By.CSS_SELECTOR, "textarea[placeholder='How can I help?']"))
    assert composer.is_displayed() and composer.is_enabled() and composer.accessible_name == "Message"
    composer.send_keys("/help")
    send = dock.find_element(By.CSS_SELECTOR, "button[aria-label='Send message']")
    wait(driver, 5).until(lambda _item: send.is_enabled()); send.click()
    wait(driver, 5).until(lambda _item: "Commands" in dock.text and "/compact" in dock.text)
    assert driver.save_screenshot(str(output / "assistant-bound-desktop.png"))
    opened = header_rect(driver)
    for key in ("height", "titleY", "titleHeight"):
        assert abs(opened[key] - before[key]) <= 1, {"before": before, "opened": opened}
    collapse = dock.find_element(By.CSS_SELECTOR,
        "button[aria-label='Collapse assistant dock'],button[aria-label='Close assistant']")
    collapse.click()
    wait(driver, 5).until(lambda _item: dock.get_attribute("aria-hidden") == "true")
    control = driver.find_element(By.CSS_SELECTOR,
        ".authorities-workspace>[data-workspace-header] button[aria-label='Assistant']")
    assert control.get_attribute("aria-expanded") == "false"; control.click()
    wait(driver, 5).until(dock_is_visible)
    reopened = driver.find_element(By.CSS_SELECTOR,
        "aside[data-assistant-dock][aria-label='Assistant dock'][aria-hidden='false']")
    reopened_composer = reopened.find_element(By.CSS_SELECTOR,
        "textarea[placeholder='How can I help?']")
    assert driver.execute_script("return arguments[0]===arguments[1]", dock, reopened)
    assert driver.execute_script("return arguments[0]===arguments[1]", composer, reopened_composer)
    assert "Commands" in reopened.text and "/compact" in reopened.text
    driver.execute_cdp_cmd("Emulation.setDeviceMetricsOverride",
        {"width": 320, "height": 800, "deviceScaleFactor": 1, "mobile": False})
    wait(driver, 5).until(lambda _item: reopened.get_attribute("role") == "dialog")
    compact = driver.execute_script(r"""
const dock=arguments[0],composer=dock.querySelector("textarea"),r=dock.getBoundingClientRect(),
  c=composer.getBoundingClientRect(),root=document.documentElement;
return {left:r.left,right:r.right,top:r.top,bottom:r.bottom,height:r.height,
  viewportWidth:innerWidth,viewportHeight:innerHeight,documentOverflow:root.scrollWidth-root.clientWidth,
  dockOverflow:dock.scrollWidth-dock.clientWidth,composer:{left:c.left,right:c.right,top:c.top,bottom:c.bottom}};
""", reopened)
    assert (compact["left"] >= 0 and compact["right"] <= 321 and compact["top"] >= 0 and
        compact["bottom"] <= 801 and compact["height"] <= compact["viewportHeight"] * .76 + 1 and
        compact["documentOverflow"] <= 1 and compact["dockOverflow"] <= 1), compact
    assert (compact["composer"]["left"] >= compact["left"] and
        compact["composer"]["right"] <= compact["right"] and
        compact["composer"]["top"] >= compact["top"] and
        compact["composer"]["bottom"] <= compact["bottom"]), compact
    assert driver.save_screenshot(str(output / "assistant-bound-320.png"))
    driver.execute_cdp_cmd("Emulation.clearDeviceMetricsOverride", {})
    wait(driver, 5).until(lambda _item: driver.execute_script("return innerWidth") > 1000)
    reopened.find_element(By.CSS_SELECTOR,
        "button[aria-label='Collapse assistant dock'],button[aria-label='Close assistant']").click()
    return {"available": True, "draft": draft_id, "before": before, "opened": opened,
            "sameMountedConversation": True, "compact": compact, "collapsedOnlyByUser": True}


def correct_parallel_citation(driver: webdriver.Chrome) -> dict[str, object]:
    before = len(citation_options(driver))
    occurrence(driver, "1986 CanLII 46").click()
    surface = review_surface(driver)
    selected = pointer_select_text(driver, surface, OAKES, backward=True)
    control = driver.find_element(By.XPATH, "//button[normalize-space(.)='Use selection as citation']")
    wait(driver, 5).until(lambda _item: control.is_enabled()); control.click(); idle(driver)
    wait(driver, 30).until(lambda _item: len(citation_options(driver)) == before - 1)
    surface = review_surface(driver)
    bounds = wait(driver, 5).until(lambda _item: authority_bounds(driver, surface, OAKES))
    surrogate_delta = bounds["start"] - surface.text.index(OAKES)
    assert surrogate_delta == 1, {"bounds": bounds, "text": surface.text}
    actions = driver.find_elements(By.XPATH,
        "//button[normalize-space(.)='Use selection as citation' or normalize-space(.)='Use selection as pinpoint']")
    assert len(actions) == 2 and all(not item.is_enabled() for item in actions), [item.is_enabled() for item in actions]
    draft_url = driver.current_url
    driver.refresh()
    wait(driver, 120).until(lambda item: item.find_elements(
        By.CSS_SELECTOR, "[role='listbox'][aria-label='Citations']"))
    occurrence(driver, "1986 CanLII 46", "In-text").click()
    persisted = authority_bounds(driver, review_surface(driver), OAKES)
    assert parse_qs(urlparse(driver.current_url).query).get("draft") == \
        parse_qs(urlparse(draft_url).query).get("draft")
    return {"pointer": selected, "utf16SurrogateDelta": surrogate_delta,
            "occurrencesBefore": before, "occurrencesAfter": before - 1,
            "stored": bounds, "persistedAfterReload": persisted, "actionsReset": True}


def correct_pinpoint(driver: webdriver.Chrome) -> dict[str, object]:
    occurrence(driver, "2009 SCC 32", "In-text").click()
    surface = review_surface(driver)
    selected = pointer_select_text(driver, surface, "para 29")
    control = driver.find_element(By.XPATH, "//button[normalize-space(.)='Use selection as pinpoint']")
    wait(driver, 5).until(lambda _item: control.is_enabled()); control.click(); idle(driver)
    wait(driver, 30).until(lambda _item: any(mark.text == "para 29" for mark in
        review_surface(driver).find_elements(By.CSS_SELECTOR, "[data-pinpoint-span]")))
    stored = authority_bounds(driver, review_surface(driver), "para 29", "pinpoint")
    assert all(not button.is_enabled() for button in driver.find_elements(By.XPATH,
        "//button[normalize-space(.)='Use selection as citation' or normalize-space(.)='Use selection as pinpoint']"))
    driver.refresh()
    wait(driver, 120).until(lambda item: item.find_elements(
        By.CSS_SELECTOR, "[role='listbox'][aria-label='Citations']"))
    occurrence(driver, "2009 SCC 32", "In-text").click()
    wait(driver, 30).until(lambda _item: any(mark.text == "para 29" for mark in
        review_surface(driver).find_elements(By.CSS_SELECTOR, "[data-pinpoint-span]")))
    persisted = authority_bounds(driver, review_surface(driver), "para 29", "pinpoint")
    return {"pointer": selected, "stored": stored, "persistedAfterReload": persisted,
            "actionsReset": True}


def reject_false_positive(driver: webdriver.Chrome, source: Path) -> dict[str, object]:
    draft_id = parse_qs(urlparse(driver.current_url).query)["draft"][0]
    before = {key: len(value) for key, value in draft_state(driver, draft_id).items()
              if key in {"occurrences", "authorities"}}
    occurrence(driver, FALSE_POSITIVE).click()
    click_button(driver, "Not a citation"); idle(driver)
    wait(driver, 30).until(lambda _item: not any(
        FALSE_POSITIVE in item.text for item in [*citation_options(driver), *authority_rows(driver)]))
    after = assert_rejected_false_positive(driver, draft_id)
    assert after == {key: value - 1 for key, value in before.items()}, {"before": before, "after": after}

    driver.refresh()
    wait(driver, 120).until(lambda item: parse_qs(urlparse(item.current_url).query).get("draft") ==
        [draft_id] and bool(citation_options(item)))
    assert not any(FALSE_POSITIVE in item.text
                   for item in [*citation_options(driver), *authority_rows(driver)])
    assert_rejected_false_positive(driver, draft_id)

    requests = len(driver.execute_script("return window.__authoritiesSmoke.requestTimings"))
    upload(driver, "Replace file", [source])
    wait(driver, 120).until(lambda item: item.execute_script(r"""
return window.__authoritiesSmoke.requestTimings.slice(arguments[0]).some(request=>
  request.status>=200&&request.status<300&&request.method==='POST'&&
  (/\/authorities\/[^/]+\/source$/.test(new URL(request.url).pathname)||
   request.url.includes('/authorities-runtime/refresh')));
""", requests))
    idle(driver)
    assert not any(FALSE_POSITIVE in item.text
                   for item in [*citation_options(driver), *authority_rows(driver)])
    assert_rejected_false_positive(driver, draft_id)

    tab(driver, "Drafts").click()
    saved = wait(driver).until(lambda item: item.find_element(By.XPATH,
        "//button[.//span[normalize-space(.)='real-authorities-smoke']]"))
    saved.click()
    wait(driver, 120).until(lambda item: parse_qs(urlparse(item.current_url).query).get("draft") ==
        [draft_id] and tab(item, "Automatic").get_attribute("aria-selected") == "true" and
        bool(citation_options(item)))
    assert not any(FALSE_POSITIVE in item.text
                   for item in [*citation_options(driver), *authority_rows(driver)])
    assert_rejected_false_positive(driver, draft_id)
    return {"citation": FALSE_POSITIVE, "before": before, "after": after,
            "survivedRefresh": True, "survivedRescan": True, "survivedReopen": True,
            "orphanAuthorityRemoved": True}


def consolidate_parallel(driver: webdriver.Chrome, location: str = "") -> dict[str, object]:
    before = len(citation_options(driver))
    occurrence(driver, "1986 CanLII 46", location).click()
    selected = pointer_select_text(driver, review_surface(driver), OAKES)
    control = driver.find_element(By.XPATH, "//button[normalize-space(.)='Use selection as citation']")
    wait(driver, 5).until(lambda _item: control.is_enabled()); control.click(); idle(driver)
    wait(driver, 30).until(lambda _item: len(citation_options(driver)) == before - 1)
    bounds = authority_bounds(driver, review_surface(driver), OAKES)
    return {"pointer": selected, "occurrencesBefore": before,
            "occurrencesAfter": before - 1, "stored": bounds}


def citation_keyboard_navigation(driver: webdriver.Chrome) -> dict[str, object]:
    items = citation_options(driver); outer = driver.execute_script("return scrollY")
    items[0].click(); driver.find_element(By.CSS_SELECTOR,
        "[role='listbox'][aria-label='Citations'] [role='option'][aria-selected='true']").send_keys(Keys.END)
    last = wait(driver, 5).until(lambda item: item.switch_to.active_element
        if item.switch_to.active_element == citation_options(item)[-1] else False)
    visible = driver.execute_script(r"""
const item=arguments[0],list=item.parentElement,a=item.getBoundingClientRect(),b=list.getBoundingClientRect();
return a.top>=b.top-1&&a.bottom<=b.bottom+1;
""", last)
    assert visible and last.get_attribute("aria-selected") == "true"
    last.send_keys(Keys.HOME)
    wait(driver, 5).until(lambda item: item.switch_to.active_element == citation_options(item)[0])
    assert driver.execute_script("return scrollY") == outer
    return {"items": len(items), "endVisible": True, "homeAndEnd": True,
            "outerScrollStable": True}


def split_merge_later_occurrence(driver: webdriver.Chrome) -> dict[str, object]:
    target = occurrence(driver, "2016 SCC 27", "Footnote")
    target.click(); options_before = citation_options(driver)
    selected_before = options_before.index(driver.find_element(By.CSS_SELECTOR,
        "[role='listbox'][aria-label='Citations'] [role='option'][aria-selected='true']"))
    cursor = set_cursor(driver, review_surface(driver), "2016 SCC 27", 5)
    split = driver.find_element(By.XPATH, "//button[normalize-space(.)='Split at cursor']")
    wait(driver, 5).until(lambda _item: split.is_enabled()); split.click(); idle(driver)
    wait(driver, 30).until(lambda _item: len(citation_options(driver)) == len(options_before) + 1)
    selected = driver.find_element(By.CSS_SELECTOR,
        "[role='listbox'][aria-label='Citations'] [role='option'][aria-selected='true']")
    split_index = citation_options(driver).index(selected)
    assert split_index == selected_before + 1 and "SCC 27" in selected.text, {
        "before": selected_before, "after": split_index, "selected": selected.text}
    merge = driver.find_element(By.XPATH, "//button[normalize-space(.)='Merge with previous']")
    wait(driver, 5).until(lambda _item: merge.is_enabled()); merge.click(); idle(driver)
    wait(driver, 30).until(lambda _item: len(citation_options(driver)) == len(options_before))
    merged = driver.find_element(By.CSS_SELECTOR,
        "[role='listbox'][aria-label='Citations'] [role='option'][aria-selected='true']")
    merged_index = citation_options(driver).index(merged)
    assert merged_index == selected_before and "2016 SCC 27" in merged.text, {
        "before": selected_before, "after": merged_index, "selected": merged.text}
    return {"cursor": cursor, "index": selected_before, "splitIndex": split_index,
            "mergedIndex": merged_index, "occurrences": len(options_before)}


def reference_keyboard_flow(driver: webdriver.Chrome) -> dict[str, object]:
    source = occurrence(driver, "Ibid", "Footnote"); source.click()
    source_index = citation_options(driver).index(source)
    link = driver.find_element(By.XPATH, "//button[normalize-space(.)='Link to authority']")
    link.click()
    focused = wait(driver, 5).until(lambda item: item.switch_to.active_element
        if item.switch_to.active_element.get_attribute("role") == "option" else False)
    focused.send_keys(Keys.ESCAPE)
    wait(driver, 5).until(lambda item: not item.find_elements(By.XPATH,
        "//*[contains(normalize-space(.), 'Choose the full citation this cross-reference points to.')]") )
    assert driver.switch_to.active_element.get_attribute("aria-selected") == "true"
    driver.find_element(By.XPATH, "//button[normalize-space(.)='Link to authority']").click()
    focused = wait(driver, 5).until(lambda item: item.switch_to.active_element
        if item.switch_to.active_element.get_attribute("role") == "option" else False)
    focused.send_keys(Keys.HOME)
    driver.switch_to.active_element.send_keys(Keys.ENTER)
    wait(driver, 30).until(lambda _item: driver.execute_script("""
return [...document.querySelectorAll('span')].some(node=>node.innerText.startsWith('Linked to '));
"""))
    selected = driver.find_element(By.CSS_SELECTOR,
        "[role='listbox'][aria-label='Citations'] [role='option'][aria-selected='true']")
    assert citation_options(driver).index(selected) == source_index
    assert driver.switch_to.active_element == selected
    click_button(driver, "Clear link"); idle(driver)
    wait(driver, 5).until(lambda _item: driver.execute_script("""
return [...document.querySelectorAll('span')].some(node=>node.innerText==='Not linked');
"""))
    assert driver.find_element(By.XPATH, "//button[normalize-space(.)='Link to authority']").is_displayed()
    return {"sourceIndex": source_index, "escapeRestoredFocus": True,
            "linkedByKeyboard": True, "sourceFocusRestored": True, "cleared": True}


def centering_proof(driver: webdriver.Chrome) -> dict[str, object]:
    outer = driver.execute_script("return scrollY")
    occurrence(driver, "2016 SCC 27", "In-text").click()
    surface = review_surface(driver)
    metrics = wait(driver, 5).until(lambda _item: driver.execute_script(r"""
const root=arguments[0],mark=root.querySelector('[data-authority-span]');if(!mark)return null;
const r=root.getBoundingClientRect(),m=mark.getBoundingClientRect(),line=parseFloat(getComputedStyle(root).lineHeight);
return {scrollTop:root.scrollTop,centerDelta:Math.abs((m.top+m.height/2)-(r.top+r.height/2)),
  lineHeight:line,outerScroll:scrollY,rootTop:r.top,markTop:m.top};
""", surface))
    assert metrics["scrollTop"] > 0 and metrics["centerDelta"] <= metrics["lineHeight"], metrics
    assert metrics["outerScroll"] == outer, {"before": outer, "after": metrics}
    return metrics


def active_review_viewports(driver: webdriver.Chrome, output: Path) -> dict[str, object]:
    proof = {}
    for name, scale in (("320", 1), ("320-200-percent", 2)):
        driver.execute_cdp_cmd("Emulation.setDeviceMetricsOverride", {
            "width": 320, "height": 900, "deviceScaleFactor": scale, "mobile": False})
        occurrence(driver, "2009 SCC 32", "In-text").click()
        occurrence(driver, "2016 SCC 27", "In-text").click()
        surface = review_surface(driver)
        wait(driver, 5).until(lambda _item: driver.execute_script(r"""
const root=arguments[0],mark=root.querySelector('[data-authority-span]');if(!mark)return false;
const a=root.getBoundingClientRect(),b=mark.getBoundingClientRect();
return Math.abs((b.top+b.bottom-a.top-a.bottom)/2)<=parseFloat(getComputedStyle(root).lineHeight);
""", surface))
        metrics = driver.execute_script(r"""
const surface=arguments[0],review=surface.closest('.authorities-review'),list=review.querySelector('[role=listbox]'),
  mark=surface.querySelector('[data-authority-span]');surface.scrollIntoView({block:'center'});
const rect=n=>{const r=n.getBoundingClientRect();return {left:r.left,right:r.right,top:r.top,bottom:r.bottom}},
  editor=rect(surface),selected=rect(mark),lineHeight=parseFloat(getComputedStyle(surface).lineHeight);
return {width:innerWidth,scale:devicePixelRatio,overflow:document.documentElement.scrollWidth-innerWidth,
  review:rect(review),list:rect(list),surface:editor,selectedSpan:selected,lineHeight,
  centerDelta:Math.abs((selected.top+selected.bottom-editor.top*2-editor.bottom+editor.top)/2),
  selected:document.querySelectorAll('[aria-label="Citations"] [aria-selected=true]').length};
""", surface)
        assert metrics["width"] == 320 and metrics["scale"] == scale and metrics["overflow"] <= 1, metrics
        for key in ("review", "list", "surface", "selectedSpan"):
            assert metrics[key]["left"] >= -1 and metrics[key]["right"] <= 321, metrics
        assert metrics["centerDelta"] <= metrics["lineHeight"] and metrics["selected"] == 1, metrics
        assert driver.save_screenshot(str(output / f"automatic-review-{name}.png"))
        proof[name] = metrics
    driver.execute_cdp_cmd("Emulation.clearDeviceMetricsOverride", {})
    driver.set_window_size(1440, 1000)
    return proof


def automatic_review(driver: webdriver.Chrome, source: Path, pdfs: list[Path], filing: Path,
                     output: Path, static_header: dict[str, float], beaver: bool) -> dict[str, object]:
    upload(driver, "Add file", [source])
    setup = wait(driver, 5).until(lambda item: item.find_element(
        By.CSS_SELECTOR, "dialog[open]"))
    click_button(driver, "Import and review", setup)
    def imported(item: webdriver.Chrome):
        errors = item.execute_script("return window.__authoritiesSmoke.errors")
        assert not errors, errors
        return next(iter(item.find_elements(
            By.CSS_SELECTOR, "[role='listbox'][aria-label='Citations']")), False)
    citations = wait(driver, 120).until(imported)
    import_timing = citation_import_timing(driver)
    options = citations.find_elements(By.CSS_SELECTOR, "[role='option']")
    labels = [option.text for option in options]
    assert all(any(citation in label for label in labels) for citation in REAL_CITATIONS), labels
    next(option for option in options if REAL_CITATIONS[0] in option.text).click()
    surface = wait(driver).until(lambda _item: review_surface(driver))
    assert surface.get_attribute("aria-label") in {"In-text citation context", "Footnote context"}
    kind_labels = [span.text for span in driver.find_elements(By.TAG_NAME, "span")
                   if span.text.startswith(("In-text citation ", "Footnote "))]
    assert kind_labels and any(label.startswith("Footnote 1") for label in labels), \
        driver.find_element(By.TAG_NAME, "body").text
    correction = correct_parallel_citation(driver)
    false_positive = reject_false_positive(driver, source)
    headers = header_cycle(driver, static_header)
    assistant = assistant_header(driver, beaver, output)
    pinpoint = correct_pinpoint(driver)
    footnote_parallel = consolidate_parallel(driver, "Footnote")
    keyboard = citation_keyboard_navigation(driver)
    split_merge = split_merge_later_occurrence(driver)
    reference = reference_keyboard_flow(driver)
    centering = centering_proof(driver)
    responsive = active_review_viewports(driver, output)
    choose(driver, "Create", "Book and Table")
    driver.find_element(By.XPATH, "//summary[normalize-space(.)='Options']").click()
    missing = driver.find_element(By.XPATH,
        "//label[.//select and starts-with(normalize-space(.), 'Missing sources')]/select")
    assert missing.get_attribute("value") == "placeholder"
    occurrence(driver, "2016 SCC 27", "Footnote").click()
    busy = build(driver, prove_busy=True)
    built_sources = driver.execute_script(
        "return window.__authoritiesSmoke.lastBuild||window.__authoritiesSmoke.lastPrepared")
    assert built_sources["settings"]["sourceMode"] == "automatic", built_sources
    assert built_sources["settings"]["missingSourcePolicy"] == "placeholder", built_sources
    assert "attached" in built_sources["sources"], built_sources
    proof = {item["citation"]: item["origin"] for item in built_sources["sourceProof"]}
    assert proof.get("2026 SCC 16") == "original" and \
        proof.get(RECONSTRUCTED) == "reconstructed", proof
    files = downloads(driver, output / "automatic")
    table_file = next(path for path in files if path.suffix.lower() == ".docx")
    book_file = next(path for path in files if path.suffix.lower() == ".pdf")
    assert driver.save_screenshot(str(output / "automatic-desktop.png"))
    book = inspect_book(book_file, output / "automatic-first-page.png")
    with fitz.open(book_file) as built:
        toc = built.get_toc()
        entry = next(index for index, item in enumerate(toc) if "2009 SCC 32" in item[1])
        level, _, first = toc[entry]
        last = next((page for depth, _, page in toc[entry + 1:] if depth <= level), len(built) + 1)
        marked = None
        for page in (built[index] for index in range(first - 1, last - 1)):
            bars = [drawing["rect"] for drawing in page.get_drawings()
                    if 2 < drawing["rect"].height < 100 and
                    (color := drawing.get("color") or drawing.get("fill")) and
                    color[0] > .5 and color[1] < .35 and color[2] < .35]
            if any(bar.y0 <= target.y1 and bar.y1 >= target.y0
                   for target in page.search_for("[29]") for bar in bars):
                marked = page
                break
        assert marked is not None, "R. v. Grant para 29 did not receive a bounded margin mark."
        marked.get_pixmap(matrix=fitz.Matrix(1.3, 1.3), alpha=False).save(
            output / "automatic-marked-paragraph.png")
    sources = driver.find_element(By.XPATH, "//summary[.//h2[normalize-space(.)='Sources']]/parent::details")
    if not sources.get_attribute("open"):
        sources.find_element(By.TAG_NAME, "summary").click()
    oakes = next(row for row in sources.find_elements(By.TAG_NAME, "article") if "1986 CanLII 46" in row.text)
    upload(driver, "Add file", [pdfs[1]], oakes); idle(driver)
    wait(driver, 30).until(lambda _item: any(pdfs[1].name in row.text
        for row in sources.find_elements(By.TAG_NAME, "article") if "1986 CanLII 46" in row.text))
    profiles = profile_builds(driver, filing, output)
    return {"labels": kind_labels, "citations": labels, "importTiming": import_timing,
            "table": inspect_table(table_file),
            "placeholderBook": book, "markedParagraph": 29,
            "sourceOrigins": built_sources["sourceOrigins"], "correction": correction,
            "sourceProof": proof, "falsePositive": false_positive,
            "pinpointCorrection": pinpoint, "footnoteParallel": footnote_parallel,
            "citationKeyboard": keyboard, "splitMerge": split_merge,
            "referenceLink": reference, "centering": centering,
            "activeReviewViewports": responsive, "busyBuild": busy, "profiles": profiles,
            "headers": headers, "assistantHeader": assistant,
            "manualPdfFallback": pdfs[1].name}


def build(driver: webdriver.Chrome, prove_busy: bool = False) -> dict[str, object]:
    control = driver.find_element(By.XPATH, "//button[normalize-space(.)='Build']")
    wait(driver, 30).until(lambda _item: control.is_enabled())
    completed = driver.execute_script("return window.__authoritiesSmoke.builds")
    started = driver.execute_script("return window.__authoritiesSmoke.buildStarts")
    failures = len(driver.execute_script("return window.__authoritiesSmoke.errors"))
    if prove_busy:
        merge = driver.find_element(By.XPATH, "//button[normalize-space(.)='Merge with previous']")
        assert merge.is_enabled(), "Busy proof requires an ordinarily enabled mutation."
        driver.execute_script("window.__authoritiesSmoke.gateBuild=true")
    control.click()
    busy = {}
    if prove_busy:
        wait(driver, 120).until(lambda item: item.execute_script(
            "return window.__authoritiesSmoke.releaseBuild!==null"))
        busy = {"starts": driver.execute_script("return window.__authoritiesSmoke.buildStarts") - started,
            "cancelShown": bool(driver.find_elements(By.XPATH, "//button[normalize-space(.)='Cancel']")),
            "settingsDisabled": all(not field.is_enabled() for field in driver.find_elements(
                By.XPATH, "//h2[normalize-space(.)='Build outputs']/following::select")),
            "reviewActionsDisabled": all(not button.is_enabled() for button in driver.find_elements(
                By.XPATH, "//button[normalize-space(.)='Use selection as citation' or normalize-space(.)='Use selection as pinpoint' or normalize-space(.)='Split at cursor' or normalize-space(.)='Merge with previous']"))}
        assert busy == {"starts": 1, "cancelShown": True, "settingsDisabled": True,
                        "reviewActionsDisabled": True}, busy
        driver.execute_script("window.__authoritiesSmoke.releaseBuild();window.__authoritiesSmoke.releaseBuild=null")
    wait(driver, 120).until(lambda _item: driver.execute_script(
        "return window.__authoritiesSmoke.builds") > completed)
    idle(driver)
    errors = driver.execute_script("return window.__authoritiesSmoke.errors")
    assert len(errors) == failures, {"errors": errors[failures:],
        "request": driver.execute_script("return window.__authoritiesSmoke.lastBuild")}
    assert driver.find_elements(By.CSS_SELECTOR, "button[aria-label^='Download ']"), [
        [status.text for status in driver.find_elements(By.CSS_SELECTOR, "[role='status']") if status.text],
        driver.execute_script("return window.__authoritiesSmoke.errors")]
    return busy


def downloads(driver: webdriver.Chrome, directory: Path) -> list[Path]:
    directory.mkdir(parents=True, exist_ok=True)
    driver.execute_cdp_cmd("Page.setDownloadBehavior",
                           {"behavior": "allow", "downloadPath": str(directory.resolve())})
    result = []
    for control in driver.find_elements(By.CSS_SELECTOR, "button[aria-label^='Download ']"):
        filename = control.get_attribute("aria-label").removeprefix("Download ")
        target = directory / filename
        control.click()
        wait(driver, 60).until(lambda _item: target.is_file() and target.stat().st_size > 0
            and not list(directory.glob("*.crdownload")))
        result.append(target)
    assert result, "No built files were offered for download."
    return result


def inspect_artifacts(paths: list[Path], preview: Path) -> list[dict[str, object]]:
    result = []
    for path in paths:
        if path.suffix.lower() == ".pdf":
            with fitz.open(path) as pdf:
                text = "\n".join(page.get_text() for page in pdf)
                links = sum(len(page.get_links()) for page in pdf)
                toc = pdf.get_toc()
                assert (pdf.is_pdf and not pdf.is_encrypted and pdf.page_count > 0 and
                        len(text) > 100 and links > 0 and toc)
                pdf[0].get_pixmap(matrix=fitz.Matrix(1.2, 1.2), alpha=False).save(
                    preview.with_name(f"{preview.stem}-{path.stem}.png"))
                result.append({"file": path.name, "kind": "pdf", "pages": pdf.page_count,
                    "textCharacters": len(text), "firstPage": pdf[0].get_text().strip()[:500],
                    "links": links, "bookmarks": len(toc), "sha256": digest(path)})
        else:
            with zipfile.ZipFile(path) as package:
                assert package.testzip() is None and "word/document.xml" in package.namelist()
            document = Document(path)
            text = "\n".join([paragraph.text for paragraph in document.paragraphs] +
                [cell.text for table in document.tables for row in table.rows for cell in row.cells])
            assert len(text) > 50 and sum(citation in text for citation in REAL_CITATIONS) >= 2, text
            result.append({"file": path.name, "kind": "docx", "tables": len(document.tables),
                           "textCharacters": len(text), "sha256": digest(path)})
    return result


def cover_details(driver: webdriver.Chrome, court_file: str,
                  roles: tuple[tuple[str, str], tuple[str, str]]) -> None:
    driver.find_element(By.XPATH,
        "//h3[normalize-space(.)='Cover and index']/following-sibling::div[1]"
        "//button[normalize-space(.)='Add details' or normalize-space(.)='Edit']").click()
    dialog = wait(driver, 5).until(lambda item: item.find_element(
        By.XPATH, "//dialog[@open and .//*[normalize-space(.)='Cover details']]"))
    number = dialog.find_element(By.XPATH,
        ".//label[starts-with(normalize-space(.), 'Court file number')]/input")
    number.clear(); number.send_keys(court_file)
    groups = dialog.find_elements(By.CSS_SELECTOR, "section[aria-label^='Party group']")
    assert len(groups) == 2
    for group, (role, names) in zip(groups, roles, strict=True):
        role_input = group.find_element(By.XPATH, ".//label[normalize-space(.)='Role']/input")
        role_input.clear(); role_input.send_keys(role)
        parties = group.find_element(By.TAG_NAME, "textarea")
        parties.clear(); parties.send_keys(names)
    dialog.find_element(By.XPATH, ".//button[normalize-space(.)='Save cover']").click()
    wait(driver, 10).until(lambda _item: not driver.find_elements(By.XPATH,
        "//dialog[@open and .//*[normalize-space(.)='Cover details']]"))
    idle(driver)


def profile_builds(driver: webdriver.Chrome, filing: Path, output: Path) -> dict[str, object]:
    proof = {}
    profiles = (
        ("abkb", "Alberta", "Court of King’s Bench of Alberta", "Book and Table", None, None),
        ("abca", "Alberta", "Court of Appeal of Alberta", None, None, None),
        ("fc", "Federal courts", "Federal Court", None, "Paper", "Respondent"),
        ("fca", "Federal courts", "Federal Court of Appeal", None, "Electronic", "Intervener"),
    )
    for slug, jurisdiction, court, create, medium, role in profiles:
        correction = None
        set_authorities_court(driver, jurisdiction, court)
        if slug == "abca":
            upload(driver, "Replace file", [filing]); idle(driver)
            wait(driver, 120).until(lambda item: filing.name in item.find_element(
                By.XPATH, "//h2[normalize-space(.)='Import and review']/parent::div").text)
            correction = consolidate_parallel(driver)
        if create:
            choose(driver, "Create", create); idle(driver)
        if medium:
            choose(driver, "Filing", medium); choose(driver, "Filed by", role); idle(driver)
            cover_details(driver, "T-123-26" if slug == "fc" else "A-123-26",
                (("Applicant" if slug == "fc" else "Appellant",
                  "North Prairie Ltd.\nRiver Holdings Inc."),
                 ("Respondent", "Attorney General of Canada")))
        build(driver)
        request = driver.execute_script(
            "return window.__authoritiesSmoke.lastBuild||window.__authoritiesSmoke.lastPrepared")
        assert request["settings"]["profileId"] == {
            "abkb": "ab-court-of-kings-bench", "abca": "ab-court-of-appeal",
            "fc": "federal-court", "fca": "federal-court-of-appeal"}[slug], request
        if medium:
            assert (request["settings"].get("filingMedium"), request["settings"].get("bookRole")) == \
                (medium.lower(), role.lower()), request
        files = downloads(driver, output / "automatic-profiles" / slug)
        assert len(files) == (2 if slug in {"abkb", "abca"} else 1), {
            "profile": slug, "files": [path.name for path in files]}
        kinds = {path.suffix.lower() for path in files}
        assert (slug != "abca" or {".docx", ".pdf"}.issubset(kinds)) and \
            (slug == "abca" or ".pdf" in kinds), {
            "profile": slug, "files": [path.name for path in files]}
        artifacts = inspect_artifacts(files,
            output / "automatic-profiles" / slug / "first-page.png")
        if role:
            assert any(f"Filed by {role}" in str(item.get("firstPage", ""))
                       for item in artifacts), artifacts
        proof[slug] = {"request": request, "artifacts": artifacts,
                       "parallelCorrection": correction}
    return proof


def add_manual_authority(driver: webdriver.Chrome, citation: str, name: str = ""):
    before = len(authority_rows(driver))
    click_button(driver, "Add authority")
    dialog = wait(driver, 5).until(lambda item: item.find_element(By.CSS_SELECTOR, "dialog[open]"))
    dialog.find_element(By.XPATH,
        ".//label[starts-with(normalize-space(.), 'Citation')]/input").send_keys(citation)
    if name:
        dialog.find_element(By.XPATH,
            ".//label[starts-with(normalize-space(.), 'Displayed title')]/input").send_keys(name)
    click_button(driver, "Add", dialog)
    wait(driver, 120).until(lambda _item: len(authority_rows(driver)) == before + 1)
    idle(driver)
    return next(row for row in authority_rows(driver) if (name or citation) in row.text)


def procedural_tabs(driver: webdriver.Chrome) -> list[str]:
    return [row.find_element(By.CSS_SELECTOR, "[aria-label^='Tab ']").get_attribute("aria-label")
            for row in authority_rows(driver)]


def open_pdf(driver: webdriver.Chrome, root) -> dict[str, object]:
    opened = len(driver.execute_script("return window.__authoritiesSmoke.openedUrls"))
    requests = len(driver.execute_script("return window.__authoritiesSmoke.requestTimings"))
    click_button(driver, "Open", root)
    url = wait(driver, 30).until(lambda item: (lambda values: values[-1]
        if len(values) > opened else False)(item.execute_script(
            "return window.__authoritiesSmoke.openedUrls")))
    assert url.startswith("blob:"), url
    fresh = driver.execute_script(
        "return window.__authoritiesSmoke.requestTimings.slice(arguments[0]).map(x=>x.url)", requests)
    return {"scheme": "blob", "requests": fresh}


def cover_row(driver: webdriver.Chrome):
    return driver.find_element(By.XPATH,
        "//h3[normalize-space(.)='Cover and index']/following-sibling::div[1]/div[.//span[normalize-space(.)='Cover']]")


def supplement_row(driver: webdriver.Chrome, filename: str):
    return driver.find_element(By.XPATH,
        f"//h3[normalize-space(.)='Other book PDFs']/parent::div/following-sibling::div[1]"
        f"//span[@title={json.dumps(filename)}]/parent::div")


def manual_book(driver: webdriver.Chrome, pdfs: list[Path], output: Path) -> dict[str, object]:
    if parse_qs(urlparse(driver.current_url).query).get("draft"):
        click_button(driver, "New")
        wait(driver, 5).until(lambda item: "draft" not in parse_qs(urlparse(
            item.current_url).query))
    tab(driver, "Manual").click()
    title = wait(driver, 5).until(lambda item: item.find_element(
        By.XPATH, "//label[starts-with(normalize-space(.), 'Book title')]/input"))
    title.clear()
    title.send_keys("Authorities Smoke Book")
    upload(driver, "Add files", pdfs[:2])
    wait(driver, 120).until(lambda _item: len(authority_rows(driver)) == 2)
    idle(driver)
    draft_id = parse_qs(urlparse(driver.current_url).query)["draft"][0]
    assert procedural_tabs(driver) == ["Tab 1", "Tab 2"]

    first = authority_rows(driver)[0]
    first_title = first.find_element(By.TAG_NAME, "h3").text
    upload(driver, "Replace", [pdfs[3]], first)
    wait(driver, 120).until(lambda _item: pdfs[3].name in
        next(row for row in authority_rows(driver) if first_title in row.text).text)
    assert procedural_tabs(driver) == ["Tab 1", "Tab 2"]
    opened_authority = open_pdf(driver,
        next(row for row in authority_rows(driver) if first_title in row.text))

    jordan = add_manual_authority(driver, "R v Jordan, 2016 SCC 27")
    handoffs = wait(driver, 120).until(lambda _item:
        (next(row for row in authority_rows(driver) if "2016 SCC 27" in row.text)
         .find_elements(By.LINK_TEXT, "Download from CanLII")) or False)
    assert len(handoffs) == 1 and urlparse(handoffs[0].get_attribute("href")).hostname == "www.canlii.org"
    assert handoffs[0].get_attribute("href").endswith(".pdf")
    handoff = handoffs[0].get_attribute("href")
    page_url, handles = driver.current_url, len(driver.window_handles)
    handoffs[0].click()
    assert driver.current_url == page_url and len(driver.window_handles) == handles
    assert driver.execute_script("return window.__authoritiesSmoke.canliiLinkEvents") == [{
        "href": handoff, "target": "_blank", "appPrevented": False}]
    assert jordan.find_element(By.XPATH,
        ".//*[self::button or self::label][normalize-space(.)='Add file']").is_displayed()
    click_button(driver, "Connect downloads folder")
    wait(driver, 5).until(lambda item: item.find_element(By.XPATH,
        "//button[normalize-space(.)='Change downloads folder']").get_attribute("aria-pressed") == "true")
    expected = Path(urlparse(handoff).path).name
    driver.execute_script("window.__authoritiesSmoke.fakeCanlii=arguments[0]", {
        "name": expected, "base64": base64.b64encode(pdfs[2].read_bytes()).decode("ascii")})
    jordan = next(row for row in authority_rows(driver) if "2016 SCC 27" in row.text)
    jordan.find_element(By.LINK_TEXT, "Download from CanLII").click()
    wait(driver, 120).until(lambda _item: any(expected in row.text for row in
        authority_rows(driver) if "2016 SCC 27" in row.text))
    idle(driver)
    assert driver.current_url == page_url and len(driver.window_handles) == handles
    assert driver.execute_script("return window.__authoritiesSmoke.canliiClicks") == [handoff]
    assert driver.execute_script("return window.__authoritiesSmoke.canliiLinkEvents") == [
        {"href": handoff, "target": "_blank", "appPrevented": False},
        {"href": handoff, "target": "_blank", "appPrevented": True}]
    assert not driver.find_elements(By.CSS_SELECTOR, "a[href*='canlii.org']")

    unknown = add_manual_authority(driver, "Unreported decision", "Unreported decision")
    assert not unknown.find_elements(By.LINK_TEXT, "Download from CanLII")
    upload(driver, "Add file", [pdfs[3]], unknown)
    wait(driver, 120).until(lambda _item: pdfs[3].name in
        next(row for row in authority_rows(driver) if "Unreported decision" in row.text).text)
    assert procedural_tabs(driver) == ["Tab 1", "Tab 2", "Tab 3", "Tab 4"]

    upload(driver, "Add file", [pdfs[6]], cover_row(driver)); idle(driver)
    wait(driver, 30).until(lambda _item: pdfs[6].name in cover_row(driver).text)
    opened_cover = open_pdf(driver, cover_row(driver))

    upload_aria(driver, "Add other book files", [pdfs[4]])
    wait(driver, 120).until(lambda _item: pdfs[4].name in driver.find_element(
        By.XPATH, "//h3[normalize-space(.)='Other book PDFs']/ancestor::div[contains(@class,'border-t')][1]").text)
    idle(driver)
    original_row = supplement_row(driver, pdfs[4].name)
    original_tab = next(span.text for span in original_row.find_elements(By.TAG_NAME, "span")
                        if span.text.upper().startswith("TAB "))
    original_parts = driver.execute_script("return window.__authoritiesSmoke.lastBookPart")
    original_part = next(part for part in original_parts["supplements"]
                         if part["filename"] == pdfs[4].name)
    upload(driver, "Replace", [pdfs[5]], original_row)
    wait(driver, 120).until(lambda _item: pdfs[5].name in supplement_row(driver, pdfs[5].name).text)
    idle(driver)
    replaced_row = supplement_row(driver, pdfs[5].name)
    replaced_tab = next(span.text for span in replaced_row.find_elements(By.TAG_NAME, "span")
                        if span.text.upper().startswith("TAB "))
    replaced_parts = driver.execute_script("return window.__authoritiesSmoke.lastBookPart")
    replaced_part = next(part for part in replaced_parts["supplements"]
                         if part["filename"] == pdfs[5].name)
    assert (replaced_tab, replaced_part["id"], replaced_part["bindingRole"]) == \
        (original_tab, original_part["id"], original_part["bindingRole"]), {
            "before": original_part, "after": replaced_part,
            "tabs": [original_tab, replaced_tab]}
    opened_supplement = open_pdf(driver, replaced_row)

    build(driver)
    first_file = next(path for path in downloads(driver, output / "manual-first")
                      if path.suffix.lower() == ".pdf")
    first_proof = inspect_book(first_file, output / "manual-first-page.png",
        required_text=("Authorities Smoke Book", "Table of Contents"),
        required_outline=("Authorities Smoke Book", "Table of Contents"))
    tab(driver, "Automatic").click()
    tab(driver, "Drafts").click()
    saved = wait(driver).until(lambda item: [button for button in item.find_elements(By.TAG_NAME, "button")
        if "Authorities Smoke Book" in button.text])
    saved[0].click()
    wait(driver).until(lambda item: parse_qs(urlparse(item.current_url).query).get("draft") == [draft_id])
    assert procedural_tabs(driver) == ["Tab 1", "Tab 2", "Tab 3", "Tab 4"]
    assert pdfs[6].name in cover_row(driver).text
    assert pdfs[5].name in supplement_row(driver, pdfs[5].name).text

    row = authority_rows(driver)[0]
    row.find_element(By.CSS_SELECTOR, "button[aria-haspopup='menu']").click()
    wait(driver, 5).until(lambda item: item.find_element(By.XPATH,
        "//*[@role='menuitem' and normalize-space(.)='Edit details']")).click()
    dialog = wait(driver, 5).until(lambda item: item.find_element(By.CSS_SELECTOR, "dialog[open]"))
    choose(driver, "Type", "Case", dialog)
    citation = dialog.find_element(By.XPATH,
        ".//label[starts-with(normalize-space(.), 'Citation')]/input")
    citation.clear()
    citation.send_keys("2009 SCC 32")
    name = dialog.find_element(By.XPATH,
        ".//label[starts-with(normalize-space(.), 'Displayed title')]/input")
    name.clear()
    name.send_keys("R v Grant")
    assert driver.save_screenshot(str(output / "manual-identity-editor.png"))
    click_button(driver, "Save", dialog)
    wait(driver).until(lambda _item: "R v Grant" in row.text and "2009 SCC 32" in row.text)
    assert procedural_tabs(driver) == ["Tab 1", "Tab 2", "Tab 3", "Tab 4"]
    idle(driver)
    build(driver)
    rebuilt = next(path for path in downloads(driver, output / "manual-rebuilt")
                   if path.suffix.lower() == ".pdf")
    rebuilt_proof = inspect_book(rebuilt, output / "manual-rebuilt-page.png",
        required_text=("Authorities Smoke Book", "Table of Contents"),
        required_outline=("Authorities Smoke Book", "Table of Contents"))
    assert digest(first_file) != digest(rebuilt)
    assert "R v Grant, 2009 SCC 32" in "\n".join(rebuilt_proof["outline"])
    responsive = manual_viewport_proof(driver, output)
    assert driver.save_screenshot(str(output / "manual-desktop.png"))
    return {"draft": draft_id, "handoff": {"url": handoff, "target": "_blank",
                "openableWithoutFolder": True},
            "canliiCapture": {"filename": expected, "clicks": 1, "requests": 0},
            "fixedTabs": procedural_tabs(driver), "replacement": pdfs[3].name,
            "openPdf": {"authority": opened_authority, "cover": opened_cover,
                "supplement": opened_supplement},
            "supplementReplacement": {"tab": replaced_tab, "id": replaced_part["id"],
                "bindingRole": replaced_part["bindingRole"], "filename": replaced_part["filename"]},
            "ordinaryPdfUpload": True,
            "correctedIdentity": {"filename": "2009scc32.pdf", "kind": "case",
                "name": "R v Grant", "citation": "2009 SCC 32"}, "first": first_proof,
            "rebuilt": rebuilt_proof, "viewports": responsive}


def inspect_table(path: Path) -> dict[str, object]:
    with zipfile.ZipFile(path) as package:
        assert package.testzip() is None
        assert {"[Content_Types].xml", "word/document.xml"}.issubset(package.namelist())
    document = Document(path)
    text = "\n".join([paragraph.text for paragraph in document.paragraphs] +
        [cell.text for table in document.tables for row in table.rows for cell in row.cells])
    assert document.tables and "Table of Authorities" in text
    assert all(citation in text for citation in REAL_CITATIONS), text
    return {"file": path.name, "sha256": digest(path), "tables": len(document.tables)}


def link_page(link: dict[str, object]) -> int:
    try:
        return int(link.get("page", -1))
    except (TypeError, ValueError):
        return -1


def inspect_book(path: Path, preview: Path, *,
                 required_text: tuple[str, ...] = ("Book of Authorities", "Table of Contents"),
                 required_outline: tuple[str, ...] = ("Book of Authorities", "Table of Contents"),
                 minimum_links: int = 3) -> dict[str, object]:
    pdf = fitz.open(path)
    try:
        text = "\n".join(page.get_text() for page in pdf)
        toc = pdf.get_toc()
        outline = [item[1] for item in toc]
        links = [link for page in pdf for link in page.get_links() if link_page(link) >= 0]
        assert pdf.is_pdf and not pdf.is_encrypted and pdf.page_count >= 5
        assert all(value in text for value in required_text), text
        assert len(links) >= minimum_links and set(required_outline).issubset(set(outline))
        pdf[0].get_pixmap(matrix=fitz.Matrix(1.2, 1.2), alpha=False).save(preview)
        return {"file": path.name, "sha256": digest(path), "pages": pdf.page_count,
                "bookmarks": len(toc), "outline": outline, "links": len(links),
                "textCharacters": len(text),
                "textSha256": hashlib.sha256(text.encode("utf-8")).hexdigest()}
    finally:
        pdf.close()


def manual_viewport_proof(driver: webdriver.Chrome, output: Path) -> dict[str, object]:
    proof = {}
    for name, width, height, scale in (("desktop", 1440, 1000, 1), ("320", 320, 900, 1),
                                        ("320-200-percent", 320, 900, 2)):
        driver.execute_cdp_cmd("Emulation.setDeviceMetricsOverride", {
            "width": width, "height": height, "deviceScaleFactor": scale, "mobile": False})
        root = driver.find_element(By.XPATH,
            "//h2[normalize-space(.)='Authorities']/ancestor::section[1]")
        driver.execute_script("arguments[0].scrollIntoView({block:'start'})", root)
        tabs = procedural_tabs(driver)
        assert tabs == ["Tab 1", "Tab 2", "Tab 3", "Tab 4"], tabs
        metrics = driver.execute_script("""
const root=arguments[0],box=root.getBoundingClientRect();
return {width:innerWidth,scale:devicePixelRatio,physicalWidth:innerWidth*devicePixelRatio,
  overflow:document.documentElement.scrollWidth-innerWidth,left:box.left,right:box.right,
  wideNodes:[...document.querySelectorAll('*')].flatMap(node=>{const style=getComputedStyle(node);
    if(!node.getClientRects().length||node.scrollWidth<=node.clientWidth+1||
      ['hidden','clip'].includes(style.overflowX))return[];
    return [{tag:node.tagName,role:node.getAttribute('role'),aria:node.getAttribute('aria-label'),
      width:node.clientWidth,scrollWidth:node.scrollWidth,className:String(node.className).slice(0,120)}];
  }).slice(0,12)};
""", root)
        assert metrics["width"] == width and metrics["scale"] == scale, metrics
        assert metrics["physicalWidth"] == width * scale and metrics["overflow"] <= 1, metrics
        assert metrics["left"] >= -1 and metrics["right"] <= width + 1, metrics
        assert not metrics["wideNodes"], metrics
        assert driver.save_screenshot(str(output / f"manual-{name}.png"))
        metrics["tabs"] = tabs
        proof[name] = metrics
    driver.execute_cdp_cmd("Emulation.clearDeviceMetricsOverride", {})
    driver.set_window_size(1440, 1000)
    return proof


def keyboard_tabs(driver: webdriver.Chrome) -> None:
    automatic = tab(driver, "Automatic")
    automatic.click()
    automatic.send_keys(Keys.ARROW_RIGHT)
    wait(driver, 5).until(lambda _item: tab(driver, "Manual").get_attribute("aria-selected") == "true")
    driver.switch_to.active_element.send_keys(Keys.ARROW_LEFT)
    wait(driver, 5).until(lambda _item: automatic.get_attribute("aria-selected") == "true")


def visit_tabs(driver: webdriver.Chrome) -> list[str]:
    names = ["Automatic", "Manual", "Drafts", "Settings"]
    for name in names:
        control = tab(driver, name)
        control.click()
        panel = control.get_attribute("aria-controls")
        wait(driver, 5).until(lambda item, control=control, panel=panel:
            control.get_attribute("aria-selected") == "true" and
            item.find_element(By.ID, panel).is_displayed())
    tab(driver, "Automatic").click()
    wait(driver, 5).until(lambda _item:
        tab(driver, "Automatic").get_attribute("aria-selected") == "true")
    return names


def viewport_proof(driver: webdriver.Chrome, output: Path) -> dict[str, object]:
    proof = {}
    for name, width, height, scale in (("desktop", 1440, 1000, 1), ("320", 320, 800, 1),
                                        ("200-percent", 320, 800, 2)):
        driver.execute_script("window.__authoritiesSmoke.cls=0;window.__authoritiesSmoke.shifts=[]")
        driver.execute_cdp_cmd("Emulation.setDeviceMetricsOverride", {
            "width": width, "height": height, "deviceScaleFactor": scale, "mobile": False})
        driver.execute_script("scrollTo(0,0)")
        visited = visit_tabs(driver)
        metrics = driver.execute_script("""
const s=window.__authoritiesSmoke||{};
return {width:innerWidth,visualWidth:visualViewport?.width||innerWidth,
  deviceScale:devicePixelRatio,physicalWidth:innerWidth*devicePixelRatio,
  overflow:document.documentElement.scrollWidth-innerWidth,
  cls:s.cls||0,longTasks:(s.longTasks||[]).length,
  selectedTabs:document.querySelectorAll('[aria-label="Authorities sections"] [role=tab][aria-selected=true]').length,
  visibleTabs:[...document.querySelectorAll('[aria-label="Authorities sections"] [role=tab]')]
    .every(n=>{const r=n.getBoundingClientRect();return r.left>=-1&&r.right<=innerWidth+1}),
  unnamed:[...document.querySelectorAll('button,a[href],input,select,[role=textbox]')]
    .filter(n=>n.getClientRects().length && !(n.getAttribute('aria-label')||n.labels?.length||n.textContent.trim()))
    .map(n=>n.outerHTML.slice(0,120))};
""")
        metrics["visitedTabs"] = visited
        assert abs(metrics["width"] - width) <= 1 and abs(metrics["visualWidth"] - width) <= 1, metrics
        assert metrics["deviceScale"] == scale and metrics["physicalWidth"] == width * scale, metrics
        assert metrics["overflow"] <= 1 and metrics["cls"] <= 0.1, metrics
        assert metrics["selectedTabs"] == 1 and metrics["visibleTabs"] and \
            visited == ["Automatic", "Manual", "Drafts", "Settings"], metrics
        assert not metrics["unnamed"], metrics
        assert driver.save_screenshot(str(output / f"authorities-{name}.png"))
        for section, heading in (("Drafts", "Saved drafts"), ("Settings", "New drafts")):
            tab(driver, section).click()
            wait(driver, 5).until(lambda item, heading=heading: item.find_elements(
                By.XPATH, f"//h2[normalize-space(.)='{heading}']"))
            assert driver.save_screenshot(str(output / f"authorities-{section.lower()}-{name}.png"))
        tab(driver, "Automatic").click()
        proof[name] = metrics
    driver.execute_cdp_cmd("Emulation.clearDeviceMetricsOverride", {})
    driver.set_window_size(1440, 1000)
    return proof


def network_urls(driver: webdriver.Chrome) -> list[str]:
    result = []
    for record in driver.get_log("performance"):
        try:
            message = json.loads(record["message"])["message"]
            if message["method"] == "Network.requestWillBeSent":
                result.append(message["params"]["request"]["url"])
        except (KeyError, TypeError, json.JSONDecodeError):
            pass
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description="Live Chrome smoke for the shared Authorities workspace.")
    parser.add_argument("--url", default="http://127.0.0.1:3000/table-of-authorities")
    parser.add_argument("--artifacts", type=Path,
                        default=Path(__file__).resolve().parents[1] / ".tmp-live-qa" / "authorities-browser-proof")
    parser.add_argument("--headed", action="store_true")
    parser.add_argument("--built-asset", type=Path,
                        help="Exact Authorities page bundle served for this proof.")
    parser.add_argument("--server-pid", type=int)
    parser.add_argument("--server-command")
    parser.add_argument("--fixtures-only", action="store_true",
                        help="Validate the local DOCX/PDF fixtures without launching Chrome.")
    parser.add_argument("--manual-only", action="store_true",
                        help="Run the manual-book journey without importing a filing document.")
    args = parser.parse_args()
    native_performance = native_citation_performance()
    standalone = urlparse(args.url).path.rstrip("/").endswith("authorities.html")
    mode = "standalone" if standalone else "beaver"
    with tempfile.TemporaryDirectory(prefix=f"authorities-{mode}-") as temporary:
        temporary_path = Path(temporary)
        run = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%fZ")
        output = args.artifacts.resolve() / mode / run
        output.mkdir(parents=True, exist_ok=True)
        fixture_dir = output / "inputs"
        fixture_dir.mkdir(parents=True)
        source, pdfs, filing = fixtures(fixture_dir)
        inputs = fixture_proof(source, pdfs, filing)
        if args.fixtures_only:
            print(json.dumps({"inputs": inputs, "native_citation_parse": native_performance}, indent=2))
            return 0
        driver = chrome(temporary_path / "chrome", args.headed)
        driver.set_window_size(1440, 1000)
        started = time.monotonic()
        try:
            driver.get(args.url)
            wait(driver, 30).until(lambda item: item.find_elements(By.XPATH,
                "//h1[normalize-space(.)='Authorities']"))
            navigation = driver.execute_script("""
const n=performance.getEntriesByType('navigation')[0]; return n&&{
 duration:n.duration,responseStart:n.responseStart,domContentLoaded:n.domContentLoadedEventEnd,
 load:n.loadEventEnd,transferSize:n.transferSize};
""")
            assert navigation and navigation["duration"] > 0
            served_build = None
            if args.built_asset:
                loaded = driver.execute_async_script(r"""
const name=arguments[0],done=arguments[arguments.length-1];
const entry=performance.getEntriesByType('resource').find(({name:url})=>
  new URL(url).pathname.endsWith(`/assets/${name}`));
if(!entry){done(null);return;}
fetch(entry.name,{cache:'no-store'}).then(async response=>{
  const bytes=await response.arrayBuffer(),hash=await crypto.subtle.digest('SHA-256',bytes);
  done({url:entry.name,sha256:[...new Uint8Array(hash)]
    .map(value=>value.toString(16).padStart(2,'0')).join('')});
}).catch(error=>done({error:String(error)}));
""", args.built_asset.name)
                built_hash = digest(args.built_asset)
                assert loaded and loaded.get("sha256") == built_hash, (loaded, built_hash)
                served_build = {"server_pid": args.server_pid,
                    "server_command": args.server_command,
                    "built_asset": str(args.built_asset.resolve()),
                    "built_sha256": built_hash, "loaded_url": loaded["url"],
                    "loaded_sha256": loaded["sha256"]}
            cold_layout_shift = driver.execute_script("return window.__authoritiesSmoke.cls||0")
            assert cold_layout_shift <= 0.1, cold_layout_shift
            static_header = header_rect(driver)
            static_surface = workspace_rect(driver)
            keyboard_tabs(driver)
            automatic = None if args.manual_only else automatic_review(
                driver, source, pdfs, filing, output, static_header, mode == "beaver")
            manual = manual_book(driver, pdfs, output)
            viewports = viewport_proof(driver, output)
            urls = network_urls(driver)
            canlii_requests = [url for url in urls
                if (urlparse(url).hostname or "").lower().endswith("canlii.org")]
            assert not canlii_requests, canlii_requests
            severe = [entry for entry in driver.get_log("browser")
                      if entry.get("level") == "SEVERE"]
            assert not severe, severe
            result = {"schema_version": "beaver.authorities-browser-proof.v3",
                "created_at": datetime.now(timezone.utc).isoformat(), "mode": mode,
                "url": args.url, "artifacts": str(output), "cold_navigation": navigation,
                "browser_version": driver.capabilities.get("browserVersion"),
                "native_citation_parse": native_performance,
                "served_build": served_build,
                "cold_layout_shift": cold_layout_shift,
                "static_surface": static_surface, "inputs": inputs,
                "automatic": automatic,
                "manual": manual, "viewports": viewports, "canlii_requests": 0,
                "network_requests": len(urls), "browser_severe": [],
                "elapsed_seconds": round(time.monotonic() - started, 3)}
            (output / "result.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
            print(json.dumps(result, indent=2))
            return 0
        except Exception as caught:
            driver.save_screenshot(str(output / "failure.png"))
            failure = {"schema_version": "beaver.authorities-browser-proof.v3",
                "created_at": datetime.now(timezone.utc).isoformat(), "mode": mode,
                "url": args.url, "current_url": driver.current_url,
                "artifacts": str(output), "error": repr(caught),
                "browser_version": driver.capabilities.get("browserVersion"),
                "native_citation_parse": native_performance,
                "browser": driver.get_log("browser"),
                "import_timing": driver.execute_script(
                    "return JSON.parse(sessionStorage.getItem('authoritiesImportTiming')||'null')"),
                "requests": driver.execute_script(
                    "return JSON.parse(JSON.stringify(window.__authoritiesSmoke||null))"),
                "elapsed_seconds": round(time.monotonic() - started, 3)}
            (output / "failure.json").write_text(json.dumps(failure, indent=2), encoding="utf-8")
            raise
        finally:
            driver.quit()


if __name__ == "__main__":
    raise SystemExit(main())
