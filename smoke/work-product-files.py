"""Browser-free invented fixtures and independent PDF/DOCX output verification.

Citation identifiers refer to published decisions (SCR/CanLII/SCC/BCCA):
https://www.canlii.org/en/ca/scc/doc/2009/2009scc32/2009scc32.html
https://www.canlii.org/en/ca/scc/doc/2016/2016scc27/2016scc27.html
https://www.canlii.org/en/ca/scc/doc/1986/1986canlii46/1986canlii46.html
The 2026 SCC 16 and 2021 BCCA 222 identifiers exercise format recognition only.
All argument,
party, record and PDF prose is invented test content, not quotations from those decisions.
The source generators remain here; their per-run files are disposable test inputs."""

from __future__ import annotations

import hashlib, json, sys, zipfile

from pathlib import Path

from xml.sax.saxutils import escape

import fitz

from docx import Document


OAKES = "R v Oakes, [1986] 1 SCR 103, 1986 CanLII 46 (SCC)"

FALSE_POSITIVE = "2024 ABKB 999"

RECONSTRUCTED = "2021 BCCA 222"

REAL_CITATIONS = ("2009 SCC 32", "2016 SCC 27", "2026 SCC 16", RECONSTRUCTED)

def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def pdf_fixture(path: Path, title: str, detail: str) -> Path:
    pdf = fitz.open()
    page = pdf.new_page(width=612, height=792)
    page.insert_text((72, 86), title, fontsize=18)
    page.insert_textbox(fitz.Rect(72, 125, 540, 700), detail, fontsize=11)
    pdf.set_metadata({"title": title, "subject": detail})
    pdf.save(path)
    pdf.close()
    return path


def add_footnote(path: Path, text: str) -> None:
    with zipfile.ZipFile(path) as source:
        files = {name: source.read(name) for name in source.namelist()}
    document = files["word/document.xml"].decode()
    assert "[[FOOTNOTE]]" in document
    files["word/document.xml"] = document.replace(
        "<w:t>[[FOOTNOTE]]</w:t>", '<w:footnoteReference w:id="2"/>').encode()
    types = files["[Content_Types].xml"].decode()
    files["[Content_Types].xml"] = types.replace("</Types>",
        '<Override PartName="/word/footnotes.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/></Types>').encode()
    rels = files["word/_rels/document.xml.rels"].decode()
    files["word/_rels/document.xml.rels"] = rels.replace("</Relationships>",
        '<Relationship Id="rIdAuthoritiesFootnote" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes" Target="footnotes.xml"/></Relationships>').encode()
    files["word/footnotes.xml"] = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f'<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
        f'<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>'
        f'<w:footnote w:id="2"><w:p><w:r><w:footnoteRef/></w:r><w:r><w:t xml:space="preserve">{escape(text)}</w:t></w:r></w:p></w:footnote>'
        f'</w:footnotes>').encode()
    patched = path.with_name(f"{path.stem}-patched.docx")
    with zipfile.ZipFile(patched, "w", zipfile.ZIP_DEFLATED) as target:
        for name, content in files.items():
            target.writestr(name, content)
    patched.replace(path)


def fixtures(directory: Path) -> tuple[Path, list[Path], Path]:
    source = directory / "real-authorities-smoke.docx"
    document = Document()
    document.add_heading("Written argument", level=1)
    document.add_paragraph("In this invented observatory-planning dispute, the inspection argument cites R v Grant, 2009 SCC 32 at para 29.")
    document.add_paragraph(f"😀 The next invented argument cites {OAKES}.")
    document.add_paragraph(f"The observatory inventory marker {FALSE_POSITIVE} is not an authority.")
    document.add_paragraph("A separate invented argument cites Ahluwalia v Ahluwalia, 2026 SCC 16.")
    document.add_paragraph("Another invented argument cites Neufeld v Hansman, 2021 BCCA 222.")
    document.add_paragraph("😀 Invented observatory context. " + "The instrument inventory supplies invented context. " * 45 +
        "An invented deadline argument cites R v Jordan, 2016 SCC 27 at para 47. " +
        "The inspection request follows from this invented schedule. " * 45)
    document.add_paragraph("[[FOOTNOTE]]")
    document.save(source)
    add_footnote(source, f"See R v Grant, 2009 SCC 32; {OAKES}; Ibid at para 31; and R v Jordan, 2016 SCC 27.")
    pdfs = []
    for filename, title, citation in (
        ("2009scc32.pdf", "R v Grant", "2009 SCC 32\n\n[29] Invented source excerpt for the telescope inspection issue."),
        ("R v Oakes.pdf", "R v Oakes", "[1986] 1 SCR 103"),
        ("R v Jordan.pdf", "R v Jordan", "2016 SCC 27"),
        ("R v Grant replacement.pdf", "R v Grant", "2009 SCC 32\n\nReplacement original."),
        ("Hearing transcript.pdf", "Hearing transcript", "Supplemental book material."),
        ("Hearing transcript replacement.pdf", "Hearing transcript",
         "Replacement supplemental book material."),
        ("Authorities Smoke Cover.pdf", "Authorities Smoke Book", "Custom cover."),
    ):
        pdfs.append(pdf_fixture(directory / filename, title,
            f"{citation}\n\nValid local source PDF used by the Authorities production smoke."))
    filing = pdf_fixture(directory / "appeal-factum.pdf", "Appeal Factum",
        f"R v Grant, 2009 SCC 32\n\n{OAKES}\n\nR v Jordan, 2016 SCC 27")
    return source, pdfs, filing


def inspect_table(path: Path) -> dict[str, object]:
    with zipfile.ZipFile(path) as package:
        assert package.testzip() is None
        assert {"[Content_Types].xml", "word/document.xml"}.issubset(package.namelist())
    document = Document(path)
    text = "\n".join([paragraph.text for paragraph in document.paragraphs] +
        [cell.text for table in document.tables for row in table.rows for cell in row.cells])
    assert document.tables and "Table of Authorities" in text
    assert all(citation in text for citation in REAL_CITATIONS), text
    return {"file": path.name, "sha256": digest(path), "tables": len(document.tables)}


def link_page(link: dict[str, object]) -> int:
    try:
        return int(link.get("page", -1))
    except (TypeError, ValueError):
        return -1


def inspect_book(path: Path, preview: Path, *,
                 required_text: tuple[str, ...] = ("Book of Authorities", "Table of Contents"),
                 required_outline: tuple[str, ...] = ("Book of Authorities", "Table of Contents"),
                 minimum_links: int = 3) -> dict[str, object]:
    pdf = fitz.open(path)
    try:
        text = "\n".join(page.get_text() for page in pdf)
        toc = pdf.get_toc()
        outline = [item[1] for item in toc]
        links = [link for page in pdf for link in page.get_links() if link_page(link) >= 0]
        assert pdf.is_pdf and not pdf.is_encrypted and pdf.page_count >= 5
        assert all(value in text for value in required_text), text
        assert len(links) >= minimum_links and set(required_outline).issubset(set(outline))
        pdf[0].get_pixmap(matrix=fitz.Matrix(1.2, 1.2), alpha=False).save(preview)
        return {"file": path.name, "sha256": digest(path), "pages": pdf.page_count,
                "bookmarks": len(toc), "outline": outline, "links": len(links),
                "textCharacters": len(text),
                "textSha256": hashlib.sha256(text.encode("utf-8")).hexdigest()}
    finally:
        pdf.close()


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def make_pdf(path: Path, lines: list[str]) -> None:
    with fitz.open() as document:
        page = document.new_page(width=612, height=792)
        remaining = page.insert_textbox(
            fitz.Rect(72, 64, 540, 735), "\n".join(lines),
            fontsize=11, fontname="Times-Roman", lineheight=1.25,
        )
        assert remaining >= 0, f"Fixture text overflowed {path.name}."
        document.set_metadata({"title": lines[0], "author": "Beaver live proof"})
        document.save(path)


def make_transcript_pdf(path: Path) -> None:
    with fitz.open() as document:
        for text in ("TRANSCRIPT OF PROCEEDINGS", "TABLE OF CONTENTS", "ORAL PROCEEDINGS TRANSCRIPT"):
            page = document.new_page(width=612, height=792)
            page.insert_text((72, 72), text, fontsize=11, fontname="Times-Roman")
        document.set_page_labels([
            {"startpage": 0, "prefix": "", "style": "", "firstpagenum": 1},
            {"startpage": 1, "prefix": "", "style": "r", "firstpagenum": 1},
            {"startpage": 2, "prefix": "", "style": "D", "firstpagenum": 1},
        ])
        document.save(path)


def make_fixtures(directory: Path) -> dict[str, Path]:
    directory.mkdir()
    fixtures = {
        "affidavit": directory / "Affidavit-of-Morgan-Vale.pdf",
        "exhibit_b": directory / "Project-ledger.pdf",
        "ambiguous": directory / "Service-agreement.pdf",
        "notice": directory / "Notice-of-motion.pdf",
        "representations": directory / "Written-representations.pdf",
        "oral_request": directory / "Request-for-oral-hearing.docx",
        
        "appeal_pleading": directory / "Statement-of-Claim.pdf",
        "appeal_reasons": directory / "Reasons-for-Judgment.pdf",
        "appeal_order": directory / "Formal-Order.pdf",
        "appeal_notice": directory / "Notice-of-Appeal.pdf",
        "appeal_transcript_prior": directory / "Prior-Transcript.pdf",
        "appeal_transcript": directory / "Telescope-hearing-transcript.pdf",
    }
    make_pdf(fixtures["affidavit"], [
        "COURT OF KING'S BENCH OF ALBERTA",
        "COURT FILE NUMBER: 2601-80808",
        "JUDICIAL CENTRE: CALGARY",
        "BETWEEN:",
        "ORBITAL OBSERVATORY SOCIETY",
        "PLAINTIFF",
        "- and -",
        "SUMMIT ROOFTOP COOPERATIVE",
        "DEFENDANT",
        "2nd Affidavit of Morgan Vale",
        "I, Morgan Vale, of Calgary, Alberta, SWEAR AND SAY THAT:",
        "1. I maintain the telescope calibration log and recorded these measurements personally.",
        "2. The observatory permit issued for the rooftop instrument is attached and marked as Exhibit A.",
        "3. The telescope calibration log is attached and marked as Exhibit B.",
        "SWORN BEFORE ME AT THE CITY OF CALGARY",
        "IN THE PROVINCE OF ALBERTA",
        "THIS 29TH DAY OF AUGUST, 2026",
        "A Commissioner for Oaths in and for Alberta",
    ])
    make_pdf(fixtures["exhibit_b"], [
        'This is Exhibit "B" referred to in the Affidavit.',
        "TELESCOPE CALIBRATION LOG",
        "August 2026",
        "The calibration log records three lens readings taken during the inspection.",
    ])
    make_pdf(fixtures["ambiguous"], [
        "OBSERVATORY PERMIT",
        "Issued February 11, 2026",
        "The rooftop observatory permits one instrument inspection between 10:00 and 11:00.",
    ])
    make_pdf(fixtures["notice"], [
        "FEDERAL COURT",
        "COURT FILE NUMBER: T-501-26",
        "BETWEEN:",
        "ORBITAL OBSERVATORY SOCIETY",
        "APPLICANT",
        "- and -",
        "SUMMIT ROOFTOP COOPERATIVE",
        "RESPONDENT",
        "NOTICE OF MOTION",
        "The Applicant requests scheduling directions for an observatory permit inspection.",
    ])
    make_pdf(fixtures["representations"], [
        "WRITTEN REPRESENTATIONS OF THE APPLICANT",
        "The proposed inspection lasts one hour and preserves the telescope calibration records.",
    ])
    word = Document()
    word.add_heading("REQUEST FOR ORAL HEARING", level=1)
    word.add_paragraph("The moving party requests an oral hearing under Rule 369.2(2).")
    word.save(fixtures["oral_request"])
    for key, heading in (
        ("appeal_pleading", "STATEMENT OF CLAIM"),
        ("appeal_reasons", "REASONS FOR JUDGMENT"),
        ("appeal_order", "FORMAL ORDER"),
        ("appeal_notice", "NOTICE OF APPEAL"),
    ):
        metadata = (["Trial Court File Number: 2501-60606",
            "Decision Maker Appealed From: The Honourable Justice R. Example",
            "Decision Date: February 17, 2026", "Decision Filing Date: February 18, 2026"]
            if key == "appeal_reasons" else [])
        make_pdf(fixtures[key], [heading, "Court File Number 2603-70707",
                                 "Orbital Observatory Society v. Summit Rooftop Cooperative", *metadata])
    make_transcript_pdf(fixtures["appeal_transcript_prior"])
    make_transcript_pdf(fixtures["appeal_transcript"])
    return fixtures


def save_pdf_page(document: fitz.Document, page_number: int, output: Path) -> None:
    document[page_number].get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False).save(output)


def inspect_affidavit(path: Path, output: Path) -> dict[str, object]:
    with fitz.open(path) as document:
        pages = [page.get_text() for page in document]
        toc = [row[1] for row in document.get_toc()]
        assert document.page_count == 4, pages
        assert [page.get_label() for page in document] == ["1", "2", "3", "4"]
        assert not document.is_encrypted and all(text.strip() for text in pages)
        assert 'This is Exhibit "A" referred to in the Affidavit of:' in pages[1]
        assert "OBSERVATORY PERMIT" in pages[2]
        assert 'This is Exhibit "B" referred to in the Affidavit.' in pages[3]
        assert "TELESCOPE CALIBRATION LOG" in pages[3]
        assert toc.count("Exhibit A certificate") == 1, toc
        assert "Exhibit B certificate" not in toc, toc
        for page_number, name in ((0, "affidavit-output"), (1, "exhibit-a-certificate"),
                                  (3, "exhibit-b-with-certificate")):
            save_pdf_page(document, page_number, output / f"{name}.png")
        return {"filename": path.name, "sha256": sha256(path), "bytes": path.stat().st_size,
                "pages": document.page_count, "bookmarks": toc}


def inspect_federal(
    path: Path, output: Path, slug: str = "federal-motion",
    court: str = "FEDERAL COURT", pages_expected: int = 4,
    extra_text: tuple[str, ...] = (),
) -> dict[str, object]:
    with fitz.open(path) as document:
        text = "\n".join(page.get_text() for page in document)
        toc = [row[1] for row in document.get_toc()]
        links = sum(len(page.get_links()) for page in document)
        assert document.page_count == pages_expected, text
        assert [page.get_label() for page in document] == [
            str(value) for value in range(1, pages_expected + 1)
        ]
        assert not document.is_encrypted
        for value in (court, "MOTION RECORD", "INDEX", "NOTICE OF MOTION",
                      "WRITTEN REPRESENTATIONS OF THE APPLICANT", *extra_text):
            assert value in text, value
        assert links >= 2 and len(toc) >= 3, (links, toc)
        for page_number, name in ((0, f"{slug}-cover"), (1, f"{slug}-index")):
            save_pdf_page(document, page_number, output / f"{name}.png")
        return {"filename": path.name, "sha256": sha256(path), "bytes": path.stat().st_size,
                "pages": document.page_count, "bookmarks": toc, "internal_links": links}


def inspect_appeal_record(
    record_path: Path, transcript_path: Path, source_transcript: Path, output: Path
) -> dict[str, object]:
    with fitz.open(record_path) as record, fitz.open(transcript_path) as transcript:
        pages = [page.get_text() for page in record]
        text = "\n".join(pages)
        toc = [row[1] for row in record.get_toc()]
        links = sum(len(page.get_links()) for page in record)
        corner = record[0].get_pixmap(alpha=False).pixel(5, 5)[:3]
        index_text = " ".join(pages[1].split())
        assert record.page_count == 6, pages
        assert [page.get_label() for page in record] == [str(value) for value in range(1, 7)]
        assert corner[0] > 240 and corner[1] < 15 and corner[2] < 15, corner
        for value in ("COURT OF APPEAL OF ALBERTA", "APPEAL RECORD", "Table of Contents",
                      "STATEMENT OF CLAIM", "REASONS FOR JUDGMENT", "FORMAL ORDER",
                      "NOTICE OF APPEAL"):
            assert value in text, value
        cover_text = " ".join(pages[0].split())
        for value in ("Rowan Counsel", "403-555-0191", "403-555-0192", "Avery Advocate",
                      "620 Telescope Avenue", "780-555-0193"):
            assert value in cover_text, (value, cover_text)
        for description in ("STATEMENT OF CLAIM", "REASONS FOR JUDGMENT", "Formal Order",
                            "NOTICE OF APPEAL"):
            assert description in index_text, (description, index_text)
        assert "Part 2 - Formal order or decision" not in index_text, index_text
        assert index_text.count("Part 2") == 1, index_text
        assert "ORAL PROCEEDINGS TRANSCRIPT" not in text
        assert links >= 4 and len(toc) >= 6, (links, toc)
        assert transcript.page_count == 3 and not transcript.is_encrypted
        assert [page.get_label() for page in transcript] == ["", "i", "1"]
        assert "ORAL PROCEEDINGS TRANSCRIPT" in transcript[2].get_text()
        assert transcript_path.name == source_transcript.name
        assert sha256(transcript_path) == sha256(source_transcript)
        save_pdf_page(record, 0, output / "appeal-record-cover.png")
        save_pdf_page(record, 1, output / "appeal-record-index.png")
        return {
            "record": {"filename": record_path.name, "sha256": sha256(record_path),
                       "pages": record.page_count, "bookmarks": toc,
                       "internal_links": links, "cover_rgb": list(corner),
                       "filename_descriptions": True},
            "transcript": {"filename": transcript_path.name,
                           "sha256": sha256(transcript_path), "pages": transcript.page_count,
                           "page_labels": [page.get_label() for page in transcript],
                           "preserved_byte_for_byte": True},
        }


def inspect_no_oral_record(path: Path, description: str, output: Path) -> dict[str, object]:
    with fitz.open(path) as document:
        text = "\n".join(page.get_text() for page in document)
        assert document.page_count == 6
        assert " ".join(description.split()) in " ".join(text.split()), text
        assert "ORAL PROCEEDINGS TRANSCRIPT" not in text
        save_pdf_page(document, 1, output / "appeal-record-no-oral-record-index.png")
        return {"filename": path.name, "sha256": sha256(path),
                "pages": document.page_count, "description": description}


def inspect_highlights(path: Path, output: Path) -> dict[str, object]:
    # PyMuPDF independently reads, edits and reopens the real Node browser export.
    with fitz.open(path) as document:
        original_text = [page.get_text() for page in document]
        annotations = [(page, annotation) for page in document for annotation in page.annots() or []]
        assert len(annotations) == 4
        assert all(annotation.type[0] == fitz.PDF_ANNOT_HIGHLIGHT for _, annotation in annotations)
        assert any("Uncited passage" in annotation.info.get("content", "") for _, annotation in annotations)
        assert list(document[-1].annots() or []), "The scanned page lost its manually drawn annotation."
        page, annotation = annotations[0]
        page.delete_annot(annotation)
        target = output / "annotation-roundtrip.pdf"
        document.save(target)
    with fitz.open(target) as reread:
        assert sum(len(list(page.annots() or [])) for page in reread) == 3
        assert [page.get_text() for page in reread] == original_text
    return {"editable_highlights": 4, "after_independent_delete": 3, "source_text_unchanged": True}


if __name__ == "__main__":
    action, directory, *paths = sys.argv[1:]
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    if action == "authorities-fixtures":
        source, pdfs, filing = fixtures(directory)
        result = {"source": str(source), "pdfs": list(map(str, pdfs)), "filing": str(filing)}
    elif action == "court-fixtures":
        result = {key: str(value) for key, value in make_fixtures(directory / "inputs").items()}
    elif action == "highlights": result = inspect_highlights(Path(paths[0]), directory)
    elif action == "table": result = inspect_table(Path(paths[0]))
    elif action == "book": result = inspect_book(Path(paths[0]), directory / "book.png")
    elif action == "affidavit": result = inspect_affidavit(Path(paths[0]), directory)
    elif action == "fc": result = inspect_federal(Path(paths[0]), directory)
    elif action == "fca-standalone": result = inspect_federal(Path(paths[0]), directory, "fca", "FEDERAL COURT OF APPEAL")
    elif action == "fca": result = inspect_federal(Path(paths[0]), directory, "fca", "FEDERAL COURT OF APPEAL", 5, ("REQUEST FOR ORAL HEARING",))
    elif action == "appeal": result = inspect_appeal_record(Path(paths[0]), Path(paths[1]), Path(paths[2]), directory)
    elif action == "no-oral": result = inspect_no_oral_record(Path(paths[0]), paths[1], directory)
    else: raise ValueError(action)
    print(json.dumps(result))
